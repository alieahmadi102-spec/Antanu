# -*- coding: utf-8 -*-
"""
extractors.py — استخراج متن از فرمت‌های رایج کتاب.

پشتیبانی: PDF، Word (docx)، متن (txt/md)، پاورپوینت (pptx)، اکسل (xlsx).

هر فرمت داخل try/except است: فایل خراب یا پسوند ناشناخته رشته‌ی خالی می‌دهد و
هیچ‌وقت برنامه را نمی‌شکند.

منطق پاک‌سازی و سنجش کیفیتِ متن فارسی از همان ابزار آزموده‌ی
tools/ingest_library.py می‌آید (اصلاح شکل‌های نمایشی عربی، رد متنی که واژه‌هایش
به هم چسبیده‌اند)، تا PDFهای اسکن‌شده و استخراج‌های بی‌کیفیت خودکار کنار بروند.
"""
import os
import sys

_TOOLS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools")
if _TOOLS not in sys.path:
    sys.path.insert(0, _TOOLS)

try:                                  # بازاستفاده از منطق فارسیِ موجود
    from ingest_library import clean, text_is_usable
except Exception:                     # اگر ابزار نبود، کمینه‌ی لازم را خودمان داریم
    import re
    import unicodedata

    def clean(text: str) -> str:
        if not text:
            return ""
        t = unicodedata.normalize("NFKC", text)
        t = t.replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
        t = re.sub(r"[ \t]+", " ", t)
        return re.sub(r"\n{3,}", "\n\n", t).strip()

    def text_is_usable(text: str):
        return (len((text or "").strip()) >= 400), "بررسی ساده"


SUPPORTED = (".pdf", ".docx", ".txt", ".md", ".pptx", ".xlsx")

# صفحه‌ای که کمتر از این نویسه متن دارد، عملاً تصویر است و رد می‌شود
_MIN_PAGE_CHARS = 40


def _from_pdf(path: str) -> str:
    """PDF — صفحه‌به‌صفحه؛ صفحه‌های بدون متن (اسکن) نادیده گرفته می‌شوند."""
    import pypdf
    reader = pypdf.PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            txt = (page.extract_text() or "").strip()
        except Exception:
            continue
        if len(txt) >= _MIN_PAGE_CHARS:
            parts.append(txt)
    return "\n\n".join(parts)


def _from_docx(path: str) -> str:
    """Word — پاراگراف‌ها به‌علاوه‌ی متن جدول‌ها."""
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_txt(path: str) -> str:
    """متن ساده — اول UTF-8، بعد cp1256 (متن‌های قدیمی فارسی/عربی)."""
    for enc in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


def _from_pptx(path: str) -> str:
    """پاورپوینت — متن همه‌ی شکل‌ها و جدول‌ها در همه‌ی اسلایدها."""
    from pptx import Presentation
    pres = Presentation(path)
    parts = []
    for slide in pres.slides:
        for shape in slide.shapes:
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
            except Exception:
                continue
    return "\n".join(parts)


def _from_xlsx(path: str) -> str:
    """اکسل — همه‌ی سلول‌های همه‌ی برگه‌ها، سطر به سطر."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"[{ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    parts.append(" ".join(vals))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n".join(parts)


_READERS = {
    ".pdf": _from_pdf,
    ".docx": _from_docx,
    ".txt": _from_txt,
    ".md": _from_txt,
    ".pptx": _from_pptx,
    ".xlsx": _from_xlsx,
}


def is_supported(filename: str) -> bool:
    return os.path.splitext(filename or "")[1].lower() in SUPPORTED


def extract_text(filepath: str) -> str:
    """متن را از هر فرمت پشتیبانی‌شده استخراج می‌کند.

    اگر فایل خراب باشد، پسوندش پشتیبانی نشود، یا متنِ به‌دردبخوری نداشته باشد
    (مثل PDF اسکن‌شده یا استخراجی که واژه‌ها را به هم چسبانده)، رشته‌ی خالی
    برمی‌گردد — بدون استثنا و بدون کرش.
    """
    try:
        ext = os.path.splitext(filepath or "")[1].lower()
        reader = _READERS.get(ext)
        if not reader or not os.path.exists(filepath):
            return ""
        raw = reader(filepath) or ""
    except Exception:
        return ""
    text = clean(raw)
    if not text:
        return ""
    try:
        ok, _why = text_is_usable(text)
    except Exception:
        ok = len(text) >= 400
    return text if ok else ""
