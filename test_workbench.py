# -*- coding: utf-8 -*-
"""
test_workbench.py — آزمون میزکارهای نرم‌افزاری (SPSS / EViews / SmartPLS)

چه چیزهایی سنجیده می‌شود:
  ۱) درستی پیکربندی منوها: هر تحلیلِ منو واقعاً در موتورها موجود و مجاز باشد
  ۲) شبکه‌ی داده: ذخیره/خواندن رفت‌وبرگشتی با ستون فارسی و نام تکراری
  ۳) صفحه‌های /spss /eviews /smartpls برای کاربر واردشده باز شوند
  ۴) فهرست داده‌ها: آپلود ثبت شود؛ داده‌ی کاربرِ دیگر قابل خواندن نباشد (۴۰۳)
  ۵) اجرای تحلیل از مسیر واقعی API (ANOVA و PLS) و ساخت نمودار PLS از کلیدهای فارسی

اجرا:  python test_workbench.py
"""
import os
import sys
import warnings

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.chdir(os.path.dirname(os.path.abspath(__file__)))

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  ✅ " if cond else "  ❌ ") + name + (f"  — {detail}" if detail else ""))


def section(title):
    print("\n" + "─" * 68 + f"\n{title}\n" + "─" * 68)


def _test_user():
    import db as D
    conn = D.get_db()
    u = conn.execute("SELECT id FROM users WHERE is_admin = 0 LIMIT 1").fetchone()
    if u:
        conn.execute("DELETE FROM usage_log WHERE user_id = ?", (u["id"],))
        conn.commit()
    conn.close()
    return u


def _sample_grid():
    """داده‌ی نمونه با سه گروه واقعی و چهار گویه‌ی لیکرت."""
    import random
    random.seed(11)
    cols = ["گروه", "نمره", "q1", "q2", "q3", "q4"]
    rows = []
    for i in range(45):
        g = ["الف", "ب", "ج"][i % 3]
        base = {"الف": 12, "ب": 14, "ج": 16}[g]
        q1 = random.randint(2, 5)
        rows.append([g, round(base + random.gauss(0, 1.5), 2), q1,
                     min(5, max(1, q1 + random.choice([-1, 0, 1]))),
                     random.randint(2, 5), random.randint(2, 5)])
    return cols, rows


# ═══════════════ ۱) درستی پیکربندی منوها ═══════════════

def test_config_integrity():
    section("۱) هر تحلیلِ منو در موتورها موجود و برای همان میزکار مجاز است")
    import workbench, stats_engine
    available = set(stats_engine.available_analyses())

    def leaves(items):
        for it in items:
            if it.get("children"):
                yield from leaves(it["children"])
            elif it.get("analysis"):
                yield it

    bad_missing, bad_disallowed, bad_fields = [], [], []
    for sw_id, sw in workbench.SOFTWARES.items():
        for menu in sw["menus"]:
            for leaf in leaves(menu.get("items", [])):
                a = leaf["analysis"]
                if a not in available:
                    bad_missing.append(f"{sw_id}:{a}")
                if a not in workbench.ALLOWED[sw_id]:
                    bad_disallowed.append(f"{sw_id}:{a}")
                for f in leaf.get("fields", []):
                    if f.get("type") not in ("col", "cols", "select", "number", "text", "check"):
                        bad_fields.append(f"{sw_id}:{a}:{f.get('key')}")
    check("همه‌ی تحلیل‌های منو در موتورها موجودند", not bad_missing, str(bad_missing))
    check("همه‌ی تحلیل‌های منو در فهرست مجاز همان میزکارند", not bad_disallowed, str(bad_disallowed))
    check("نوع همه‌ی فیلدهای پنجره‌ها معتبر است", not bad_fields, str(bad_fields))
    check("سه میزکار تعریف شده‌اند", set(workbench.SOFTWARES) == {"spss", "eviews", "smartpls"})


# ═══════════════ ۲) شبکه‌ی داده: ذخیره/خواندن ═══════════════

def test_grid_roundtrip():
    section("۲) ذخیره و خواندن شبکه‌ی داده (ستون فارسی + نام تکراری)")
    import tempfile
    import workbench
    p = os.path.join(tempfile.gettempdir(), "wb-roundtrip.csv")
    cols, rows = _sample_grid()
    info = workbench.save_grid(p, cols, rows)
    check("ذخیره انجام شد", info["rows"] == 45 and info["columns"] == 6)
    g = workbench.load_grid(p)
    check("تعداد سطر و ستون بعد از خواندن درست است",
          g["total_rows"] == 45 and len(g["columns"]) == 6)
    names = [c["name"] for c in g["columns"]]
    check("نام ستون فارسی سالم ماند", "گروه" in names and "نمره" in names, str(names))
    types = {c["name"]: c["type"] for c in g["columns"]}
    check("نوع ستون‌ها درست تشخیص داده شد",
          types.get("گروه") == "text" and types.get("نمره") == "numeric")
    check("آماره‌های Variable View ساخته شدند",
          all(k in g["columns"][1] for k in ("mean", "std", "min", "max")))

    # نام تکراری ستون باید خودکار یکتا شود
    workbench.save_grid(p, ["x", "x", "x"], [[1, 2, 3]])
    g2 = workbench.load_grid(p)
    n2 = [c["name"] for c in g2["columns"]]
    check("نام‌های تکراری ستون یکتا شدند", len(set(n2)) == 3, str(n2))

    # سقف‌ها
    try:
        workbench.save_grid(p, ["a"], [[1]] * (workbench.MAX_SAVE_ROWS + 1))
        check("سقف سطرها اعمال می‌شود", False)
    except ValueError:
        check("سقف سطرها اعمال می‌شود", True)


# ═══════════════ ۳) صفحه‌ها و API با کاربر واقعی ═══════════════

def test_pages_and_api():
    section("۳) صفحه‌های میزکار و APIهای شبکه/فهرست داده")
    from fastapi.testclient import TestClient
    import main

    u = _test_user()
    if not u:
        print("  ⏭  کاربر آزمایشی پیدا نشد — این بخش رد شد.")
        return
    c = TestClient(main.app)

    r = c.get("/spss", follow_redirects=False)
    check("بدون ورود، به صفحه‌ی ورود می‌فرستد", r.status_code in (302, 303, 307))

    c.cookies.set("antanu_session", main.make_session(u["id"]))
    for path, name in (("/spss", "ANTANU SPSS"), ("/eviews", "ANTANU EViews"),
                       ("/smartpls", "ANTANU SmartPLS")):
        r = c.get(path)
        check(f"صفحه‌ی {path} باز می‌شود", r.status_code == 200 and name in r.text
              and "workbench.js" in r.text, str(r.status_code))

    # ذخیره‌ی شبکه → فهرست → خواندن
    cols, rows = _sample_grid()
    r = c.post("/api/workbench/save",
               json={"title": "داده‌ی آزمون میزکار", "columns": cols, "rows": rows})
    check("ذخیره‌ی شبکه از API", r.status_code == 200 and r.json().get("file"), r.text[:150])
    fname = r.json()["file"]

    r = c.get("/api/workbench/datasets")
    ds = r.json().get("datasets", [])
    check("داده‌ی ذخیره‌شده در فهرست هست",
          any(d["file"] == fname and d["source"] == "edit" for d in ds))

    r = c.get("/api/workbench/data", params={"file": fname})
    check("خواندن شبکه از API", r.status_code == 200 and r.json()["total_rows"] == 45)

    r = c.get("/api/workbench/data", params={"file": "../../etc/passwd"})
    check("نام فایل خطرناک رد می‌شود", r.status_code == 400)

    # مالکیت: کاربر دیگر نتواند بخواند
    import db as D
    conn = D.get_db()
    other = conn.execute("SELECT id FROM users WHERE id != ? AND is_admin = 0 LIMIT 1",
                         (u["id"],)).fetchone()
    if not other:
        import secrets as _s
        from db import hash_pw
        conn.execute("INSERT INTO users (username, password, stars) VALUES (?, ?, 1)",
                     ("wb_other_" + _s.token_hex(3), hash_pw("x12345"), ))
        other = conn.execute("SELECT id FROM users ORDER BY id DESC LIMIT 1").fetchone()
        conn.commit()
    conn.close()
    c2 = TestClient(main.app)
    c2.cookies.set("antanu_session", main.make_session(other["id"]))
    r = c2.get("/api/workbench/data", params={"file": fname})
    check("داده‌ی کاربر دیگر خوانده نمی‌شود (۴۰۳)", r.status_code == 403, str(r.status_code))

    return fname


# ═══════════════ ۴) اجرای تحلیل از مسیر واقعی ═══════════════

def test_run_analyses(fname):
    section("۴) اجرای ANOVA و PLS از مسیر واقعی /api/stats/run")
    if not fname:
        print("  ⏭  فایل داده‌ای از بخش قبل نیست — رد شد.")
        return
    from fastapi.testclient import TestClient
    import main

    u = _test_user()

    async def fake_ai(c, prompt=None, system=None, max_tokens=1500, **kw):
        return "تفسیر آزمایشی: اثر معنادار است."
    main._call_model_once = fake_ai
    # حتی اگر در پایگاه‌داده کلید سرویس هوش مصنوعی نباشد، مسیر تفسیر آزموده شود
    main.pick_model_for = lambda task: {"id": "t", "key": "k", "model": "m",
                                        "base": "https://example.invalid/v1", "name": "آزمایشی"}

    c = TestClient(main.app)
    c.cookies.set("antanu_session", main.make_session(u["id"]))

    r = c.post("/api/stats/run", json={"file": fname, "analysis": "anova",
                                       "params": {"dependent": "نمره", "factor": "گروه"}})
    d = r.json()
    res = d.get("result") or {}
    check("ANOVA بدون خطا اجرا شد", r.status_code == 200 and not res.get("error"),
          str(res.get("error", ""))[:120])
    check("تفسیر دانشگاهی برگشت", "آزمایشی" in (d.get("interpretation") or ""))

    r = c.post("/api/stats/run", json={
        "file": fname, "analysis": "pls_sem",
        "params": {"factors": {"کیفیت": ["q1", "q2"], "رضایت": ["q3", "q4"]},
                   "structural": [["کیفیت", "رضایت"]], "bootstrap": 60}})
    res = (r.json().get("result") or {})
    check("PLS-SEM اجرا شد و سازه‌ها/مسیرها را داد",
          "سازه‌ها" in res and "مسیرهای ساختاری" in res, str(res.get("error", ""))[:120])
    boot = res.get("بوت‌استرپ") or {}
    check("بوت‌استرپ واقعاً اجرا شد", bool(boot), str(list(res.keys()))[:150])
    try:
        import matplotlib  # noqa: F401
        check("نمودار مسیر PLS از کلیدهای فارسی ساخته شد", bool(res.get("diagram")),
              str(list(res.keys()))[:150])
    except ImportError:
        print("  ⏭  matplotlib نصب نیست — آزمون نمودار رد شد (روی سرور اصلی نصب است).")

    r = c.post("/api/stats/run", json={"file": fname, "analysis": "unit_root",
                                       "params": {"cols": ["نمره"]}})
    res = (r.json().get("result") or {})
    check("آزمون ریشه‌ی واحد (EViews) اجرا شد", "نتایج" in res, str(res.get("error", ""))[:120])


# ═══════════════ ۵) بخش‌های کشویی پنل مدیریت ═══════════════

def test_admin_accordion():
    section("۵) پنل مدیریت: هر بخش کشویی است")
    from fastapi.testclient import TestClient
    import main
    import db as D

    conn = D.get_db()
    a = conn.execute("SELECT id FROM users WHERE is_admin = 1 LIMIT 1").fetchone()
    conn.close()
    if not a:
        print("  ⏭  کاربر مدیر پیدا نشد — این بخش رد شد.")
        return
    c = TestClient(main.app)
    c.cookies.set("antanu_session", main.make_session(a["id"]))
    r = c.get("/admin")
    check("پنل مدیریت باز می‌شود", r.status_code == 200)
    html = r.text
    check("اسکریپت کشویی‌کردن بخش‌ها هست", "card-head" in html and "card-body" in html)
    check("جعبه‌ی جستجوی بخش هست", "accSearch" in html)
    check("دکمه‌های باز/بستن همه هستند", "accOpenAll" in html and "accCloseAll" in html)
    check("وضعیت باز/بسته در مرورگر ذخیره می‌شود", "antanu_admin_open" in html)

    css = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "static", "style.css"), encoding="utf-8").read()
    check("بدنه‌ی بخشِ بسته پنهان است", ".card.acc > .card-body { display: none;" in css)
    check("بدنه‌ی بخشِ باز نمایان است", ".card.acc.open > .card-body { display: block; }" in css)

    # همه‌ی کارت‌های پنل عنوان دارند، وگرنه کشویی نمی‌شوند
    tpl = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "templates", "admin.html"), encoding="utf-8").read()
    import re as _re
    cards = _re.findall(r'<div class="card"[^>]*>(.*?)(?=\n  <!--|\n<script)', tpl, _re.DOTALL)
    without_h2 = [i for i, body in enumerate(cards) if "<h2>" not in body]
    check("همه‌ی کارت‌های پنل عنوان <h2> دارند", not without_h2, str(without_h2))


# ═══════════════ ۶) راهنمای چندزبانه ═══════════════

def test_help_multilang():
    section("۶) راهنمای استفاده و ترجمه‌اش به زبان‌های آنتانو")
    import asyncio
    from fastapi.testclient import TestClient
    import main, help_docs, i18n

    # متن‌های دستی موجود و به‌روزند
    fa, fa_lang = help_docs.source_text("fa")
    en, en_lang = help_docs.source_text("en")
    check("راهنمای فارسی موجود است", fa_lang == "fa" and len(fa) > 3000)
    check("راهنمای انگلیسی موجود است", en_lang == "en" and len(en) > 3000)
    for name, txt in (("فارسی", fa), ("انگلیسی", en)):
        check(f"راهنمای {name} نرم‌افزارهای تازه را دارد",
              all(k in txt for k in ("SPSS", "EViews", "SmartPLS")))
    check("راهنمای فارسی قاعده‌ی پانوشتِ نام لاتین را توضیح داده",
          "Times New Roman" in fa and "لاتین" in fa)

    # زبان بدون فایل دستی → مبدأ انگلیسی است، نه فارسی
    src, src_lang = help_docs.source_text("ja")
    check("مبدأ ترجمه برای زبان‌های دیگر، انگلیسی است", src_lang == "en")

    # تکه‌بندی: هیچ تکه‌ای نباید خالی یا غول‌آسا باشد
    chunks = help_docs.split_chunks(en)
    check("راهنما برای ترجمه تکه‌تکه می‌شود", len(chunks) >= 3)
    check("هیچ تکه‌ای بیش از حد بزرگ نیست",
          all(len(c) < help_docs.CHUNK_CHARS * 2.2 for c in chunks),
          str([len(c) for c in chunks]))
    check("جمعِ تکه‌ها چیزی از متن را نینداخته",
          sum(len(c) for c in chunks) >= len(en) - len(chunks) * 2)

    # همه‌ی زبان‌های سایت نام ترجمه‌ای دارند
    missing = [c for c in i18n.LANGUAGES if c not in help_docs.LANG_NAMES]
    check("برای همه‌ی ۱۲ زبان سایت، نام مقصد تعریف شده", not missing, str(missing))

    # اثرانگشت: با تغییر متن، ترجمه‌ی انباری باطل می‌شود
    k1 = help_docs.cache_key("tr")
    check("کلید انبار شاملِ اثرانگشتِ متن است", help_docs.fingerprint() in k1)

    # صفحه: فارسی و انگلیسی بدون هیچ فراخوانیِ هوش مصنوعی سرو شوند
    calls = []

    async def fake_fb(task, prompt=None, system=None, max_tokens=1800, messages=None,
                      keep_foreign=False):
        calls.append(task)
        body = prompt.split("--- DOCUMENT START ---")[1].split("--- DOCUMENT END ---")[0]
        return "```markdown\n[TR]" + body.strip() + "\n```"
    main._call_model_with_fallback = fake_fb

    c = TestClient(main.app)
    r = c.get("/help", headers={"Accept-Language": "fa"})
    check("راهنمای فارسی باز می‌شود", r.status_code == 200 and "نرم‌افزارهای آماری" in r.text)
    r = c.get("/help", headers={"Accept-Language": "en"})
    check("راهنمای انگلیسی باز می‌شود",
          r.status_code == 200 and "Statistical software inside ANTANU" in r.text)
    check("برای فارسی/انگلیسی هیچ ترجمه‌ای فراخوانی نشد", not calls, str(calls))

    # ترجمه‌ی واقعی یک زبان: انبار می‌شود و جعبه‌ی کد پاک می‌شود
    db = main.get_db()
    db.execute("DELETE FROM settings WHERE key LIKE 'helpdoc:%'")
    db.commit(); db.close()
    txt = asyncio.run(main.translate_help_doc("tr"))
    check("ترجمه انجام شد", len(txt) > 3000 and "[TR]" in txt)
    check("جعبه‌ی کدِ دورِ خروجی حذف شد", "```markdown" not in txt)
    check("ترجمه در پایگاه‌داده انبار شد",
          len(main.get_setting(help_docs.cache_key("tr"), "")) > 3000)
    r = c.get("/help", headers={"Accept-Language": "tr"})
    check("نسخه‌ی انباری بدون ترجمه‌ی دوباره سرو می‌شود", "[TR]" in r.text)

    # شکستِ سرویس: متن مبدأ نشان داده شود و چیزی انبار نشود
    async def fail_fb(task, prompt=None, system=None, max_tokens=1800, messages=None,
                      keep_foreign=False):
        raise main.ModelError(402, "insufficient balance")
    main._call_model_with_fallback = fail_fb
    out = asyncio.run(main.translate_help_doc("ja"))
    check("اگر ترجمه شکست بخورد، متن مبدأ برمی‌گردد (نه صفحه‌ی خالی)", len(out) > 3000)
    check("ترجمه‌ی ناقص انبار نمی‌شود",
          main.get_setting(help_docs.cache_key("ja"), "") == "")

    db = main.get_db()
    db.execute("DELETE FROM settings WHERE key LIKE 'helpdoc:%'")
    db.commit(); db.close()

    # اندپوینت‌های مدیریت
    import db as D
    conn = D.get_db()
    a = conn.execute("SELECT id FROM users WHERE is_admin = 1 LIMIT 1").fetchone()
    conn.close()
    if a:
        ca = TestClient(main.app)
        ca.cookies.set("antanu_session", main.make_session(a["id"]))
        r = ca.get("/admin/help/status")
        st = r.json()
        check("وضعیت ترجمه‌ها برای همه‌ی زبان‌ها گزارش می‌شود",
              r.status_code == 200 and len(st["langs"]) == len(i18n.LANGUAGES))
        check("فارسی و انگلیسی «دستی» شمرده می‌شوند",
              all(l["state"] == "written" for l in st["langs"] if l["code"] in ("fa", "en")))
        r = ca.post("/admin/help/clear", json={})
        check("پاک‌کردن ترجمه‌ها کار می‌کند", r.status_code == 200 and r.json()["ok"])
        rn = TestClient(main.app).get("/admin/help/status")
        check("کاربر غیرمدیر به وضعیت ترجمه دسترسی ندارد", rn.status_code in (401, 403, 302, 307))


# ═══════════════ ۷) نوار ورودی هرگز از صفحه بیرون نرود ═══════════════

def test_composer_never_pushed_offscreen():
    section("۷) با تیک «مقاله بلند»، دکمه‌ی ارسال از پایین صفحه بیرون نیفتد")
    here = os.path.dirname(os.path.abspath(__file__))
    css = open(os.path.join(here, "static", "style.css"), encoding="utf-8").read()
    js = open(os.path.join(here, "static", "app.js"), encoding="utf-8").read()

    # ریشه‌ی باگ: در فلکس‌باکس، ناحیه‌ی پیام‌ها زیر اندازه‌ی محتوایش کوچک نمی‌شد
    check("ناحیه‌ی پیام‌ها اجازه‌ی کوچک‌شدن دارد (min-height:0)",
          "#messages { flex: 1; min-height: 0;" in css)
    check("ظرفِ اصلی هم اجازه‌ی کوچک‌شدن دارد",
          "main { flex: 1; display: flex; flex-direction: column; min-width: 0; min-height: 0; }" in css)

    # سقفِ ارتفاع پنل — تضمین می‌کند فوتر بلندتر از صفحه نشود
    panel = css.split("#ctypePanel {", 1)[1].split("}", 1)[0]
    check("پنل نوع محتوا سقف ارتفاع دارد", "max-height" in panel, panel[:80])
    check("پنل به‌جای بزرگ‌شدن، خودش اسکرول می‌شود", "overflow-y: auto" in panel)
    check("سقف پنل با ارتفاع صفحه سنجیده می‌شود (vh/dvh)",
          "vh" in panel or "dvh" in panel)

    chips = css.split("#chips {", 1)[1].split("}", 1)[0]
    check("چیپ‌های پیوست هم سقف و اسکرول دارند",
          "max-height" in chips and "overflow-y: auto" in chips)

    check("سرِ پنل هنگام اسکرول می‌ماند (دکمه‌ی پاک‌کردن در دسترس)",
          "#ctypePanel .ctype-head" in css and "position: sticky" in css)
    check("تنظیمات خروجیِ تازه‌ظاهرشده خودکار جلوی چشم می‌آید",
          "scrollIntoView" in js and "ctypeDocOpts" in js)


# ═══════════════ ۸) میزکارها روی گوشی ═══════════════

def test_workbench_mobile_and_autobuild():
    section("۸) میزکارها روی گوشی + ساخت خودکار سازه‌ها")
    here = os.path.dirname(os.path.abspath(__file__))
    css = open(os.path.join(here, "static", "workbench.css"), encoding="utf-8").read()
    js = open(os.path.join(here, "static", "workbench.js"), encoding="utf-8").read()
    tpl = open(os.path.join(here, "templates", "workbench.html"), encoding="utf-8").read()

    # باگ: روی گوشی پنل کناری مخفی بود و هیچ راهی برای بازکردنش نبود
    check("دکمه‌ی باز/بستن پنل در نوار ابزار هست", 'data-act="panel"' in tpl)
    check("دکمه فقط روی صفحه‌های باریک دیده می‌شود",
          ".wb-tool.panel-toggle { display: none; }" in css
          and ".wb-tool.panel-toggle { display: inline-block; }" in css)
    check("پنل روی گوشی با کلاس show باز می‌شود",
          ".wb-side.show, .pls-panel.show { display: flex; }" in css)
    check("دکمه‌ی بستنِ پنل هم هست", 'class="panel-close"' in tpl)
    check("جاوااسکریپت پنل را باز/بسته می‌کند", "function togglePanel" in js)

    # ساخت خودکار سازه‌ها از روی نام ستون‌ها
    check("دکمه‌ی ساخت خودکار سازه‌ها هست", 'id="plsAuto"' in tpl)
    check("منطق تشخیص سازه از نام ستون هست", "function guessConstructs" in js)
    check("راهنمای «سازه یعنی چه» داخل پنل هست",
          "سازه یعنی چه" in tpl and "گویه" in tpl)

    # همان قاعده‌ی سرور و مرورگر باید یک نتیجه بدهد
    import analysis_planner
    cols = ["ID"] + [f"FA{i}" for i in range(1, 6)] + [f"IA{i}" for i in range(1, 9)] \
        + [f"FT{i}" for i in range(1, 4)] + ["جنسیت"]
    g = analysis_planner.guess_constructs(cols)
    check("سازه‌ها از نام ستون‌های پرسشنامه درست تشخیص داده می‌شوند",
          set(g) == {"FA", "IA", "FT"}, str(sorted(g)))
    check("ستون‌های بی‌شماره وارد سازه نمی‌شوند",
          all("ID" not in v and "جنسیت" not in v for v in g.values()))
    check("هر سازه همه‌ی گویه‌هایش را دارد",
          len(g["FA"]) == 5 and len(g["IA"]) == 8 and len(g["FT"]) == 3)

    # بوم مدل باید در عرض صفحه جا شود
    check("بوم مدل برای جاشدن در صفحه مقیاس می‌خورد",
          "function fitCanvas" in js and "viewBox" in js)
    check("کشیدن سازه با انگشت هم کار می‌کند", "pointerdown" in js and "pointermove" in js)
    check("مقیاسِ بوم در جابه‌جایی سازه لحاظ می‌شود", "vb.width / box.width" in js)

    # راهنما هم قدم‌به‌قدم توضیح داده باشد
    import help_docs
    fa, _ = help_docs.source_text("fa")
    en, _ = help_docs.source_text("en")
    check("راهنمای فارسی روش کار با SmartPLS را قدم‌به‌قدم دارد",
          "ساخت خودکار سازه‌ها" in fa and "گویه" in fa and "مسیر" in fa)
    check("راهنمای انگلیسی هم همین را دارد",
          "Detect constructs automatically" in en and "Indicator" in en)


# ═══════════════ ۹) خروجی Word از پنجره‌ی نتایج ═══════════════

def test_output_export():
    section("۹) خروجی Word از پنجره‌ی خروجی میزکار")
    import export_utils, design_engine, glob
    here = os.path.dirname(os.path.abspath(__file__))
    js = open(os.path.join(here, "static", "workbench.js"), encoding="utf-8").read()

    check("دکمه‌ی خروجی Word در پنجره‌ی نتایج هست", "wbOutExport" in js)
    check("نمودار مسیر به‌جای نشانی، به‌صورت تصویر در سند می‌رود",
          "![نمودار مسیر مدل]" in js and "delete res.diagram" in js)
    check("سرفصل‌های عمیق‌تر از سه سطح خام در متن نمی‌افتند",
          'depth <= 3 ? "#".repeat(depth)' in js)

    # تصویر: فقط فایل‌های خودِ آنتانو، نه مسیرهای بیرونی یا خطرناک
    png = sorted(glob.glob(os.path.join(export_utils.EXPORT_DIR, "*.png")),
                 key=os.path.getmtime)
    if png:
        good = "/download/" + os.path.basename(png[-1])
        check("تصویرِ معتبرِ خودِ آنتانو پذیرفته می‌شود",
              export_utils.resolve_image(good) is not None)
    for bad in ("../../etc/passwd", "https://evil.example/x.png",
                "/download/nope-not-here.png", "run.exe", ""):
        check(f"مسیر ناامن رد می‌شود ({bad or 'خالی'})",
              export_utils.resolve_image(bad) is None)

    # ساخت واقعی سند با تصویر و جدول
    if not png:
        print("  ⏭  تصویری برای آزمون درج نبود.")
        return
    md = ("## نتیجه‌ی آزمون\n\n"
          f"![نمودار مسیر مدل](/download/{os.path.basename(png[-1])})\n\n"
          "| شاخص | مقدار |\n|---|---|\n| n | ۹۰ |\n")
    blocks = export_utils.md_to_blocks(md)
    check("مارک‌داونِ تصویر به بلوکِ تصویر تبدیل می‌شود",
          any(k == "img" for k, _ in blocks), str([k for k, _ in blocks]))

    import docx
    name = export_utils.build_docx(blocks, title="آزمون خروجی")
    d = docx.Document(os.path.join(export_utils.EXPORT_DIR, name))
    body = "\n".join(p.text for p in d.paragraphs)
    check("تصویر واقعاً داخل فایل Word درج شد", len(d.inline_shapes) == 1)
    check("زیرنویس تصویر نوشته شد", "نمودار مسیر مدل" in body)
    check("نشانی خام فایل در متن نیامد", "/download/" not in body and "![" not in body)
    check("جدول هم سالم ساخته شد", len(d.tables) == 1)

    spec = design_engine.build_spec("academic")
    n2 = design_engine.build_designed_docx(blocks, spec, title="آزمون طراحی‌شده")
    d2 = docx.Document(os.path.join(export_utils.EXPORT_DIR, n2))
    check("در حالت «طراحی هوشمند» هم تصویر درج می‌شود", len(d2.inline_shapes) == 1)

    # بقیه‌ی قالب‌ها نباید با بلوکِ تازه کرش کنند
    for fn in ("build_txt", "build_md"):
        try:
            getattr(export_utils, fn)(blocks, "آزمون")
            ok = True
        except Exception as e:
            ok = False; print("   ", fn, e)
        check(f"{fn} با بلوکِ تصویر کرش نمی‌کند", ok)
    try:
        export_utils.build_xlsx(blocks)
        ok = True
    except Exception as e:
        ok = False; print("    build_xlsx", e)
    check("build_xlsx با بلوکِ تصویر کرش نمی‌کند", ok)


def main_run():
    test_config_integrity()
    test_grid_roundtrip()
    fname = test_pages_and_api()
    test_run_analyses(fname)
    test_admin_accordion()
    test_help_multilang()
    test_composer_never_pushed_offscreen()
    test_workbench_mobile_and_autobuild()
    test_output_export()

    print("\n" + "═" * 68)
    print(f"نتیجه: {len(PASS)} موفق، {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق‌ها:", FAIL)
        sys.exit(1)


if __name__ == "__main__":
    main_run()
