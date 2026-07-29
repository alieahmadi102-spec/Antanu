# -*- coding: utf-8 -*-
"""
main.py — هسته اصلی هوش مصنوعی آنتانو (ANTANU)
اجرا:  uvicorn main:app --host 0.0.0.0 --port 8000
"""
import os
import re
import json
import asyncio
import secrets

import httpx
from fastapi import FastAPI, Request, Form, HTTPException, UploadFile, File
from starlette.concurrency import run_in_threadpool
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse, FileResponse, Response, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader, select_autoescape

from db import get_db, init_db, hash_pw, verify_pw, generate_code, DB_PATH

# ---------------- بارگذاری فایل .env ----------------
# کلید API و تنظیمات محرمانه در فایل .env نگهداری می‌شوند (کنار main.py).
# این فایل هرگز نباید روی گیت‌هاب آپلود شود — در .gitignore مستثنی شده است.

def _load_env_file(path: str = ".env"):
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


_load_env_file()

# ---------------- تنظیمات ----------------
# آنتانو با هر سرویس سازگار با OpenAI کار می‌کند: Groq، GitHub Models، Gemini، OpenRouter و...
# فقط کافی است ANTANU_PROVIDER و ANTANU_API_KEY را تنظیم کنید (راهنما در README).

PROVIDERS = {
    # سریع و رایگان — پیشنهادی برای تعداد کاربر زیاد
    "groq": {
        "base": "https://api.groq.com/openai/v1",
        "model": "llama-3.3-70b-versatile",
    },
    # رایگان با توکن گیت‌هاب (محدودیت روزانه کم دارد)
    "github": {
        "base": "https://models.github.ai/inference",
        "model": "openai/gpt-4o-mini",
    },
    # بهترین کیفیت فارسی در بین گزینه‌های رایگان
    "gemini": {
        "base": "https://generativelanguage.googleapis.com/v1beta/openai",
        "model": "gemini-2.0-flash",
    },
    # انویدیا NIM — رایگان با ثبت‌نام در build.nvidia.com
    "nvidia": {
        "base": "https://integrate.api.nvidia.com/v1",
        "model": "meta/llama-3.3-70b-instruct",
    },
    # سربراس — رایگان و بسیار سریع
    "cerebras": {
        "base": "https://api.cerebras.ai/v1",
        "model": "llama-3.3-70b",
    },
    # میسترال — پلن رایگان دارد
    "mistral": {
        "base": "https://api.mistral.ai/v1",
        "model": "mistral-small-latest",
    },
    # توگدر — اعتبار رایگان اولیه
    "together": {
        "base": "https://api.together.xyz/v1",
        "model": "meta-llama/Llama-3.3-70B-Instruct-Turbo-Free",
    },
    # دیپ‌سیک رسمی (پولی ولی بسیار ارزان)
    "deepseek": {
        "base": "https://api.deepseek.com/v1",
        "model": "deepseek-chat",
    },
    # API مخصوص آنتانو — برای آینده که مدل اختصاصی خودت را وصل کنی
    "antanu": {
        "base": os.environ.get("ANTANU_OWN_BASE", "https://api.antanu.ai/v1"),
        "model": os.environ.get("ANTANU_OWN_MODEL", "antanu-1"),
    },
    # کلادفلر Workers AI — رایگان با حساب Cloudflare (نیاز به Account ID دارد)
    "cloudflare": {
        "base": "https://api.cloudflare.com/client/v4/accounts/ACCOUNT_ID/ai/v1",
        "model": "@cf/meta/llama-3.3-70b-instruct-fp8-fast",
    },
    # چت جی‌پی‌تی رسمی (پولی)
    "openai": {
        "base": "https://api.openai.com/v1",
        "model": "gpt-4o-mini",
    },
    # چندین مدل رایگان — openrouter/free خودش مدل رایگانِ در دسترس را انتخاب می‌کند
    "openrouter": {
        "base": "https://openrouter.ai/api/v1",
        "model": "openrouter/free, meta-llama/llama-3.3-70b:free, openai/gpt-oss-120b:free",
    },
}

PROVIDER = os.environ.get("ANTANU_PROVIDER", "groq").lower()
_p = PROVIDERS.get(PROVIDER, PROVIDERS["groq"])
API_BASE = os.environ.get("ANTANU_API_BASE", _p["base"]).rstrip("/")
API_KEY = os.environ.get("ANTANU_API_KEY", "")
# می‌توان چند مدل را با کاما جدا کرد؛ اگر اولی در دسترس نبود خودکار سراغ بعدی می‌رود
MODELS = [m.strip() for m in os.environ.get("ANTANU_MODEL", _p["model"]).split(",") if m.strip()]
ADMIN_CONTACT = os.environ.get("ANTANU_ADMIN_CONTACT", "آیدی تلگرام ادمین: @Anuyouka")

# ---------------- فهرست هوش مصنوعی‌های قابل انتخاب توسط کاربر ----------------
# دو روش تعریف:
# روش ساده: فقط ANTANU_PROVIDER و ANTANU_API_KEY را بدهید → فهرست پیش‌فرض همان سرویس ساخته می‌شود.
# روش چندسرویسه: در فایل .env هر هوش مصنوعی را در یک خط تعریف کنید (تا ۱۰ عدد):
#   ANTANU_AI_1=نام نمایشی | سرویس | نام مدل | کلید
#   «سرویس» یکی از این‌هاست: groq / gemini / github / openrouter / یا آدرس کامل API
# مثال:
#   ANTANU_AI_1=جمنای گوگل | gemini | gemini-2.0-flash | AIza...
#   ANTANU_AI_2=لاما (گراک) | groq | llama-3.3-70b-versatile | gsk_...
#   ANTANU_AI_3=دیپ‌سیک | openrouter | deepseek/deepseek-r1-distill:free | sk-or-...

def _parse_custom_ais():
    out = []
    for i in range(1, 11):
        raw = os.environ.get(f"ANTANU_AI_{i}", "").strip()
        if not raw:
            continue
        parts = [p.strip() for p in raw.split("|")]
        if len(parts) != 4:
            print(f"[ANTANU] قالب ANTANU_AI_{i} نادرست است — باید ۴ بخش جداشده با | باشد")
            continue
        name, prov, model, key = parts
        base = PROVIDERS[prov.lower()]["base"] if prov.lower() in PROVIDERS else prov.rstrip("/")
        out.append({"id": f"ai{i}", "name": name, "base": base, "model": model, "key": key})
    return out


_customs = _parse_custom_ais()

AI_CATALOG = []
if API_KEY:
    # «آنتانو (خودکار)» — هوش مصنوعی اصلی خود سایت
    AI_CATALOG.append(
        {"id": "auto", "name": "آنتانو (خودکار)", "base": API_BASE, "model": MODELS[0], "key": API_KEY}
    )
    if PROVIDER == "openrouter" and not _customs:
        # فهرست پیش‌فرض مدل‌های رایگان OpenRouter (با همان یک کلید)
        AI_CATALOG += [
            {"id": "llama",    "name": "Llama 3.3 70B",          "base": API_BASE, "model": "meta-llama/llama-3.3-70b:free",    "key": API_KEY},
            {"id": "gptoss",   "name": "GPT-OSS 120B (OpenAI)",  "base": API_BASE, "model": "openai/gpt-oss-120b:free",         "key": API_KEY},
            {"id": "deepseek", "name": "DeepSeek R1",            "base": API_BASE, "model": "deepseek/deepseek-r1-distill:free","key": API_KEY},
            {"id": "gptnano",  "name": "GPT-5.4 Nano (سریع)",    "base": API_BASE, "model": "openai/gpt-5.4-nano:free",         "key": API_KEY},
        ]
    elif not _customs and len(MODELS) > 1:
        AI_CATALOG += [
            {"id": f"m{i}", "name": m.split("/")[-1], "base": API_BASE, "model": m, "key": API_KEY}
            for i, m in enumerate(MODELS[1:], 1)
        ]

# هوش مصنوعی‌های تعریف‌شده توسط مدیر — دقیقاً همین‌ها به کاربر نمایش داده می‌شوند
AI_CATALOG += _customs

if not AI_CATALOG:
    AI_CATALOG = [{"id": "auto", "name": "آنتانو", "base": API_BASE, "model": MODELS[0], "key": ""}]

# کاتالوگ ساخته‌شده از فایل .env (به‌عنوان پشتیبان)
ENV_CATALOG = AI_CATALOG


def resolve_base(service: str) -> str:
    """نام سرویس (groq/gemini/...) یا آدرس کامل → آدرس پایه API"""
    s = (service or "").strip()
    if s.lower() in PROVIDERS:
        return PROVIDERS[s.lower()]["base"]
    return s.rstrip("/")


def get_ai_catalog():
    """فهرست نهایی هوش مصنوعی‌ها — اولویت با تنظیماتی است که ادمین در پنل ذخیره کرده"""
    from db import get_db as _gdb
    try:
        db = _gdb()
        row = db.execute("SELECT value FROM settings WHERE key = 'ai_config'").fetchone()
        db.close()
    except Exception:
        row = None
    if row:
        try:
            cfg = json.loads(row["value"])
        except json.JSONDecodeError:
            cfg = None
        if cfg:
            catalog = []
            main_key = (cfg.get("api_key") or "").strip()
            if main_key:
                prov = (cfg.get("provider") or "groq").lower()
                model = (cfg.get("model") or "").strip() or PROVIDERS.get(prov, PROVIDERS["groq"])["model"].split(",")[0].strip()
                catalog.append({
                    "id": "auto", "name": "آنتانو (خودکار)",
                    "base": resolve_base(prov), "model": model, "key": main_key,
                    "role": (cfg.get("main_role") or "general").strip().lower(),
                })
            for i, ai in enumerate(cfg.get("ais") or [], 1):
                if not (ai.get("key") or "").strip() or not (ai.get("model") or "").strip():
                    continue
                catalog.append({
                    "id": f"ai{i}",
                    "name": (ai.get("name") or f"مدل {i}").strip(),
                    "base": resolve_base(ai.get("service")),
                    "model": ai["model"].strip(),
                    "key": ai["key"].strip(),
                    "role": (ai.get("role") or "").strip().lower(),
                })
            if catalog:
                # اگر کلید اصلی خالی بود، «آنتانو (خودکار)» با اولین هوش مصنوعیِ دارای کلید کار کند
                if catalog[0]["id"] != "auto":
                    first = catalog[0]
                    catalog.insert(0, {"id": "auto", "name": "آنتانو (خودکار)",
                                       "base": first["base"], "model": first["model"], "key": first["key"]})
                return catalog
    # پشتیبان .env — اگر آن هم بی‌کلید بود ولی مدل دیگری کلید داشت، از آن استفاده کن
    cat = [dict(c) for c in ENV_CATALOG]
    if cat and not cat[0].get("key"):
        withkey = next((c for c in cat if c.get("key")), None)
        if withkey:
            cat[0].update({"base": withkey["base"], "model": withkey["model"], "key": withkey["key"]})
    return cat


# ---------------- نوار تبلیغاتی متحرک (تیکر) ----------------

# نشانی‌های داخل متن تیکر خودکار به لینک تبدیل می‌شوند
_TICKER_URL_RE = re.compile(
    r"((?:https?://|www\.)[^\s<>\"']+)", re.IGNORECASE
)
# نویسه‌های خط عربی/فارسی برای تشخیص جهت حرکت متن
_RTL_CHARS_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
_LTR_CHARS_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]")


def _ticker_linkify(text: str) -> str:
    """متن تیکر را امن می‌کند و نشانی‌ها را به لینک قابل‌کلیک تبدیل می‌کند.
    اول کل متن escape می‌شود، بعد فقط نشانی‌ها به <a> تبدیل می‌شوند — پس
    هیچ HTMLی از متن مدیر اجرا نمی‌شود."""
    safe = _html.escape(text or "")

    def to_link(m):
        url = m.group(1)
        href = url if url.lower().startswith("http") else "https://" + url
        return (f'<a href="{href}" target="_blank" rel="noopener noreferrer">{url}</a>')

    return _TICKER_URL_RE.sub(to_link, safe)


def ticker_direction(text: str) -> str:
    """جهت حرکت متن: فارسی/عربی → به راست، بقیه → به چپ.
    نشانی‌ها از شمارش کنار گذاشته می‌شوند؛ وگرنه حروف لاتینِ یک لینک،
    یک تبلیغِ فارسی را اشتباهاً «انگلیسی» نشان می‌داد."""
    body = _TICKER_URL_RE.sub(" ", text or "")
    rtl = len(_RTL_CHARS_RE.findall(body))
    ltr = len(_LTR_CHARS_RE.findall(body))
    return "rtl" if rtl > ltr else "ltr"


def _ticker_cache_key(text: str, lang: str) -> str:
    """کلید انبار ترجمه. چون به محتوای متن گره خورده، وقتی مدیر متن تبلیغ را عوض
    کند ترجمه‌های قدیمی خودبه‌خود بی‌اثر می‌شوند."""
    import hashlib
    h = hashlib.sha1((text or "").encode("utf-8")).hexdigest()[:12]
    return f"ticker_tr:{h}:{lang}"


def ticker_source_lang(text: str) -> str:
    """زبانی که مدیر متن تبلیغ را با آن نوشته است."""
    code, sure = i18n.detect_message_lang(text, fallback="fa")
    return code if sure else "fa"


def _ticker_payload(text: str, speed: int, needs_tr: bool = False) -> dict:
    return {
        "html": _ticker_linkify(text),
        "dir": ticker_direction(text),
        "speed": speed,
        "needs_tr": needs_tr,
    }


def build_ticker(lang: str | None = None):
    """اطلاعات نوار تبلیغاتی برای قالب، به زبان کاربر.

    مدیر متن را (معمولاً به فارسی) در پنل می‌نویسد؛ اینجا اگر زبان کاربر فرق داشته
    باشد، ترجمه‌ی ذخیره‌شده نشان داده می‌شود. اگر هنوز ترجمه‌ای نداریم، متن اصلی
    برمی‌گردد و با نشانه‌ی needs_tr، مرورگر ترجمه را از /api/ticker می‌گیرد — پس
    بالا آمدن صفحه هیچ‌وقت منتظر مدل نمی‌ماند.
    """
    text = (get_setting("ticker_text", "") or "").strip()
    if not text:
        return None
    try:
        speed = max(5, min(int(get_setting("ticker_speed", "25") or 25), 180))
    except (TypeError, ValueError):
        speed = 25

    target = i18n.normalize(lang) if lang else None
    if not target or target == ticker_source_lang(text):
        return _ticker_payload(text, speed)

    cached = get_setting(_ticker_cache_key(text, target), "")
    if cached:
        return _ticker_payload(cached, speed)
    return _ticker_payload(text, speed, needs_tr=True)


def pick_model_for(task: str, catalog=None):
    """مدلِ مناسبِ یک کار مشخص را برمی‌گرداند (اولین گزینه‌ی دارای کلید).

    اگر مدیر برای آن کار مدلی تعیین نکرده باشد، به مدل اصلی برمی‌گردد —
    پس هر قابلیت همیشه کار می‌کند، حتی بدون تنظیم نقش‌ها.
    """
    cat = catalog if catalog is not None else get_ai_catalog()
    if not cat:
        return None
    try:
        import router as _router
        ordered = _router.order_catalog(cat, task)
    except Exception:
        ordered = list(cat)
    return next((c for c in ordered if c.get("key")), cat[0])


def _library_context(query: str, limit: int = 4, budget: int = 4000) -> str:
    """بخش‌های مرتبط از «کتابخانه‌ی دائمی» را برای افزودن به پرامپت آماده می‌کند.

    اگر کتابخانه خالی باشد یا چیزی پیدا نشود، رشته‌ی خالی برمی‌گردد تا پرامپت
    بی‌دلیل بزرگ نشود.
    """
    try:
        import knowledge_base as _kb
        ctx = _kb.context_for(query, limit=limit, budget=budget)
    except Exception:
        return ""
    if not ctx:
        return ""
    return (
        "\n\n📚 از کتابخانه‌ی آنتانو (کتاب‌ها و جزوه‌های معتبری که مالک سایت افزوده است).\n"
        "اگر به پرسش مربوط است، از همین‌ها استفاده کن و نامِ منبع را ذکر کن؛ "
        "اگر ربطی ندارد، نادیده بگیر و از دانش خودت پاسخ بده. چیزی از خودت به این منابع نبند.\n"
        + ctx
    )


def add_knowledge(question: str, answer: str):
    """ذخیره خودکار پرسش‌وپاسخ در پایگاه دانش آنتانو (بدون تکرار)"""
    q = (question or "").strip()
    a = (answer or "").strip()
    if len(q) < 4 or len(a) < 20 or a.startswith("⚠️") or a.startswith("🧠"):
        return
    try:
        db = get_db()
        exists = db.execute("SELECT 1 FROM knowledge WHERE question = ?", (q[:500],)).fetchone()
        if not exists:
            db.execute("INSERT INTO knowledge (question, answer) VALUES (?, ?)", (q[:500], a[:8000]))
            db.commit()
        db.close()
    except Exception:
        pass


def search_knowledge(query: str, limit: int = 3):
    """جستجوی ساده در پایگاه دانش برای یافتن پاسخ‌های مشابه قبلی"""
    q = (query or "").strip()
    if len(q) < 4:
        return []
    try:
        db = get_db()
        rows = db.execute(
            "SELECT question, answer FROM knowledge WHERE question LIKE ? OR ? LIKE '%' || question || '%' "
            "ORDER BY length(question) DESC LIMIT ?",
            (f"%{q[:100]}%", q[:200], limit),
        ).fetchall()
        db.close()
        return [dict(r) for r in rows]
    except Exception:
        return []


def get_setting(key: str, default: str = "") -> str:
    try:
        db = get_db()
        row = db.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
        db.close()
        return row["value"] if row else default
    except Exception:
        return default


def set_setting(key: str, value: str):
    db = get_db()
    db.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value),
    )
    db.commit()
    db.close()

# ---------------- سطوح اشتراک ----------------
# سقف روزانه: chat بر حسب «توکن» (تقریب دلار)، بقیه بر حسب تعداد
# ۱ دلار توکن ≈ ۱ میلیون توکن ورودی/خروجی برای مدل‌های ارزان — اینجا سخاوتمندانه حساب می‌کنیم
DOLLAR_TOKENS = 200_000  # هر «دلار» = ۲۰۰ هزار توکن مصرفی (ورودی+خروجی)
QUOTAS = {
    # ستاره ۰ = مهمان رایگان (بدون ثبت‌نام): فقط گفتگو، متن و مقاله — بدون ساخت عکس/ویدیو
    0: {"tokens": DOLLAR_TOKENS // 2, "image": 0, "video": 0, "article": 1, "tts": 0},
    1: {"tokens": 1 * DOLLAR_TOKENS, "image": 0,  "video": 0, "article": 2, "tts": 0},
    2: {"tokens": 2 * DOLLAR_TOKENS, "image": 0,  "video": 0, "article": 3, "tts": 3},
    3: {"tokens": 3 * DOLLAR_TOKENS, "image": 10, "video": 1, "article": 4, "tts": 5},
    4: {"tokens": 4 * DOLLAR_TOKENS, "image": 15, "video": 3, "article": 6, "tts": 10},
}
# ستاره ۵ = نامحدود، فقط برای مدیران تیم (فروخته و نمایش داده نمی‌شود)
ADMIN_STARS = 5

QUOTA_NAMES = {"chat": "پیام", "image": "ساخت عکس", "video": "ساخت ویدیو",
               "article": "مقاله بلند", "tts": "صدای حرفه‌ای"}

# مدت اعتبار اشتراک از لحظه ثبت‌نام/تمدید (روز)
SUBSCRIPTION_DAYS = 30

MAX_TOKENS = {0: 700, 1: 700, 2: 1000, 3: 1600, 4: 2400}

app = FastAPI(title="ANTANU")

# هندلر سراسری خطا: علت واقعی را در ترمینال سرور چاپ می‌کند تا اشکال‌زدایی ساده شود
import traceback as _tb
from starlette.requests import Request as _Req
from fastapi.responses import JSONResponse as _JR


@app.exception_handler(Exception)
async def _global_error_handler(request: _Req, exc: Exception):
    from fastapi import HTTPException as _HE
    if isinstance(exc, _HE):
        return _JR({"detail": str(exc.detail)}, status_code=exc.status_code)
    print("\n" + "=" * 60)
    print(f"❌ خطای سرور در مسیر: {request.method} {request.url.path}")
    _tb.print_exc()
    print("=" * 60 + "\n")
    return _JR(
        {"detail": "خطای داخلی سرور. لطفاً متن خطا را که در پنجره سیاه (ترمینال) چاپ شده به پشتیبانی بدهید."},
        status_code=500,
    )

app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/sw.js")
def service_worker():
    """سرویس‌ورکر باید از ریشه سرو شود تا کل دامنه را پوشش دهد (PWA)"""
    return FileResponse("static/sw.js", media_type="application/javascript",
                        headers={"Cache-Control": "no-cache"})


def _build_report() -> dict:
    """گزارشِ «سرور کدام ساخت را اجرا می‌کند».

    وقتی تغییری می‌دهیم ولی روی گوشی دیده نمی‌شود، معمولاً یعنی سرور هنوز
    نسخه‌ی قدیمی را بالا آورده. این گزارش هم در پنل مدیریت نشان داده می‌شود
    و هم از مسیر /version (فقط برای مدیر) قابل دیدن است.
    """
    checks = [
        ("نوار تبلیغاتی: حرکت فریم‌به‌فریم", "static/app.js", "requestAnimationFrame(frame)"),
        ("کتابخانه‌ی دائمی آنتانو", "knowledge_base.py", "seed_knowledge_if_needed"),
        ("تحلیل خودکار آماری (SPSS/EViews/SmartPLS)", "analysis_planner.py", "validate_plan"),
        ("پاسخ به زبان کاربر", "i18n.py", "detect_message_lang"),
        ("ترجمه‌ی خودکار نوار تبلیغاتی", "main.py", "_ticker_cache_key"),
        ("نوار تبلیغاتی: اندازه‌گیری پیکسلی", "static/app.js", "initTicker"),
        ("نوار تبلیغاتی: مسیر دقیق در CSS", "static/style.css", "--ticker-from"),
        ("کتابخانه‌ی marked روی سرور خودمان", "static/vendor/marked.min.js", None),
        ("کتابخانه‌ی DOMPurify روی سرور خودمان", "static/vendor/purify.min.js", None),
        ("فونت وزیرمتن روی سرور خودمان", "static/vendor/vazirmatn/vazirmatn.css", None),
    ]
    items, all_ok = [], True
    for label, path, needle in checks:
        try:
            with open(path, encoding="utf-8", errors="ignore") as fh:
                ok = True if needle is None else (needle in fh.read())
        except OSError:
            ok = False
        all_ok = all_ok and ok
        items.append({"label": label, "ok": ok})
    return {
        "app_version": get_setting("app_version", "1.0"),
        "asset_ver": _jinja.globals.get("asset_ver"),
        "checks": items,
        "all_ok": all_ok,
        "summary": "همه‌چیز به‌روز است." if all_ok else
                   "سرور هنوز نسخه‌ی قدیمی را اجرا می‌کند — فایل‌ها را جایگزین کن و دوباره build بگیر.",
    }


@app.get("/version", response_class=PlainTextResponse)
def version_info(request: Request):
    """همان گزارش، به‌صورت متن ساده. فقط برای مدیر — کاربر عادی نباید
    اطلاعات ساخت و مسیر فایل‌های سرور را ببیند."""
    require_admin(request)
    rep = _build_report()
    lines = [
        f"ANTANU — نسخه‌ی برنامه: {rep['app_version']}",
        f"نسخه‌ی فایل‌های استاتیک (asset_ver): {rep['asset_ver']}",
        "",
        "وضعیت آخرین تغییرها روی این سرور:",
    ]
    for it in rep["checks"]:
        lines.append(f"  {'✅' if it['ok'] else '❌'}  {it['label']}")
    lines += ["", ("✅ " if rep["all_ok"] else "❌ ") + rep["summary"]]
    return "\n".join(lines)


# رندر مستقیم قالب‌ها با Jinja2 (مستقل از نسخه starlette — بدون خطای ناسازگاری)
_jinja = Environment(
    loader=FileSystemLoader("templates"),
    autoescape=select_autoescape(["html"]),
)


def _compute_asset_ver() -> str:
    """نسخه‌ی خودکار فایل‌های استاتیک بر پایه‌ی محتوای آن‌ها.
    با هر تغییر در app.js یا style.css این مقدار عوض می‌شود، پس مرورگر کاربر
    بدون نیاز به پاک‌کردن کش، خودکار نسخه‌ی جدید را می‌گیرد."""
    import hashlib
    h = hashlib.md5()
    files = ["static/app.js", "static/style.css",
             "static/vendor/marked.min.js", "static/vendor/purify.min.js",
             "static/vendor/vazirmatn/vazirmatn.css"]
    for f in files:
        try:
            with open(f, "rb") as fh:
                h.update(fh.read())
        except OSError:
            pass
    return h.hexdigest()[:10]


# یک‌بار هنگام شروع محاسبه می‌شود؛ چون هر دیپلوی کانتینر را از نو می‌سازد،
# همان لحظه نسخه‌ی تازه محاسبه و به همه‌ی صفحه‌ها تزریق می‌شود.
_jinja.globals["asset_ver"] = _compute_asset_ver()


# ---------------- چندزبانه‌سازی (زبان پیش‌فرض: انگلیسی) ----------------

import i18n

LANG_COOKIE = "antanu_lang"


def resolve_lang(request: Request | None = None, user=None) -> str:
    """زبان این درخواست را تعیین می‌کند، به این ترتیب:
    ۱) زبان ذخیره‌شده در حساب کاربر (روی همه‌ی دستگاه‌ها یکسان)
    ۲) کوکی مرورگر (برای کاربر مهمان و صفحه‌ی ورود)
    ۳) زبان مرورگر (Accept-Language)
    ۴) زبان پیش‌فرض سایت = انگلیسی
    """
    if user is not None:
        try:
            u_lang = user["lang"] if "lang" in user.keys() else None
        except Exception:
            u_lang = None
        if u_lang and i18n.is_supported(u_lang):
            return u_lang
    if request is not None:
        cookie = request.cookies.get(LANG_COOKIE)
        if cookie and i18n.is_supported(cookie):
            return cookie
        header = request.headers.get("accept-language", "")
        if header:
            return i18n.pick_from_header(header)
    return i18n.DEFAULT_LANG


def render(name: str, status_code: int = 200, request: Request | None = None,
           lang: str | None = None, **context) -> HTMLResponse:
    """قالب را با زبان درست رندر می‌کند و تابع t() را در اختیار قالب می‌گذارد.
    زبان از حساب کاربر (اگر در context باشد) یا کوکی/مرورگر گرفته می‌شود."""
    code = lang or resolve_lang(request, context.get("user"))
    ctx = {
        "lang": code,
        "dir": i18n.direction(code),
        "is_rtl": i18n.direction(code) == "rtl",
        "t": lambda key, **kw: i18n.t(key, code, **kw),
        # واژه‌نامه برای پیام‌های سمت مرورگر (app.js) — با window.T(key) خوانده می‌شود
        "i18n_js": i18n.catalog(code),
        "languages": i18n.language_list(),
        "current_lang": code,
        "lang_native": i18n.native_name(code),
        **context,
    }
    html = _jinja.get_template(name).render(**ctx)
    return HTMLResponse(html, status_code=status_code)


init_db()

# بارگذاری «مغز واژگان» (ارمنی) در پایگاه داده اگر خالی باشد
try:
    import glossary as _glossary
    _seeded = _glossary.seed_glossary_if_empty()
    if _seeded:
        print(f"[ANTANU] واژه‌نامه بارگذاری شد: {_seeded} واژه")
except Exception as _e:
    print("[ANTANU] بارگذاری واژه‌نامه انجام نشد:", _e)

# بارگذاری «کتابخانه‌ی دائمی» (کتاب‌ها و جزوه‌های data/knowledge)
try:
    import knowledge_base as _kb
    _kb_seeded = _kb.seed_knowledge_if_needed()
    if _kb_seeded:
        print(f"[ANTANU] کتابخانه بارگذاری شد: {_kb_seeded} قطعه متن")
except Exception as _e:
    print("[ANTANU] بارگذاری کتابخانه انجام نشد:", _e)

import datetime as _dt


# ---------- پشتیبان‌گیری خودکار شبانه از پایگاه داده ----------
# روی سرورهای رایگان (که ممکن است ری‌استارت شوند) هر شب یک نسخه‌ی سالم از
# antanu.db کنار گذاشته می‌شود و فقط چند نسخه‌ی آخر نگه داشته می‌شود.
def _backup_db_now() -> str | None:
    """یک نسخه‌ی سازگار (hot backup) از پایگاه داده می‌سازد و مسیرش را برمی‌گرداند."""
    import sqlite3
    if not os.path.exists(DB_PATH):
        return None
    backup_dir = os.environ.get(
        "ANTANU_BACKUP_DIR",
        os.path.join(os.path.dirname(os.path.abspath(DB_PATH)) or ".", "backups"),
    )
    os.makedirs(backup_dir, exist_ok=True)
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    out_path = os.path.join(backup_dir, f"antanu-backup-{stamp}.db")
    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(out_path)
    try:
        with dst:
            src.backup(dst)
    finally:
        src.close()
        dst.close()
    # فقط N نسخه‌ی آخر را نگه دار
    keep = int(os.environ.get("ANTANU_BACKUP_KEEP", "7"))
    files = sorted(
        f for f in os.listdir(backup_dir)
        if f.startswith("antanu-backup-") and f.endswith(".db")
    )
    for old in files[:-keep] if keep > 0 else []:
        try:
            os.remove(os.path.join(backup_dir, old))
        except Exception:
            pass
    return out_path


def _send_backup_to_telegram(path: str | None) -> tuple[bool, str]:
    """ارسال فایل پشتیبان به کانال/چت تلگرام (اگر توکن ربات و آیدی چت تنظیم شده باشد).
    برای فعال‌سازی: ANTANU_TG_BOT_TOKEN و ANTANU_TG_BACKUP_CHAT را در .env بگذار."""
    # اول از تنظیمات پنل مدیریت (پایگاه داده) می‌خوانیم، بعد از متغیرهای محیطی
    token = (get_setting("tg_bot_token", "") or os.environ.get("ANTANU_TG_BOT_TOKEN", "")).strip()
    chat = (get_setting("tg_backup_chat", "") or os.environ.get("ANTANU_TG_BACKUP_CHAT", "")).strip()
    if not token or not chat:
        return False, "توکن ربات یا آیدی کانال تنظیم نشده — در پنل مدیریت وارد و ذخیره کنید."
    if not path or not os.path.exists(path):
        return False, "فایل پشتیبان پیدا نشد."
    try:
        caption = "🗄 پشتیبان خودکار آنتانو — " + _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
        with open(path, "rb") as f:
            r = httpx.post(
                f"https://api.telegram.org/bot{token}/sendDocument",
                data={"chat_id": chat, "caption": caption},
                files={"document": (os.path.basename(path), f, "application/octet-stream")},
                timeout=httpx.Timeout(180, connect=15),
            )
        if r.status_code == 200 and r.json().get("ok"):
            return True, "پشتیبان با موفقیت به تلگرام ارسال شد."
        return False, f"تلگرام خطا داد ({r.status_code}): {r.text[:200]}"
    except Exception as e:
        return False, f"ارسال به تلگرام ناموفق بود: {e}"


async def _auto_backup_loop():
    """هر ۲۴ ساعت یک‌بار پشتیبان می‌گیرد (قابل خاموش‌کردن با ANTANU_AUTO_BACKUP=0)
    و در صورت تنظیم بودن تلگرام، یک نسخه هم به کانال تلگرام می‌فرستد."""
    interval = int(os.environ.get("ANTANU_BACKUP_INTERVAL_HOURS", "24")) * 3600
    while True:
        await asyncio.sleep(interval)
        try:
            path = await run_in_threadpool(_backup_db_now)
            if path:
                await run_in_threadpool(_send_backup_to_telegram, path)
        except Exception:
            pass


@app.on_event("startup")
async def _start_auto_backup():
    if os.environ.get("ANTANU_AUTO_BACKUP", "1") != "0":
        asyncio.create_task(_auto_backup_loop())

# ---------- تاریخ و ساعت زنده (شمسی + میلادی، به وقت ایران) ----------

_J_MONTHS = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
             "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
_WEEKDAYS = ["دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه", "شنبه", "یکشنبه"]


def _to_jalali(gy, gm, gd):
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    gy2, gm2, gd2 = gy - 1600, gm - 1, gd - 1
    g_day_no = 365 * gy2 + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400
    g_day_no += g_d_m[gm2]
    if gm2 > 1 and ((gy % 4 == 0 and gy % 100 != 0) or gy % 400 == 0):
        g_day_no += 1
    g_day_no += gd2
    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    if j_day_no < 186:
        jm, jd = 1 + j_day_no // 31, 1 + j_day_no % 31
    else:
        jm, jd = 7 + (j_day_no - 186) // 30, 1 + (j_day_no - 186) % 30
    return jy, jm, jd


def now_string() -> str:
    """مثل: جمعه ۱۹ تیر ۱۴۰۵ برابر با 10 July 2026، ساعت ۱۴:۳۰ به وقت ایران"""
    now = _dt.datetime.now(_dt.timezone(_dt.timedelta(hours=3, minutes=30)))
    jy, jm, jd = _to_jalali(now.year, now.month, now.day)
    weekday = _WEEKDAYS[now.weekday()]
    return (f"{weekday} {jd} {_J_MONTHS[jm - 1]} {jy} هجری شمسی، برابر با "
            f"{now.day} {now.strftime('%B')} {now.year} میلادی، ساعت {now.strftime('%H:%M')} به وقت ایران")


# ---------- فیلتر حروف غیرفارسی (هندی، چینی، کره‌ای، ژاپنی، تایلندی و...) ----------

_FOREIGN_RE = re.compile(
    r"[\u00C0-\u024F"            # لاتین لهجه‌دار (فرانسوی، ویتنامی، ترکی و...)
    r"\u0370-\u03FF"             # یونانی
    r"\u0400-\u04FF"             # روسی و سیریلیک
    r"\u0900-\u097F\u0980-\u0DFF\u0E00-\u0E7F"  # هندی، بنگالی، تایلندی
    r"\u1100-\u11FF\u1E00-\u1EFF"                  # کره‌ای قدیم، ویتنامی
    r"\u3040-\u30FF\u3130-\u318F\u3400-\u4DBF\u4E00-\u9FFF\uAC00-\uD7AF]+"  # ژاپنی، چینی، کره‌ای
)


# الگوی پاسخ‌های نامعتبر: بعضی مدل‌های رایگان فقط برچسب ایمنی می‌دهند (مثل «User Safety: safe»)
JUNK_RE = re.compile(
    r"^(user\s*safety|response\s*safety|safety\s*[:=]|\s*(un)?safe\s*$|i cannot|i can't assist)",
    re.IGNORECASE,
)


# زبان جاری این درخواست. فیلتر «حروف غیرفارسی» فقط وقتی معنا دارد که خروجی باید فارسی باشد؛
# اگر کاربر زبان دیگری انتخاب کرده (آلمانی، چینی، ژاپنی…) فیلتر باید کامل خاموش شود،
# وگرنه حتی «ä ö ü é ñ» آلمانی/فرانسوی/اسپانیایی هم پاک می‌شود.
import contextvars

_lang_ctx = contextvars.ContextVar("antanu_lang", default="fa")


def current_lang_code() -> str:
    try:
        return _lang_ctx.get()
    except Exception:
        return "fa"


def clean_foreign(text: str) -> str:
    if not text:
        return text
    if current_lang_code() != "fa":
        return text          # زبان کاربر فارسی نیست → دست به خروجی نزن
    return _FOREIGN_RE.sub("", text)


# بلوک ترجمه [[TR: ... ]] عمداً حاوی حروف غیرفارسی است (چینی، ژاپنی، روسی، پین‌یین لهجه‌دار و…)
# پس هرگز نباید از فیلتر بالا رد شود؛ وگرنه «你好 / nǐ hǎo» به «/ N ho» تبدیل می‌شود.
_TR_OPEN, _TR_CLOSE = "[[TR:", "]]"
_TR_BLOCK_RE = re.compile(r"(\[\[TR:[\s\S]*?\]\])")


def clean_foreign_keep_tr(text: str) -> str:
    """مثل clean_foreign، ولی محتوای بلوک‌های [[TR: ...]] را دست‌نخورده نگه می‌دارد."""
    if not text:
        return text
    return "".join(
        part if part.startswith(_TR_OPEN) else clean_foreign(part)
        for part in _TR_BLOCK_RE.split(text)
    )


def _partial_marker_tail(s: str, marker: str) -> int:
    """طول دنباله‌ای از s که می‌تواند آغازِ نیمه‌کاره‌ی marker باشد (برای استریم)."""
    for k in range(min(len(marker) - 1, len(s)), 0, -1):
        if s.endswith(marker[:k]):
            return k
    return 0


class ForeignFilter:
    """نسخه‌ی حالت‌دارِ clean_foreign برای استریم.
    چون پاسخ تکه‌تکه می‌رسد و یک بلوک [[TR: ...]] ممکن است بین چند تکه بشکند،
    وضعیت «داخل بلوک بودن» نگه داشته می‌شود و محتوای بلوک بدون فیلتر رد می‌شود.

    enabled=False یعنی هیچ فیلتری اعمال نشود — برای وقتی زبان کاربر فارسی نیست
    (مثلاً چینی یا روسی) و خروجی عمداً با خط دیگری نوشته می‌شود."""

    def __init__(self, enabled: bool = True):
        self._buf = ""
        self._inside = False
        self._enabled = enabled

    def feed(self, chunk: str) -> str:
        if not self._enabled:
            return chunk or ""
        self._buf += chunk or ""
        out = []
        while self._buf:
            if not self._inside:
                i = self._buf.find(_TR_OPEN)
                if i == -1:
                    keep = _partial_marker_tail(self._buf, _TR_OPEN)
                    cut = len(self._buf) - keep
                    out.append(clean_foreign(self._buf[:cut]))
                    self._buf = self._buf[cut:]
                    break
                out.append(clean_foreign(self._buf[:i]))
                out.append(_TR_OPEN)
                self._buf = self._buf[i + len(_TR_OPEN):]
                self._inside = True
            else:
                j = self._buf.find(_TR_CLOSE)
                if j == -1:
                    keep = _partial_marker_tail(self._buf, _TR_CLOSE)
                    cut = len(self._buf) - keep
                    out.append(self._buf[:cut])        # داخل بلوک: بدون فیلتر
                    self._buf = self._buf[cut:]
                    break
                out.append(self._buf[:j + len(_TR_CLOSE)])
                self._buf = self._buf[j + len(_TR_CLOSE):]
                self._inside = False
        return "".join(out)

    def flush(self) -> str:
        """ته‌مانده‌ی نگه‌داشته‌شده را در پایان استریم بیرون می‌دهد."""
        rest, self._buf = self._buf, ""
        if not self._enabled:
            return rest
        return rest if self._inside else clean_foreign(rest)


# مثل clean_foreign ولی الفبای لاتین (انگلیسی) را نگه می‌دارد — مخصوص ارجاع‌ها و DOI
_FOREIGN_RE_NONLATIN = re.compile(
    r"[Ͱ-Ͽ"            # یونانی
    r"Ѐ-ӿ"             # سیریلیک/روسی
    r"ऀ-ॿঀ-෿฀-๿"  # هندی، بنگالی، تایلندی
    r"ᄀ-ᇿ"                              # کره‌ای قدیم
    r"぀-ヿ㄰-㆏㐀-䶿一-鿿가-힯]+"  # ژاپنی، چینی، کره‌ای
)


def clean_foreign_keep_latin(text: str) -> str:
    return _FOREIGN_RE_NONLATIN.sub("", text) if text else text


# واژه‌هایی که یعنی کاربر اطلاعات «روز» می‌خواهد → جستجوی وب خودکار روشن می‌شود
AUTO_SEARCH_WORDS = [
    "امروز", "دیروز", "فردا", "اخبار", "خبر", "قیمت", "نرخ", "دلار", "یورو",
    "سکه", "طلا", "بیت‌کوین", "بیتکوین", "آب‌وهوا", "آب و هوا", "هوای",
    "الان", "اکنون", "هم‌اکنون", "جدیدترین", "آخرین", "چه خبر", "نتیجه بازی",
    "این هفته", "این ماه", "امسال", "چه سالی", "چندم", "تاریخ امروز", "ساعت چند",
]

def is_unlimited(user) -> bool:
    """مدیران تیم: نامحدود"""
    return bool(user["is_admin"]) or user["stars"] >= ADMIN_STARS


def is_guest(user) -> bool:
    """کاربر مهمان (رایگان، بدون ثبت‌نام) — ستاره ۰"""
    try:
        return int(user["stars"]) == 0
    except Exception:
        return False


def sub_status(user):
    """وضعیت اشتراک: (فعال؟، روزهای باقی‌مانده، تاریخ پایان شمسی)
    اولویت با expires_at است (که با تمدید ادمین به‌روز می‌شود)"""
    if is_unlimited(user):
        return True, None, None
    endd = None
    exp = user["expires_at"] if "expires_at" in user.keys() else None
    if exp:
        try:
            endd = _dt.date(*map(int, exp[:10].split("-")))
        except Exception:
            endd = None
    if endd is None:
        created = (user["created_at"] or "")[:10]
        try:
            y, m, d0 = map(int, created.split("-"))
            endd = _dt.date(y, m, d0) + _dt.timedelta(days=SUBSCRIPTION_DAYS)
        except Exception:
            return True, None, None
    today = _dt.datetime.now(_dt.timezone(_dt.timedelta(hours=3, minutes=30))).date()
    left = (endd - today).days
    jy, jm, jd = _to_jalali(endd.year, endd.month, endd.day)
    return left >= 0, left, f"{jd} {_J_MONTHS[jm - 1]} {jy}"


def quota_used(db, user_id: int, kind: str) -> int:
    """مصرف امروزِ یک نوع فعالیت (به وقت ایران — ریست هر شب ساعت ۱۲).
    برای chat مجموع توکن، برای بقیه تعداد دفعات."""
    # «امروز» به وقت ایران محاسبه می‌شود (ریست نیمه‌شب تهران)
    row = db.execute(
        "SELECT COALESCE(SUM(amount), 0) AS c FROM usage_log "
        "WHERE user_id = ? AND kind = ? AND date(created_at, '+3 hours', '+30 minutes') = date('now', '+3 hours', '+30 minutes')",
        (user_id, kind),
    ).fetchone()
    return row["c"] if row else 0


def quota_check(db, user, kind: str):
    """بررسی اشتراک و سهمیه — در صورت مشکل، HTTPException می‌دهد"""
    if is_unlimited(user):
        return
    active, left, endj = sub_status(user)
    if not active:
        raise HTTPException(
            403,
            f"⏳ اشتراک شما به پایان رسیده است (تاریخ پایان: {endj}). "
            f"برای شارژ مجدد به مدیر پیام دهید — {ADMIN_CONTACT}",
        )
    qkey = "tokens" if kind == "chat" else kind
    limit = QUOTAS.get(user["stars"], QUOTAS[1]).get(qkey, 0)
    used = quota_used(db, user["id"], kind)
    if limit <= 0 and kind in ("image", "video", "tts"):
        if is_guest(user):
            raise HTTPException(
                403,
                f"در حالت رایگان و بدون ثبت‌نام فقط گفتگو، متن و مقاله در دسترس است؛ "
                f"«{QUOTA_NAMES[kind]}» نیاز به اشتراک دارد. برای تهیه اشتراک — {ADMIN_CONTACT}",
            )
        raise HTTPException(
            403,
            f"اشتراک {user['stars']} ستاره شما امکان «{QUOTA_NAMES[kind]}» ندارد. "
            f"برای ارتقا به مدیر پیام دهید — {ADMIN_CONTACT}",
        )
    if used >= limit:
        # پیام بدون افشای عدد دقیق سقف (کاربر سطح مصرف را نبیند)
        raise HTTPException(
            429,
            f"سهمیه امروز «{QUOTA_NAMES[kind]}» شما به پایان رسید. هر شب ساعت ۱۲ بامداد دوباره فعال می‌شود.",
        )


def quota_add(db, user, kind: str, amount: int = 1):
    if is_unlimited(user):
        return
    db.execute("INSERT INTO usage_log (user_id, kind, amount) VALUES (?, ?, ?)",
               (user["id"], kind, amount))


def check_subscription(user):
    """فقط اعتبار زمانی (۳۰ روزه) اشتراک را بررسی می‌کند"""
    if is_unlimited(user):
        return
    active, left, endj = sub_status(user)
    if not active:
        raise HTTPException(
            403,
            f"⏳ اشتراک شما به پایان رسیده است (تاریخ پایان: {endj}). "
            f"برای شارژ مجدد به مدیر پیام دهید — {ADMIN_CONTACT}",
        )


BASE_SYSTEM_PROMPT = (
    "تو «آنتانو» (ANTANU) هستی؛ دستیار هوشمند فارسی‌زبان برای دانشجویان و پژوهشگران. "
    "هویت و معرفی تو: آنتانو محصول شرکت بین‌المللی آنتانو است؛ شرکتی گسترده در حوزه فناوری و هوش مصنوعی "
    "که مرکز اصلی آن در کشور ارمنستان است و در ایران و چندین کشور دیگر نیز فعالیت و شعبه دارد. "
    "اگر پرسیدند «تو کی هستی؟» بگو: «من آنتانو هستم؛ یک هوش مصنوعی گسترده برای گفتگو، پژوهش، مقاله‌نویسی، "
    "تحلیل داده و تولید محتوا که توسط شرکت آنتانو ساخته شده‌ام — شرکتی با مرکزیت ارمنستان و حضور فعال در ایران و کشورهای دیگر.» "
    "اگر درباره معنا یا ریشه نام «آنتانو» پرسیدند، توضیح بده که آنتانو نام برند این شرکت است، ریشه در فرهنگ ارمنی دارد "
    "و نماد دانایی و همراهی هوشمند است. "
    "هرگز خود را ساخته گوگل، اوپن‌ای‌آی، متا، آنتروپیک یا هر شرکت دیگری معرفی نکن و نام مدل زیربنایی را فاش نکن. ""\n\nتوانمندی‌های تخصصی پژوهشی تو (وقتی کاربر خواست، کامل و دانشگاهی انجام بده): "
    "موضوع‌یابی و ساخت عنوان پایان‌نامه؛ نوشتن پروپوزال؛ ساخت فهرست و سرفصل‌های پایان‌نامه؛ نوشتن پیشینه پژوهش و مبانی نظری؛ "
    "ارجاع‌دهی و فوت‌نوت به سبک APA و IEEE؛ نگارش فصل چهارم و تفسیر نتایج آماری؛ نقد علمی فصل‌ها در نقش داور. "
    "برای تحلیل آماری واقعی (رگرسیون، SEM، آزمون‌ها، پایایی، میانجی‌گری)، به کاربر بگو فایل داده (Excel/CSV/SPSS) را با «📊 تحلیل آماری» آپلود کند تا محاسبه واقعی انجام شود. "
    "\n\nبسیار مهم — تشخیص نیاز کاربر: پیش از پاسخ، دقیق تشخیص بده کاربر «چه می‌خواهد». "
    "اگر کاربر سؤال ساده پرسید، سلام/گفتگوی معمولی کرد، یا یک درخواست کوتاه داشت، فقط و فقط همان را کوتاه، مستقیم و دقیق پاسخ بده؛ "
    "دقیقاً به همان چیزی که پرسیده جواب بده و از توضیح اضافی، مقدمه‌چینی، نتیجه‌گیری و بسط دادن خودداری کن. "
    "هرگز یک پرسش ساده را مثل مقاله یا متن پژوهشی و طولانی پاسخ نده. "
    "فقط زمانی محتوای بلند، ساختارمند و دانشگاهی (با عنوان‌بندی و بخش‌بندی) تولید کن که کاربر صراحتاً «مقاله»، «پایان‌نامه»، «متن بلند»، «توضیح کامل/مفصل» یا یکی از ابزارهای پژوهشی را خواسته باشد. "
    "طول و لحن پاسخ را با طول و لحن درخواست کاربر متناسب کن: سؤال کوتاه → پاسخ کوتاه. "
    "قواعد نگارش که همیشه باید رعایت کنی: "
    "۱) به فارسیِ معیار، روان و طبیعی بنویس؛ از ترجمه تحت‌اللفظی و جمله‌بندی انگلیسی‌مآب جداً پرهیز کن. "
    "۲) دستور زبان، املا و نشانه‌گذاری فارسی را کامل رعایت کن: نیم‌فاصله در «می‌شود» و «کتاب‌ها»، فعل در انتهای جمله، حروف اضافه درست. "
    "۳) اصطلاحات تخصصی را به فارسی بنویس و در اولین اشاره، معادل انگلیسی را داخل پرانتز بیاور. "
    "۴) در حوزه‌های دانشگاهی — به‌ویژه حسابداری، مدیریت، آمار و روش تحقیق — دقیق و علمی پاسخ بده. "
    "۵) اگر چیزی را نمی‌دانی صادقانه بگو نمی‌دانم و حدس نزن. "
    "۶) پاسخ را ساختارمند ارائه کن و اگر از نتایج جستجوی وب استفاده کردی، منبع را ذکر کن. "
    "۷) از تکرار واژه‌ها، عبارت‌ها و مطالب پرهیز کن؛ همیشه واژگان متنوع و مطالب تازه به کار ببر. "
    "۸) بسیار مهم: خروجی فقط با حروف فارسی (و در صورت نیاز، معادل انگلیسی داخل پرانتز) باشد؛ "
    "هرگز واژه‌های روسی، هندی، چینی، ویتنامی، فرانسوی، اسپانیایی یا هر زبان دیگری را وسط متن فارسی نیاور. "
    "اگر واژه‌ای را نمی‌دانی، ساده‌ترین معادل فارسی را بنویس. جمله ناتمام یا شکسته ننویس. "
    "استثنای مهم این قاعده: وقتی کاربر «ترجمه» خواسته است، متنِ ترجمه‌شده باید با خطِ اصلیِ همان زبان مقصد "
    "نوشته شود (چینی با حروف چینی 你好، ژاپنی با ژاپنی، روسی با سیریلیک، ارمنی با ارمنی و…) "
    "و این متن را داخل بلوک [[TR: ...]] بگذار. این تنها جایی است که نوشتن خطِ غیرفارسی لازم و درست است. "
    "۹) نگارش تمیز و مرتب: متن را با ساختار روشن بنویس؛ برای عنوان‌ها از # و ## و ### استفاده کن، "
    "برای فهرست‌ها از «- »، و جدول‌ها را با قالب استاندارد مارک‌داون (خط سرستون و خط جداکننده |---|) بساز. "
    "از خطوط خالیِ اضافه، نویسه‌های درهم و کاراکترهای زائد پرهیز کن تا خروجی Word و PDF مرتب باشد. "
    "۱۰) پانوشت: هر جا لازم شد برای یک اصطلاح تخصصی یا منبع، پانوشت بگذار؛ در متن بعد از واژه بنویس [^۱] "
    "و در انتها تعریف را در خطی جدا بیاور: «[^۱]: معادل انگلیسی یا توضیح». برای اصطلاحات، معادل انگلیسی را در پانوشت بده. "
    "۱۱) قالب ترجمه — بسیار مهم: هر وقت معنی یا ترجمه‌ی یک «واژه» یا یک «متن» را می‌دهی (به هر زبانی، از جمله ارمنی)، "
    "به‌جای نوشتنِ ترجمه داخل جمله، آن را حتماً در این قالب ویژه بده تا کاربر بتواند فقط ترجمه را کپی کند:\n"
    "[[TR: واژه یا عبارت مبدأ || ترجمه || تلفظ فنگلیش]]\n"
    "قواعد این قالب: بخش اول = خود واژه/عبارت اصلی؛ بخش دوم = ترجمه (همین در باکس کپی نمایش داده می‌شود)؛ "
    "بخش سوم = تلفظ. جداکننده‌ی بخش‌ها دقیقاً دو خط عمودی «||» است. "
    "بسیار مهم درباره‌ی بخش سوم (تلفظ): تلفظ باید تلفظِ «متنِ ترجمه‌شده» (بخش دوم، یعنی زبان مقصد) با حروف لاتین/فینگلیش باشد — "
    "هرگز تلفظِ متنِ سؤالِ کاربر (زبان مبدأ) را نده. "
    "مثال درست: کاربر می‌گوید «آدرس این هتل کجاست به چینی» → "
    "[[TR: آدرس این هتل کجاست || 这家酒店的地址在哪里 || zhè jiā jiǔdiàn de dìzhǐ zài nǎlǐ]] "
    "(یعنی تلفظِ همان جمله‌ی چینی، نه تلفظِ جمله‌ی فارسی). "
    "مثال درست دیگر: «سلام به ارمنی» → [[TR: سلام || Բարև || Barev]]. "
    "اگر زبان مقصد با حروف لاتین نوشته می‌شود (مثل انگلیسی، آلمانی، فرانسوی)، تلفظِ خوانا و ساده‌شده‌ی همان واژه‌ی مقصد را بنویس. "
    "برای ترجمه‌ی «کلمه‌به‌کلمه»، برای هر واژه یک بلوک [[TR:...]] جداگانه بده و بخش سوم (تلفظ) را حتماً پر کن؛ خالی نگذار. "
    "برای «متن طولانی»، یک بلوک [[TR:...]] بده که بخش دومش کل متنِ ترجمه‌شده و بخش سومش تلفظِ همان متنِ ترجمه‌شده باشد. "
    "خودِ ترجمه فقط داخل [[TR:...]] باشد، نه داخل جمله. "
    "۱۲) درخواست ترجمه = فقط ترجمه؛ بدون گفتگو. اگر از پیام کاربر فهمیدی که «ترجمه» می‌خواهد "
    "(مثلاً «… به ارمنی»، «… به چینی»، «ترجمه کن»، «معنی … چیست»)، آن‌وقت کلِ متنِ پیام او "
    "«محتوایی است که باید ترجمه شود» — نه صحبتی که با تو شده باشد. "
    "پس حتی اگر آن متن شامل سلام و احوالپرسی باشد (مثل «سلام خوب هستی؟»)، به هیچ وجه به آن پاسخ نده: "
    "نه سلام کن، نه بنویس «بله من خوبم، ممنون»، نه هیچ جمله‌ی گفتگویی، تعارف، مقدمه یا نتیجه‌گیری. "
    "مستقیم و فقط بلوک‌های [[TR:...]] را بده. "
    "۱۳) درک درست نیت کاربر: پیش از پاسخ، دقیق بفهم کاربر چه می‌خواهد و به همان پاسخ بده. "
    "اگر درخواست کوتاه یا نصفه بود، از متن پیام و پیام‌های پیشین نیت را دریاب. "
    "هرگز پاسخ بی‌ربط، سرسری یا حدسی نده؛ اگر واقعاً مبهم بود، فقط یک پرسش کوتاه و دقیق بپرس. "
    "۱۴) بی‌مقدمه و مستقیم: اگر پیام کاربر شامل پرسش یا درخواستی است، پاسخ را با سلام، تعارف، "
    "«حتماً»، «سؤال خوبی پرسیدی» یا تکرار صورتِ پرسش شروع نکن؛ از همان جمله‌ی اول برو سر اصل مطلب. "
    "(فقط اگر پیام کاربر تنها یک سلامِ خالی و بدون هیچ درخواستی بود، یک سلامِ کوتاه پاسخ بده.)"
)


# ---------------- تشخیص «درخواست ترجمه» در پیام کاربر ----------------
# هدف: وقتی کاربر می‌نویسد «سلام خوب هستی؟ ترجمه ارمنی»، آنتانو به سلام جواب ندهد؛
# کل متن را محتوای مورد ترجمه بداند و مستقیم ترجمه بدهد.

_TR_CUE_WORDS = ("ترجمه", "معنی", "معنا", "معادل", "برگردان", "translate", "بگو به")

_LANG_ALIASES_CACHE = {}


def _lang_alias_map():
    """نام فارسی زبان → کد ISO (برای تشخیص «… به چینی»). یک‌بار ساخته و کش می‌شود."""
    if _LANG_ALIASES_CACHE:
        return _LANG_ALIASES_CACHE
    try:
        import translate_engine as te
        for code, name in te.LANGUAGES.items():
            if code == "auto":
                continue
            _LANG_ALIASES_CACHE[name] = code
            # «ترکی استانبولی» → «ترکی» هم پذیرفته شود
            first = name.split()[0]
            _LANG_ALIASES_CACHE.setdefault(first, code)
    except Exception:
        pass
    return _LANG_ALIASES_CACHE


def _detect_translate_intent(message: str):
    """اگر پیام «درخواست ترجمه» باشد، نام زبان مقصد (یا رشته‌ی خالی) را برمی‌گرداند؛ وگرنه None."""
    msg = (message or "").strip()
    if not msg or len(msg) > 4000:
        return None
    aliases = _lang_alias_map()
    if not aliases:
        return None
    # ۱) الگوی روشن «به <زبان>» / «به زبان <زبان>»
    for name in aliases:
        if re.search(r"به\s*(?:زبان\s*)?" + re.escape(name) + r"\b", msg):
            return name
    # ۲) واژه‌ی نشانه‌ی ترجمه + نام یک زبان در متن («ترجمه ارمنی»)
    if any(cue in msg for cue in _TR_CUE_WORDS):
        for name in aliases:
            if re.search(re.escape(name) + r"\b", msg):
                return name
        return ""
    return None


TRANSLATE_ONLY_NOTE = (
    "\n\n[تشخیص سیستم: این پیام یک «درخواست ترجمه» است{tgt}.] "
    "کل متن پیام کاربر «محتوایی است که باید ترجمه شود»، نه گفتگویی که با تو شده باشد. "
    "اگر در آن سلام یا احوالپرسی هست، آن را هم فقط ترجمه کن و به هیچ وجه به آن پاسخ نده. "
    "هیچ سلام، تعارف، جمله‌ی گفتگویی، مقدمه یا نتیجه‌گیری ننویس. "
    "فقط و فقط بلوک‌های [[TR: متن مبدأ || ترجمه || تلفظ لاتینِ همان ترجمه]] را بده."
)


# ---------------- ابزارهای احراز هویت ----------------

def current_user(request: Request):
    token = request.cookies.get("antanu_session")
    if not token:
        return None
    db = get_db()
    row = db.execute(
        "SELECT u.* FROM sessions s JOIN users u ON u.id = s.user_id WHERE s.token = ?",
        (token,),
    ).fetchone()
    db.close()
    return row


def require_user(request: Request):
    user = current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="ابتدا وارد حساب خود شوید")
    return user


def require_admin(request: Request):
    user = require_user(request)
    if not user["is_admin"]:
        raise HTTPException(status_code=403, detail="دسترسی فقط برای مدیر")
    return user


def make_session(user_id: int) -> str:
    token = secrets.token_hex(32)
    db = get_db()
    db.execute("INSERT INTO sessions (token, user_id) VALUES (?, ?)", (token, user_id))
    db.commit()
    db.close()
    return token


# ---------------- محافظت امنیتی (ضدحمله) ----------------
# ردیاب درون‌حافظه‌ای برای جلوگیری از حمله‌ی حدس‌زدن رمز و ارسال سیل‌آسا
import time as _time
from collections import deque, defaultdict

_LOGIN_FAILS = {}          # کلید: نام‌کاربری|IP → [تعداد, زمان اولین خطا]
LOGIN_MAX_FAILS = 6        # پس از ۶ تلاش ناموفق
LOGIN_LOCK_SECONDS = 300   # قفل ۵ دقیقه‌ای

_CHAT_HITS = defaultdict(deque)  # کلید: user_id → صف زمان درخواست‌ها
CHAT_MAX_PER_MIN = 25            # حداکثر ۲۵ درخواست چت در دقیقه برای هر کاربر


def _client_ip(request: Request) -> str:
    fwd = request.headers.get("x-forwarded-for", "")
    if fwd:
        return fwd.split(",")[0].strip()
    return request.client.host if request.client else "?"


def login_locked(key: str):
    """آیا این کلید (کاربر+IP) قفل است؟ → (قفل؟، ثانیه‌های باقی‌مانده)"""
    rec = _LOGIN_FAILS.get(key)
    if not rec:
        return False, 0
    count, first = rec
    if count < LOGIN_MAX_FAILS:
        return False, 0
    elapsed = _time.time() - first
    if elapsed >= LOGIN_LOCK_SECONDS:
        _LOGIN_FAILS.pop(key, None)  # پنجره تمام شد
        return False, 0
    return True, int(LOGIN_LOCK_SECONDS - elapsed)


def login_fail(key: str):
    rec = _LOGIN_FAILS.get(key)
    if not rec or (_time.time() - rec[1]) >= LOGIN_LOCK_SECONDS:
        _LOGIN_FAILS[key] = [1, _time.time()]
    else:
        rec[0] += 1


def login_ok(key: str):
    _LOGIN_FAILS.pop(key, None)


def chat_rate_ok(user_id: int) -> bool:
    now = _time.time()
    dq = _CHAT_HITS[user_id]
    while dq and now - dq[0] > 60:
        dq.popleft()
    if len(dq) >= CHAT_MAX_PER_MIN:
        return False
    dq.append(now)
    return True


# ---------------- صفحات ----------------

@app.get("/", response_class=HTMLResponse)
def root(request: Request):
    return RedirectResponse("/chat" if current_user(request) else "/login")


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if current_user(request):
        return RedirectResponse("/chat")
    return render("login.html", request=request, error=None, version=get_setting("app_version", "1.0"), contact=ADMIN_CONTACT)


@app.post("/login", response_class=HTMLResponse)
def login(request: Request, username: str = Form(...), password: str = Form(...), fingerprint: str = Form("")):
    lock_key = f"{username.strip().lower()}|{_client_ip(request)}"
    locked, remain = login_locked(lock_key)
    if locked:
        mins = max(1, remain // 60)
        return render("login.html", request=request,
                      error=f"به‌دلیل تلاش‌های ناموفق زیاد، ورود موقتاً قفل شد. حدود {mins} دقیقه دیگر دوباره تلاش کنید.",
                      status_code=429, version=get_setting("app_version", "1.0"), contact=ADMIN_CONTACT)

    db = get_db()
    # ورود با نام کاربری یا ایمیل (مهمان‌های رایگان با ایمیل ثبت‌نام می‌کنند)
    ident = username.strip()
    user = db.execute(
        "SELECT * FROM users WHERE username = ? OR email = ?", (ident, ident.lower())
    ).fetchone()

    def fail(msg):
        db.close()
        return render("login.html", request=request, error=msg, status_code=400, version=get_setting("app_version", "1.0"), contact=ADMIN_CONTACT)

    if not user or not verify_pw(password, user["password"]):
        login_fail(lock_key)
        _, left = login_locked(lock_key)
        rec = _LOGIN_FAILS.get(lock_key)
        tries_left = max(0, LOGIN_MAX_FAILS - (rec[0] if rec else 0))
        extra = f" ({tries_left} تلاش دیگر تا قفل موقت)" if tries_left <= 3 and tries_left > 0 else ""
        return fail("نام کاربری یا گذرواژه نادرست است." + extra)

    login_ok(lock_key)

    # قفل یک‌دستگاهی: هر حساب فقط روی همان دستگاهی که اولین‌بار وارد شده کار می‌کند (ادمین معاف است)
    if not user["is_admin"]:
        if user["device_fp"] and fingerprint and user["device_fp"] != fingerprint:
            return fail("این حساب به دستگاه دیگری متصل است. برای انتقال به دستگاه جدید با ادمین تماس بگیرید.")
        if not user["device_fp"] and fingerprint:
            db.execute("UPDATE users SET device_fp = ? WHERE id = ?", (fingerprint, user["id"]))
            db.commit()

    db.close()
    token = make_session(user["id"])
    resp = RedirectResponse("/chat", status_code=303)
    resp.set_cookie("antanu_session", token, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 30)
    return resp


@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request):
    if current_user(request):
        return RedirectResponse("/chat")
    return render("register.html", request=request, error=None, version=get_setting("app_version", "1.0"))


@app.post("/register", response_class=HTMLResponse)
def register(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    code: str = Form(...),
    fingerprint: str = Form(""),
):
    username, code = username.strip(), code.strip()
    db = get_db()

    def fail(msg):
        db.close()
        return render("register.html", request=request, error=msg, status_code=400, version=get_setting("app_version", "1.0"))

    if len(username) < 3:
        return fail("نام کاربری باید حداقل ۳ حرف باشد.")
    if len(password) < 6:
        return fail("گذرواژه باید حداقل ۶ حرف باشد.")
    if not fingerprint:
        return fail("شناسه دستگاه دریافت نشد. صفحه را نوسازی کنید.")
    if db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        return fail("این نام کاربری قبلاً گرفته شده است.")

    code_row = db.execute("SELECT * FROM codes WHERE code = ?", (code,)).fetchone()
    if not code_row:
        return fail("کد ثبت‌نام نامعتبر است.")
    if code_row["used"]:
        return fail("این کد قبلاً استفاده شده است. هر کد فقط برای یک کاربر و یک دستگاه معتبر است.")

    cur = db.execute(
        "INSERT INTO users (username, password, stars, device_fp, code_used, expires_at) "
        "VALUES (?, ?, ?, ?, ?, datetime('now', '+30 days'))",
        (username, hash_pw(password), code_row["stars"], fingerprint, code),
    )
    user_id = cur.lastrowid
    db.execute("UPDATE codes SET used = 1, used_by = ? WHERE id = ?", (username, code_row["id"]))
    db.commit()
    db.close()

    token = make_session(user_id)
    resp = RedirectResponse("/chat", status_code=303)
    resp.set_cookie("antanu_session", token, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 30)
    return resp


EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")


@app.get("/guest", response_class=HTMLResponse)
def guest_page(request: Request):
    """صفحه‌ی ثبت‌نام رایگان مهمان (با ایمیل)."""
    if current_user(request):
        return RedirectResponse("/chat")
    return render("guest.html", request=request, error=None, version=get_setting("app_version", "1.0"))


@app.post("/guest", response_class=HTMLResponse)
def guest_register(
    request: Request,
    email: str = Form(...),
    username: str = Form(...),
    password: str = Form(...),
):
    """ثبت‌نام رایگان مهمان با ایمیل — بدون کد. هر ایمیل فقط یک‌بار.
    فقط گفتگو، متن و مقاله در دسترس است؛ ساخت عکس و ویدیو نیاز به اشتراک دارد."""
    email = (email or "").strip().lower()
    username = (username or "").strip()
    password = password or ""

    def fail(msg):
        return render("guest.html", request=request, error=msg, status_code=400,
                      version=get_setting("app_version", "1.0"))

    if not EMAIL_RE.match(email):
        return fail("ایمیل معتبر وارد کنید (مثلاً name@example.com).")
    if len(username) < 3:
        return fail("نام کاربری باید حداقل ۳ نویسه باشد.")
    if len(password) < 6:
        return fail("گذرواژه باید حداقل ۶ نویسه باشد.")

    db = get_db()
    if db.execute("SELECT 1 FROM users WHERE email = ?", (email,)).fetchone():
        db.close()
        return fail("این ایمیل قبلاً استفاده شده است. با همین ایمیل وارد شوید یا ایمیل دیگری بزنید.")
    if db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        db.close()
        return fail("این نام کاربری قبلاً گرفته شده است؛ نام دیگری انتخاب کنید.")

    cur = db.execute(
        "INSERT INTO users (username, password, email, stars, code_used, expires_at) "
        "VALUES (?, ?, ?, 0, 'guest', datetime('now', '+3650 days'))",
        (username, hash_pw(password), email),
    )
    user_id = cur.lastrowid
    db.commit()
    db.close()

    token = make_session(user_id)
    resp = RedirectResponse("/chat", status_code=303)
    resp.set_cookie("antanu_session", token, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 30)
    return resp


@app.post("/logout")
def logout(request: Request):
    token = request.cookies.get("antanu_session")
    if token:
        db = get_db()
        db.execute("DELETE FROM sessions WHERE token = ?", (token,))
        db.commit()
        db.close()
    resp = RedirectResponse("/login", status_code=303)
    resp.delete_cookie("antanu_session")
    return resp


@app.get("/buy", response_class=HTMLResponse)
def buy_page(request: Request):
    return render("buy.html", request=request, contact=ADMIN_CONTACT)


# ---------------- صندوق ایده‌ها و نظرات کاربران ----------------

FEEDBACK_CATS = {"idea", "bug", "feature", "other"}


@app.post("/api/feedback")
async def api_feedback(request: Request):
    """ثبت ایده/نظر کاربر برای پیشرفت آنتانو (همه‌ی کاربران — از جمله مهمان‌ها)."""
    user = require_user(request)
    body = await request.json()
    content = (body.get("content") or "").strip()
    category = (body.get("category") or "idea").strip()
    if category not in FEEDBACK_CATS:
        category = "idea"
    if len(content) < 3:
        raise HTTPException(400, "متن نظر خیلی کوتاه است.")
    if len(content) > 4000:
        content = content[:4000]
    # جلوگیری از ارسال سیل‌آسا: حداکثر ۵ نظر در ساعت برای هر کاربر
    db = get_db()
    recent = db.execute(
        "SELECT COUNT(*) AS c FROM feedback WHERE user_id = ? "
        "AND created_at > datetime('now', '-1 hour')",
        (user["id"],),
    ).fetchone()["c"]
    if recent >= 5:
        db.close()
        raise HTTPException(429, "نظرهای زیادی ثبت کرده‌اید؛ کمی بعد دوباره تلاش کنید.")
    db.execute(
        "INSERT INTO feedback (user_id, username, category, content) VALUES (?, ?, ?, ?)",
        (user["id"], user["username"], category, content),
    )
    db.commit()
    db.close()
    return {"ok": True, "message": "🙏 نظر شما ثبت شد و به دست تیم آنتانو می‌رسد. سپاسگزاریم!"}


# ---------------- ترجمه‌ی چند‌لحنه ----------------

@app.get("/api/translate/langs")
def translate_langs(request: Request):
    """فهرست زبان‌ها و لحن‌ها برای رابط کاربری ترجمه."""
    require_user(request)
    import translate_engine as te
    return {"languages": te.LANGUAGES, "tones": te.TONES}


@app.post("/api/translate")
async def api_translate(request: Request):
    """ترجمه‌ی متن به یک زبان با چند لحن مختلف تا کاربر بهترین را انتخاب و کپی کند."""
    user = require_user(request)
    check_subscription(user)
    if not chat_rate_ok(user["id"]):
        raise HTTPException(429, "درخواست‌ها را خیلی سریع می‌فرستید؛ کمی صبر کنید.")
    body = await request.json()
    text = (body.get("text") or "").strip()
    target = (body.get("target") or "en").strip()
    source = (body.get("source") or "auto").strip()
    tones = body.get("tones") or ["formal", "polite", "friendly", "casual"]

    import translate_engine as te
    if not text:
        raise HTTPException(400, "متنی برای ترجمه وارد نشده است.")
    if len(text) > te.MAX_TEXT:
        raise HTTPException(400, f"متن طولانی است (حداکثر {te.MAX_TEXT} نویسه). آن را به بخش‌های کوچک‌تر تقسیم کنید.")
    if target not in te.LANGUAGES or target == "auto":
        raise HTTPException(400, "زبان مقصد نامعتبر است.")
    tones = [t for t in tones if t in te.TONES][:6] or ["formal", "polite", "friendly", "casual"]

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    # اگر ارمنی یا انگلیسی در کار باشد (مبدأ/مقصد/متن) قاعده‌ی واژه‌نامه اعمال می‌شود
    arm_hint = ""
    try:
        import glossary as _gl
        force = {l for l in (target, source) if l in _gl.GLOSSARY_LANGS}
        arm_hint = _gl.prompt_hint_for_text(text, force_langs=force)
    except Exception:
        pass

    prompt = te.build_translate_prompt(text, target, source, tones, arm_hint)
    c = pick_model_for("translate")
    if not c or not c.get("key"):
        raise HTTPException(503, "سرویس ترجمه فعلاً در دسترس نیست؛ کمی بعد تلاش کنید.")

    # خروجی نباید از فیلتر clean_foreign رد شود؛ متن مقصد عمداً غیرفارسی است
    out = await _call_model_once(c, prompt, max_tokens=2600, keep_foreign=True)

    translations = []
    try:
        m = re.search(r"\{.*\}", out or "", re.S)
        data = json.loads(m.group(0)) if m else {}
        for item in data.get("translations", []):
            if isinstance(item, dict) and item.get("text"):
                translations.append({"tone": item.get("tone", ""), "text": item["text"].strip()})
    except Exception:
        translations = []
    if not translations:
        # اگر JSON نبود، کل خروجی را به‌عنوان یک ترجمه بده
        translations = [{"tone": "ترجمه", "text": (out or "").strip()}]

    dbq = get_db()
    quota_add(dbq, user, "chat", max(500, len(text) * 2))
    dbq.commit(); dbq.close()

    return {"target": te.lang_name(target), "translations": translations}


# ---------------- متن‌به‌گفتار حرفه‌ای (TTS) ----------------

TTS_VOICES = {"self": "صدای آنتانو", "female": "صدای خانم", "male": "صدای آقا"}


@app.post("/api/tts")
async def api_tts(request: Request):
    """تبدیل متن به فایل صوتی حرفه‌ای (ElevenLabs). کلید و شناسه‌ی صداها از پنل مدیریت."""
    user = require_user(request)
    check_subscription(user)
    body = await request.json()
    text = (body.get("text") or "").strip()
    voice = (body.get("voice") or "self").strip()
    if not text:
        raise HTTPException(400, "متنی برای تبدیل به صدا نیست.")
    if len(text) > 1000:
        text = text[:1000]

    token = (get_setting("elevenlabs_key", "") or os.environ.get("ELEVENLABS_API_KEY", "")).strip()
    if not token:
        raise HTTPException(400, "صدای حرفه‌ای فعال نیست؛ فعلاً از دکمه «🔊 خواندن» استفاده کنید.")

    default_vid = (get_setting("tts_voice_self", "") or os.environ.get("ELEVENLABS_VOICE_ID", "21m00Tcm4TlvDq8ikWAM")).strip()
    vmap = {
        "self": get_setting("tts_voice_self", "").strip() or default_vid,
        "female": get_setting("tts_voice_female", "").strip() or default_vid,
        "male": get_setting("tts_voice_male", "").strip() or default_vid,
    }
    voice_id = vmap.get(voice) or default_vid

    db = get_db()
    quota_check(db, user, "tts")
    db.close()

    try:
        import export_utils
        async with httpx.AsyncClient(timeout=httpx.Timeout(120, connect=15)) as client:
            r = await client.post(
                f"https://api.elevenlabs.io/v1/text-to-speech/{voice_id}",
                headers={"xi-api-key": token, "Content-Type": "application/json"},
                json={"text": text, "model_id": "eleven_multilingual_v2",
                      "voice_settings": {"stability": 0.5, "similarity_boost": 0.75}},
            )
        if r.status_code != 200:
            if r.status_code in (401, 403):
                raise HTTPException(400, "کلید صدای حرفه‌ای نامعتبر است؛ مدیر سیستم آن را بررسی کند.")
            if r.status_code == 429:
                raise HTTPException(429, "سهمیه‌ی ماهانه‌ی سرویس صدای حرفه‌ای پر شده است؛ بعداً تلاش کنید.")
            raise HTTPException(400, "ساخت صدا ناموفق بود؛ کمی بعد دوباره تلاش کنید.")
        name = f"antanu-tts-{secrets.token_hex(6)}.mp3"
        with open(os.path.join(export_utils.EXPORT_DIR, name), "wb") as f:
            f.write(r.content)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "ارتباط با سرویس صدا برقرار نشد؛ کمی بعد تلاش کنید.")

    dbq = get_db()
    quota_add(dbq, user, "tts")
    dbq.commit(); dbq.close()
    return {"url": f"/download/{name}", "voice": TTS_VOICES.get(voice, voice)}


# ---------------- تبدیل گفتار به متن (STT) ----------------

@app.post("/api/stt")
async def api_stt(request: Request, file: UploadFile = File(...), lang: str = Form("")):
    """تبدیل فایل صوتی به متن با Whisper (سازگار با OpenAI/Groq). کلید از پنل مدیریت."""
    user = require_user(request)
    check_subscription(user)
    key = (get_setting("stt_key", "") or os.environ.get("ANTANU_STT_KEY", "")).strip()
    base = (get_setting("stt_base", "") or os.environ.get("ANTANU_STT_BASE", "https://api.groq.com/openai/v1")).rstrip("/")
    model = (get_setting("stt_model", "") or os.environ.get("ANTANU_STT_MODEL", "whisper-large-v3")).strip()
    if not key:
        raise HTTPException(400, "تبدیل صدا به متن فعلاً فعال نیست؛ مدیر باید کلید را در پنل مدیریت بگذارد.")

    raw = await file.read()
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل صوتی زیاد است (حداکثر ۲۵ مگابایت).")

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    try:
        files = {"file": (file.filename or "audio.webm", raw, file.content_type or "audio/webm")}
        data = {"model": model, "response_format": "json"}
        if lang and lang not in ("", "auto"):
            data["language"] = lang
        async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=15)) as client:
            r = await client.post(f"{base}/audio/transcriptions",
                                  headers={"Authorization": f"Bearer {key}"}, data=data, files=files)
        if r.status_code != 200:
            if r.status_code in (401, 403):
                raise HTTPException(400, "کلید تبدیل صدا نامعتبر است؛ مدیر سیستم بررسی کند.")
            raise HTTPException(400, "تبدیل صدا به متن ناموفق بود؛ کمی بعد دوباره تلاش کنید.")
        text = (r.json().get("text") or "").strip()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "ارتباط با سرویس تبدیل صدا برقرار نشد؛ کمی بعد تلاش کنید.")

    dbq = get_db()
    quota_add(dbq, user, "chat", 600)
    dbq.commit(); dbq.close()
    return {"text": text}


# ---------------- کمک‌تابع‌های صدا (برای زنجیره‌ی ویس‌به‌ویس) ----------------

async def _stt_transcribe(raw: bytes, filename: str, content_type: str, source: str = "") -> str:
    key = (get_setting("stt_key", "") or os.environ.get("ANTANU_STT_KEY", "")).strip()
    base = (get_setting("stt_base", "") or os.environ.get("ANTANU_STT_BASE", "https://api.groq.com/openai/v1")).rstrip("/")
    model = (get_setting("stt_model", "") or os.environ.get("ANTANU_STT_MODEL", "whisper-large-v3")).strip()
    if not key:
        raise HTTPException(400, "تبدیل صدا به متن فعال نیست؛ مدیر باید کلید Whisper را در پنل بگذارد.")
    files = {"file": (filename or "a.webm", raw, content_type or "audio/webm")}
    data = {"model": model, "response_format": "json"}
    if source and source not in ("", "auto"):
        data["language"] = source
    async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=15)) as client:
        r = await client.post(f"{base}/audio/transcriptions",
                              headers={"Authorization": f"Bearer {key}"}, data=data, files=files)
    if r.status_code != 200:
        raise HTTPException(400, "تبدیل صدا به متن ناموفق بود؛ کمی بعد تلاش کنید.")
    return (r.json().get("text") or "").strip()


async def _tts_generate(text: str, voice: str = "self") -> str | None:
    key = (get_setting("elevenlabs_key", "") or os.environ.get("ELEVENLABS_API_KEY", "")).strip()
    if not key or not text:
        return None
    default_vid = (get_setting("tts_voice_self", "") or os.environ.get("ELEVENLABS_VOICE_ID", "21m00Tcm4TlvDq8ikWAM")).strip()
    vmap = {
        "self": get_setting("tts_voice_self", "").strip() or default_vid,
        "female": get_setting("tts_voice_female", "").strip() or default_vid,
        "male": get_setting("tts_voice_male", "").strip() or default_vid,
    }
    vid = vmap.get(voice) or default_vid
    try:
        import export_utils
        async with httpx.AsyncClient(timeout=httpx.Timeout(120, connect=15)) as client:
            r = await client.post(
                f"https://api.elevenlabs.io/v1/text-to-speech/{vid}",
                headers={"xi-api-key": key, "Content-Type": "application/json"},
                json={"text": text[:1000], "model_id": "eleven_multilingual_v2",
                      "voice_settings": {"stability": 0.5, "similarity_boost": 0.75}},
            )
        if r.status_code == 200:
            name = f"antanu-tts-{secrets.token_hex(6)}.mp3"
            with open(os.path.join(export_utils.EXPORT_DIR, name), "wb") as f:
                f.write(r.content)
            return f"/download/{name}"
    except Exception:
        return None
    return None


@app.post("/api/voice_translate")
async def api_voice_translate(request: Request, file: UploadFile = File(...),
                              target: str = Form("fa"), source: str = Form("auto"),
                              voice: str = Form("self")):
    """ویس‌به‌ویس نوبتی: صدا → متن → ترجمه → صدا. همه در یک درخواست."""
    user = require_user(request)
    check_subscription(user)
    import translate_engine as te
    raw = await file.read()
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل صوتی زیاد است (حداکثر ۲۵ مگابایت).")

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    transcript = await _stt_transcribe(raw, file.filename or "voice.webm", file.content_type or "audio/webm", source)
    if not transcript:
        return {"transcript": "", "translation": "", "audio_url": None}

    translation = transcript
    if target and target != "auto" and target in te.LANGUAGES:
        c = pick_model_for("translate")
        if c and c.get("key"):
            arm_hint = ""
            try:
                import glossary as _gl
                force = {l for l in (target, source) if l in _gl.GLOSSARY_LANGS}
                arm_hint = _gl.prompt_hint_for_text(transcript, force_langs=force)
            except Exception:
                pass
            prompt = te.build_translate_prompt(transcript, target, source, ["friendly"], arm_hint)
            out = await _call_model_once(c, prompt, max_tokens=1500, keep_foreign=True)
            try:
                m = re.search(r"\{.*\}", out or "", re.S)
                d = json.loads(m.group(0)) if m else {}
                trs = d.get("translations") or []
                translation = ((trs[0].get("text") if trs else "") or "").strip() or transcript
            except Exception:
                translation = (out or "").strip() or transcript
            dbq = get_db()
            quota_add(dbq, user, "chat", 1500)
            dbq.commit(); dbq.close()

    audio_url = await _tts_generate(translation, voice)
    return {"transcript": transcript, "translation": translation, "audio_url": audio_url}


# ---------------- ساخت آهنگ / موسیقی (Replicate) ----------------

@app.post("/api/song")
async def api_song(request: Request):
    """ساخت قطعه‌ی موسیقی از توضیح متنی با Replicate (MusicGen). کلید از پنل مدیریت."""
    user = require_user(request)
    check_subscription(user)
    body = await request.json()
    prompt = (body.get("prompt") or "").strip()
    try:
        duration = max(3, min(int(body.get("duration") or 8), 30))
    except (TypeError, ValueError):
        duration = 8
    if not prompt:
        raise HTTPException(400, "توضیح آهنگ را بنویس (مثلاً: یک ملودی آرام با پیانو و ویولن).")

    key = (get_setting("replicate_key", "") or os.environ.get("REPLICATE_API_TOKEN", "")).strip()
    if not key:
        raise HTTPException(400, "ساخت آهنگ فعال نیست؛ مدیر باید کلید Replicate را در پنل مدیریت بگذارد.")

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    try:
        import export_utils
        headers = {"Authorization": f"Token {key}", "Content-Type": "application/json", "Prefer": "wait"}
        payload = {"input": {"prompt": prompt, "duration": duration, "output_format": "mp3"}}
        async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=15)) as client:
            r = await client.post("https://api.replicate.com/v1/models/meta/musicgen/predictions",
                                  headers=headers, json=payload)
            if r.status_code not in (200, 201):
                if r.status_code in (401, 403):
                    raise HTTPException(400, "کلید Replicate نامعتبر است؛ مدیر سیستم بررسی کند.")
                raise HTTPException(400, "ساخت آهنگ ناموفق بود؛ کمی بعد تلاش کنید.")
            data = r.json()
            # اگر هنوز آماده نبود، چند بار وضعیت را می‌پرسیم
            for _ in range(30):
                if data.get("status") == "succeeded":
                    break
                if data.get("status") in ("failed", "canceled"):
                    raise HTTPException(400, "ساخت آهنگ ناموفق بود؛ توضیح دیگری امتحان کنید.")
                get_url = (data.get("urls") or {}).get("get")
                if not get_url:
                    break
                await asyncio.sleep(3)
                data = (await client.get(get_url, headers={"Authorization": f"Token {key}"})).json()
            out = data.get("output")
            if isinstance(out, list):
                out = out[0] if out else None
            if not out:
                raise HTTPException(400, "خروجی آهنگ دریافت نشد؛ کمی بعد دوباره تلاش کنید.")
            au = await client.get(out)
            name = f"antanu-song-{secrets.token_hex(6)}.mp3"
            with open(os.path.join(export_utils.EXPORT_DIR, name), "wb") as f:
                f.write(au.content)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "ارتباط با سرویس آهنگ برقرار نشد؛ کمی بعد تلاش کنید.")

    dbq = get_db()
    quota_add(dbq, user, "chat", 2000)
    dbq.commit(); dbq.close()
    return {"url": f"/download/{name}"}


# ---------------- اسکلت منشی/تماس تلفنی (Twilio) — برای اپلیکیشن آینده ----------------

def _twilio_configured() -> bool:
    return bool(get_setting("twilio_sid", "").strip() and get_setting("twilio_token", "").strip())


@app.api_route("/twilio/voice", methods=["GET", "POST"])
async def twilio_voice(request: Request):
    """Webhook تماس ورودی Twilio — پاسخ TwiML.
    فعلاً «خاموش» است تا وقتی اپلیکیشن و شماره‌ی تلفن آماده شود؛ کلیدها از پنل خوانده می‌شوند."""
    greeting = get_setting("phone_greeting", "") or "سلام، شما با دستیار هوشمند آنتانو تماس گرفته‌اید. پیام خود را بگویید."
    # TwiML ساده: خوش‌آمد + ضبط پیام. توسعه‌ی کامل (STT/ترجمه/پاسخ) در نسخه‌ی اپ.
    twiml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Response>'
        f'<Say language="fa-IR">{_html.escape(greeting)}</Say>'
        '<Record maxLength="120" playBeep="true"/>'
        '<Say language="fa-IR">پیام شما ثبت شد. خداحافظ.</Say>'
        '</Response>'
    )
    return Response(content=twiml, media_type="application/xml")


@app.get("/api/phone/status")
def phone_status(request: Request):
    """وضعیت آمادگی منشی تلفنی (برای پنل/اپ آینده)."""
    require_user(request)
    return {"configured": _twilio_configured(), "phone": get_setting("twilio_phone", "")}


# ---------------- ربات گفتگوی نوبتی (صوتی/متنی) — منشی درون‌برنامه‌ای ----------------

@app.post("/api/converse")
async def api_converse(request: Request, file: UploadFile | None = File(None),
                       text: str = Form(""), persona: str = Form(""),
                       history: str = Form("[]"), voice: str = Form("self"),
                       lang: str = Form("auto")):
    """یک نوبت از گفتگوی صوتی/متنی با منشی هوشمند.
    ورودی: صدای کاربر (یا متن) + دستور شخصیت (persona) + تاریخچه‌ی گفتگو.
    خروجی: متنِ گفته‌ی کاربر، پاسخ منشی، و صدای پاسخ (در صورت فعال بودن TTS).
    شروع/پایان و مدت گفتگو را سمت کاربر مدیریت می‌کند؛ این مسیر فقط یک نوبت را جلو می‌برد."""
    user = require_user(request)
    check_subscription(user)

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    # ۱) گفته‌ی کاربر: از صدا یا از متن
    said = (text or "").strip()
    if file is not None:
        raw = await file.read()
        if len(raw) > 25 * 1024 * 1024:
            raise HTTPException(400, "حجم فایل صوتی زیاد است (حداکثر ۲۵ مگابایت).")
        if raw:
            said = await _stt_transcribe(raw, file.filename or "turn.webm",
                                         file.content_type or "audio/webm",
                                         "" if lang == "auto" else lang)
    if not said:
        raise HTTPException(400, "چیزی نگفتید؛ دوباره تلاش کنید.")

    # ۲) ساخت پیام‌ها با شخصیت دلخواه و تاریخچه
    try:
        hist = json.loads(history) if history else []
        if not isinstance(hist, list):
            hist = []
    except Exception:
        hist = []
    persona = (persona or "").strip() or \
        "تو «منشی هوشمند آنتانو» هستی؛ مؤدب، خلاصه‌گو و کمک‌کننده. کوتاه و طبیعی پاسخ بده."
    messages = [{"role": "system", "content": BASE_SYSTEM_PROMPT + "\n\n" + persona}]
    for m in hist[-12:]:
        role = m.get("role") if isinstance(m, dict) else None
        content = (m.get("content") if isinstance(m, dict) else "") or ""
        if role in ("user", "assistant") and content:
            messages.append({"role": role, "content": str(content)[:2000]})
    messages.append({"role": "user", "content": said})

    # ۳) پاسخ مدل — منشی باید سریع جواب بدهد
    c = pick_model_for("fast")
    if not c:
        raise HTTPException(503, "سرویس گفتگو موقتاً در دسترس نیست.")
    reply = await _call_model_once(c, messages=messages, max_tokens=600)
    reply = (reply or "").strip() or "متوجه نشدم، می‌شود دوباره بگویید؟"

    dbq = get_db()
    quota_add(dbq, user, "chat", 1200)
    dbq.commit(); dbq.close()

    # ۴) صدای پاسخ (اختیاری)
    audio_url = await _tts_generate(reply, voice)
    return {"you_said": said, "reply": reply, "audio_url": audio_url}


# ---------------- صدای حیوانات ----------------

def _animal_text_to_map(text: str) -> str:
    """متنِ «نام = آدرس» (هر خط یکی) را به JSON تبدیل می‌کند.
    فقط آدرس‌های داخلی /download/... یا https:// پذیرفته می‌شوند."""
    mapping = {}
    for line in (text or "").splitlines():
        line = line.strip()
        if not line or "=" not in line:
            continue
        name, _, url = line.partition("=")
        name, url = name.strip(), url.strip()
        if not name or not url:
            continue
        if not (url.startswith("/download/") or url.startswith("https://")):
            continue
        mapping[name] = url
    return json.dumps(mapping, ensure_ascii=False)


def _animal_map_to_text(raw: str) -> str:
    """JSON ذخیره‌شده را برای نمایش در پنل به متنِ «نام = آدرس» برمی‌گرداند."""
    try:
        mapping = json.loads(raw) if raw else {}
        if not isinstance(mapping, dict):
            return ""
    except Exception:
        return ""
    return "\n".join(f"{k} = {v}" for k, v in mapping.items())


@app.get("/api/animal_sounds")
def animal_sounds(request: Request):
    """فهرست صداهای حیوانات که مدیر در پنل تعریف کرده (JSON: {"گربه":"/download/cat.mp3", ...}).
    اگر چیزی تعریف نشده باشد، فهرست خالی و پیام آماده‌سازی برمی‌گرداند."""
    require_user(request)
    raw = get_setting("animal_sounds", "").strip()
    try:
        mapping = json.loads(raw) if raw else {}
        if not isinstance(mapping, dict):
            mapping = {}
    except Exception:
        mapping = {}
    return {"ready": bool(mapping), "sounds": mapping,
            "note": "" if mapping else "هنوز صدایی تعریف نشده؛ مدیر می‌تواند از پنل اضافه کند."}


# ---------------- دوبله/ترجمه‌ی ویدیو (اسکلت — نیازمند ffmpeg + کلیدها) ----------------

def _ffmpeg_available() -> bool:
    import shutil
    return bool(shutil.which("ffmpeg"))


@app.get("/api/video_dub/status")
def video_dub_status(request: Request):
    """آیا زیرساخت دوبله‌ی ویدیو آماده است؟ (ffmpeg + STT + TTS)"""
    require_user(request)
    stt_ready = bool((get_setting("stt_key", "") or os.environ.get("ANTANU_STT_KEY", "")).strip())
    tts_ready = bool((get_setting("elevenlabs_key", "") or os.environ.get("ELEVENLABS_API_KEY", "")).strip())
    ff = _ffmpeg_available()
    return {"ready": ff and stt_ready and tts_ready,
            "ffmpeg": ff, "stt": stt_ready, "tts": tts_ready,
            "note": "برای فعال‌شدن دوبله، ffmpeg روی سرور و کلیدهای Whisper/ElevenLabs لازم است."}


@app.post("/api/video_dub")
async def api_video_dub(request: Request, file: UploadFile = File(...),
                        target: str = Form("fa"), source: str = Form("auto"),
                        voice: str = Form("self")):
    """دوبله‌ی پایه‌ی ویدیو: استخراج صدا → متن → ترجمه → صدای جدید → چسباندن روی ویدیو.
    اگر ffmpeg/کلیدها آماده نباشند، پیام روشن می‌دهد (بدون خطای ۵۰۰)."""
    user = require_user(request)
    check_subscription(user)
    if not _ffmpeg_available():
        raise HTTPException(400, "دوبله‌ی ویدیو هنوز روی سرور فعال نیست (نیازمند ffmpeg).")
    if not (get_setting("stt_key", "") or os.environ.get("ANTANU_STT_KEY", "")).strip():
        raise HTTPException(400, "برای دوبله، مدیر باید کلید تبدیل صدا به متن را در پنل بگذارد.")

    import translate_engine as te, export_utils, subprocess, tempfile
    raw = await file.read()
    if len(raw) > 200 * 1024 * 1024:
        raise HTTPException(400, "حجم ویدیو زیاد است (حداکثر ۲۰۰ مگابایت).")

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    workdir = tempfile.mkdtemp()
    try:
        src_path = os.path.join(workdir, "in" + (os.path.splitext(file.filename or "")[1] or ".mp4"))
        with open(src_path, "wb") as f:
            f.write(raw)
        # ۱) استخراج صدا
        wav_path = os.path.join(workdir, "audio.wav")
        subprocess.run(["ffmpeg", "-y", "-i", src_path, "-vn", "-ac", "1", "-ar", "16000", wav_path],
                       check=True, capture_output=True, timeout=300)
        with open(wav_path, "rb") as f:
            audio_bytes = f.read()
        transcript = await _stt_transcribe(audio_bytes, "audio.wav", "audio/wav",
                                           "" if source == "auto" else source)
        if not transcript:
            raise HTTPException(400, "گفتاری در ویدیو پیدا نشد.")
        # ۲) ترجمه
        translation = transcript
        if target and target != "auto" and target in te.LANGUAGES:
            c = pick_model_for("translate")
            if c:
                prompt = te.build_translate_prompt(transcript, target, source, ["friendly"], "")
                out = await _call_model_once(c, prompt, max_tokens=2000, keep_foreign=True)
                try:
                    m = re.search(r"\{.*\}", out or "", re.S)
                    d = json.loads(m.group(0)) if m else {}
                    trs = d.get("translations") or []
                    translation = ((trs[0].get("text") if trs else "") or "").strip() or transcript
                except Exception:
                    translation = (out or "").strip() or transcript
        # ۳) صدای جدید
        dub_url = await _tts_generate(translation, voice)
        if not dub_url:
            raise HTTPException(400, "ساخت صدای دوبله فعال نیست؛ مدیر باید کلید ElevenLabs را بگذارد.")
        dub_path = os.path.join(export_utils.EXPORT_DIR, os.path.basename(dub_url))
        # ۴) چسباندن صدای جدید روی ویدیو
        out_name = f"antanu-dub-{secrets.token_hex(6)}.mp4"
        out_path = os.path.join(export_utils.EXPORT_DIR, out_name)
        subprocess.run(["ffmpeg", "-y", "-i", src_path, "-i", dub_path,
                        "-map", "0:v:0", "-map", "1:a:0", "-c:v", "copy",
                        "-shortest", out_path],
                       check=True, capture_output=True, timeout=300)
        dbq = get_db()
        quota_add(dbq, user, "chat", 2500)
        dbq.commit(); dbq.close()
        return {"transcript": transcript, "translation": translation,
                "url": f"/download/{out_name}"}
    except HTTPException:
        raise
    except subprocess.TimeoutExpired:
        raise HTTPException(400, "پردازش ویدیو طول کشید؛ ویدیوی کوتاه‌تری امتحان کنید.")
    except Exception:
        raise HTTPException(400, "دوبله‌ی ویدیو ناموفق بود؛ کمی بعد تلاش کنید.")
    finally:
        try:
            import shutil
            shutil.rmtree(workdir, ignore_errors=True)
        except Exception:
            pass


# ---------------- پاورقی‌گذاری روی فایل کاربر ----------------

@app.post("/api/footnote")
async def api_footnote(request: Request, file: UploadFile = File(...)):
    """به متن فایلِ کاربر پاورقی علمی اضافه می‌کند و فایل Word با پانوشت واقعی می‌سازد."""
    user = require_user(request)
    check_subscription(user)
    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های لازم نصب نیستند: {_e}")

    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل زیاد است (حداکثر ۱۰ مگابایت).")
    ext = os.path.splitext(file.filename or "f")[1].lower() or ".txt"
    src = os.path.join(export_utils.EXPORT_DIR, f"fn-src-{secrets.token_hex(6)}{ext}")
    with open(src, "wb") as f:
        f.write(raw)
    try:
        text = export_utils.extract_markdown(src)
    except Exception:
        text = ""
    finally:
        try:
            os.remove(src)
        except OSError:
            pass
    if not text or len(text.strip()) < 20:
        raise HTTPException(400, "متنی برای پاورقی‌گذاری در فایل پیدا نشد (شاید فایل اسکن‌شده یا خالی است).")
    if len(text) > 12000:
        text = text[:12000]

    db = get_db()
    quota_check(db, user, "chat")
    db.close()

    c = pick_model_for("article")
    if not c or not c.get("key"):
        raise HTTPException(503, "سرویس فعلاً در دسترس نیست؛ کمی بعد تلاش کنید.")

    prompt = (
        "به متن زیر پاورقی (فوت‌نوت) علمی و دقیق اضافه کن. قواعد الزامی:\n"
        "۱) متنِ اصلی را کلمه‌به‌کلمه و بدون تغییر حفظ کن؛ فقط بعد از هر اصطلاح تخصصی یا مفهوم مهم، نشانه‌ی «[^n]» بگذار.\n"
        "۲) برای اصطلاحات تخصصی، معادل انگلیسی؛ برای مفاهیم مهم، توضیح کوتاه و در صورت لزوم منبع بنویس.\n"
        "۳) شماره‌ی پاورقی‌ها را از ۱ و به‌ترتیب ظهور در متن بگذار.\n"
        "۴) در انتهای متن، تعریف هر پاورقی را در خطی جداگانه بیاور: «[^n]: توضیح».\n"
        "۵) هیچ مقدمه، توضیح یا جمله‌ای بیرون از خود متن و پاورقی‌ها اضافه نکن.\n\n"
        "متن:\n" + text
    )
    out = await _call_model_once(c, prompt, system=BASE_SYSTEM_PROMPT, max_tokens=4000)

    dbq = get_db()
    quota_add(dbq, user, "chat", 4000)
    dbq.commit(); dbq.close()

    blocks = export_utils.md_to_blocks(out or text)
    fn_count = sum(len(t) for k, t in blocks if k == "footnotes")
    title = os.path.splitext(file.filename or "سند")[0][:60] or "سند پاورقی‌دار"
    # real_footnotes=True → پاورقی واقعیِ پایین صفحه‌ی Word
    name = await run_in_threadpool(export_utils.build_docx, blocks, "Vazirmatn", 14, title,
                                   "right", False, False, True)
    return {"url": f"/download/{name}", "count": fn_count}


import html as _html


def render_markdown_html(md: str) -> str:
    """تبدیل مارک‌داون راهنما به HTML امن در سمت سرور (بدون نیاز به CDN)"""
    def inline(t):
        t = _html.escape(t)
        t = re.sub(r"\[([^\]]+)\]\((https?://[^\s)]+)\)",
                   r'<a href="\2" target="_blank" rel="noopener">\1</a>', t)
        t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
        return t

    lines = (md or "").split("\n")
    out, i, n, in_code = [], 0, len(md.split("\n")) if md else 0, False
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            out.append("<pre><code>" if not in_code else "</code></pre>")
            in_code = not in_code
            i += 1
            continue
        if in_code:
            out.append(_html.escape(line))
            i += 1
            continue
        if not s:
            i += 1
            continue
        if s in ("---", "***"):
            out.append("<hr>")
            i += 1
            continue
        if "|" in s and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|:- ") and "-" in lines[i + 1]:
            header = [c.strip() for c in s.strip().strip("|").split("|")]
            rows, j = [], i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            th = "".join(f"<th>{inline(c)}</th>" for c in header)
            trs = "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in rows)
            out.append(f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>")
            i = j
            continue
        mh = re.match(r"^(#{1,6})\s+(.*)$", s)
        if mh:
            lvl = len(mh.group(1))
            out.append(f"<h{lvl}>{inline(mh.group(2))}</h{lvl}>")
            i += 1
            continue
        if s.startswith("> "):
            out.append(f"<blockquote>{inline(s[2:])}</blockquote>")
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(inline(re.sub(r"^[-*]\s+", "", lines[i].strip())))
                i += 1
            out.append("<ul>" + "".join(f"<li>{it}</li>" for it in items) + "</ul>")
            continue
        if s.startswith("<div") or s == "</div>":
            i += 1
            continue
        out.append(f"<p>{inline(s)}</p>")
        i += 1
    if in_code:
        out.append("</code></pre>")
    return "\n".join(out)


@app.get("/help", response_class=HTMLResponse)
def help_page(request: Request):
    """راهنمای کامل استفاده — عمومی، بدون نیاز به ورود"""
    content = "# راهنما\nفعلاً در دسترس نیست."
    for path in ("راهنمای-کاربران.md", "راهنما.md"):
        if os.path.exists(path):
            try:
                with open(path, encoding="utf-8") as f:
                    content = f.read()
                break
            except Exception:
                pass
    return render("help.html", request=request, content_html=render_markdown_html(content))


@app.get("/chat", response_class=HTMLResponse)
def chat_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login")
    active, days_left, end_date = sub_status(user)
    limit_txt = "نامحدود"  # مصرف از کاربر مخفی است
    sub_warn = ""
    if not is_unlimited(user):
        if not active:
            sub_warn = (f"⛔ اشتراک شما به پایان رسیده است (تاریخ پایان: {end_date}). "
                        f"برای شارژ مجدد به مدیر پیام دهید — {ADMIN_CONTACT}")
        elif days_left is not None and days_left <= 5:
            sub_warn = (f"⏳ تنها {days_left} روز از اشتراک شما باقی مانده است (پایان: {end_date}). "
                        f"برای تمدید به مدیر پیام دهید — {ADMIN_CONTACT}")
    ticker = build_ticker(resolve_lang(request, user))
    return render("chat.html", request=request, user=user, daily_limit=limit_txt,
                  version=get_setting("app_version", "1.0"),
                  announcement=get_setting("announcement", ""),
                  ticker=ticker,
                  sub_warn=sub_warn)


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login")
    if not user["is_admin"]:
        return RedirectResponse("/chat")
    return render("admin.html", request=request, user=user)


# ---------------- API پروفایل کاربر ----------------

@app.get("/api/profile")
def api_profile(request: Request):
    user = require_user(request)

    created = (user["created_at"] or "")[:10]
    joined = created
    try:
        y, m, d0 = map(int, created.split("-"))
        jy, jm, jd = _to_jalali(y, m, d0)
        joined = f"{jd} {_J_MONTHS[jm - 1]} {jy}"
    except Exception:
        pass

    active, days_left, end_date = sub_status(user)
    avatar = user["avatar"] if "avatar" in user.keys() else None
    return {
        "username": user["username"],
        "stars": user["stars"],
        "avatar": avatar,
        "joined": joined,
        "unlimited": is_unlimited(user),
        "sub_active": active,
        "days_left": days_left,
        "end_date": end_date,
        "contact": ADMIN_CONTACT,
    }


@app.post("/api/profile/avatar")
async def api_profile_avatar(request: Request, file: UploadFile = File(...)):
    user = require_user(request)
    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(400, "حجم عکس نباید بیشتر از ۵ مگابایت باشد")
    try:
        import io
        import base64
        from PIL import Image
        img = Image.open(io.BytesIO(raw)).convert("RGB")
        img.thumbnail((128, 128))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=82)
        dataurl = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        raise HTTPException(400, "فایل عکس معتبر نیست")
    db = get_db()
    db.execute("UPDATE users SET avatar = ? WHERE id = ?", (dataurl, user["id"]))
    db.commit()
    db.close()
    return {"avatar": dataurl}


@app.post("/api/profile/password")
async def api_profile_password(request: Request):
    user = require_user(request)
    body = await request.json()
    old = body.get("old") or ""
    new = body.get("new") or ""
    if len(new) < 6:
        raise HTTPException(400, "گذرواژه جدید باید حداقل ۶ حرف باشد")
    if not verify_pw(old, user["password"]):
        raise HTTPException(400, "گذرواژه فعلی نادرست است")
    token = request.cookies.get("antanu_session") or ""
    db = get_db()
    db.execute("UPDATE users SET password = ? WHERE id = ?", (hash_pw(new), user["id"]))
    # همه نشست‌های دیگر این کاربر بسته می‌شوند (امنیت)
    db.execute("DELETE FROM sessions WHERE user_id = ? AND token <> ?", (user["id"], token))
    db.commit()
    db.close()
    return {"ok": True}


# ---------------- API گفتگوها ----------------

@app.get("/api/conversations")
def list_conversations(request: Request):
    user = require_user(request)
    db = get_db()
    rows = db.execute(
        "SELECT id, title, created_at FROM conversations WHERE user_id = ? ORDER BY id DESC",
        (user["id"],),
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.get("/api/conversations/search")
def search_conversations(request: Request, q: str = ""):
    """جستجوی تمام‌متنی در محتوای پیام‌های همه گفتگوهای کاربر (نه فقط عنوان)"""
    user = require_user(request)
    q = q.strip()
    if len(q) < 2:
        return []
    db = get_db()
    rows = db.execute(
        """SELECT c.id AS conv_id, c.title AS title, c.created_at AS created_at, m.content AS content
           FROM messages m JOIN conversations c ON c.id = m.conversation_id
           WHERE c.user_id = ? AND m.content LIKE ?
           ORDER BY m.id DESC LIMIT 200""",
        (user["id"], f"%{q}%"),
    ).fetchall()
    db.close()

    seen = set()
    results = []
    for r in rows:
        if r["conv_id"] in seen:
            continue
        seen.add(r["conv_id"])
        content = r["content"] or ""
        pos = content.lower().find(q.lower())
        if pos == -1:
            snippet = content[:90]
        else:
            start = max(0, pos - 30)
            snippet = ("…" if start > 0 else "") + content[start:pos + len(q) + 60]
        results.append({
            "id": r["conv_id"],
            "title": r["title"] or "گفتگوی بدون عنوان",
            "created_at": r["created_at"],
            "snippet": snippet.strip(),
        })
        if len(results) >= 30:
            break
    return results


@app.get("/api/conversations/{conv_id}/messages")
def conversation_messages(conv_id: int, request: Request):
    user = require_user(request)
    db = get_db()
    conv = db.execute(
        "SELECT id FROM conversations WHERE id = ? AND user_id = ?", (conv_id, user["id"])
    ).fetchone()
    if not conv:
        db.close()
        raise HTTPException(404, "گفتگو یافت نشد")
    rows = db.execute(
        "SELECT role, content FROM messages WHERE conversation_id = ? ORDER BY id", (conv_id,)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.post("/api/conversations/{conv_id}/rename")
async def rename_conversation(conv_id: int, request: Request):
    user = require_user(request)
    body = await request.json()
    title = (body.get("title") or "").strip()[:60]
    if not title:
        raise HTTPException(400, "نام خالی است")
    db = get_db()
    db.execute(
        "UPDATE conversations SET title = ? WHERE id = ? AND user_id = ?",
        (title, conv_id, user["id"]),
    )
    db.commit()
    db.close()
    return {"ok": True}


@app.delete("/api/conversations/{conv_id}")
def delete_conversation(conv_id: int, request: Request):
    user = require_user(request)
    db = get_db()
    db.execute("DELETE FROM conversations WHERE id = ? AND user_id = ?", (conv_id, user["id"]))
    db.commit()
    db.close()
    return {"ok": True}


@app.post("/api/conversations/{conv_id}/share")
def share_conversation(conv_id: int, request: Request):
    """یک لینک عمومی فقط‌خواندنی برای گفتگو می‌سازد (یا اگر قبلاً ساخته شده، همان را برمی‌گرداند)"""
    user = require_user(request)
    db = get_db()
    conv = db.execute(
        "SELECT id, share_token FROM conversations WHERE id = ? AND user_id = ?",
        (conv_id, user["id"]),
    ).fetchone()
    if not conv:
        db.close()
        raise HTTPException(404, "گفتگو یافت نشد")
    token = conv["share_token"]
    if not token:
        token = secrets.token_urlsafe(10)
        db.execute("UPDATE conversations SET share_token = ? WHERE id = ?", (token, conv_id))
        db.commit()
    db.close()
    return {"token": token, "url": f"/share/{token}"}


@app.post("/api/conversations/{conv_id}/unshare")
def unshare_conversation(conv_id: int, request: Request):
    user = require_user(request)
    db = get_db()
    db.execute("UPDATE conversations SET share_token = NULL WHERE id = ? AND user_id = ?",
               (conv_id, user["id"]))
    db.commit()
    db.close()
    return {"ok": True}


@app.get("/share/{token}", response_class=HTMLResponse)
def shared_conversation(token: str, request: Request):
    """صفحه‌ی عمومی فقط‌خواندنی یک گفتگوی به‌اشتراک‌گذاشته‌شده (بدون نیاز به ورود)"""
    if not re.fullmatch(r"[A-Za-z0-9_-]+", token or ""):
        raise HTTPException(404, "لینک نامعتبر است")
    db = get_db()
    conv = db.execute(
        "SELECT id, title FROM conversations WHERE share_token = ?", (token,)
    ).fetchone()
    if not conv:
        db.close()
        raise HTTPException(404, "این گفتگو یافت نشد یا اشتراک آن لغو شده است")
    rows = db.execute(
        "SELECT role, content FROM messages WHERE conversation_id = ? ORDER BY id", (conv["id"],)
    ).fetchall()
    db.close()
    msgs = [{"role": r["role"], "content": r["content"]} for r in rows]
    return render("shared.html", request=request,
                  title=conv["title"] or "گفتگوی آنتانو",
                  messages=json.dumps(msgs, ensure_ascii=False))


# ---------------- API حافظه بلندمدت ----------------

@app.get("/api/memories")
def list_memories(request: Request):
    user = require_admin(request)
    db = get_db()
    rows = db.execute(
        "SELECT id, content, created_at FROM memories WHERE user_id = ? ORDER BY id DESC",
        (user["id"],),
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.post("/api/memory")
async def save_memory(request: Request):
    user = require_admin(request)
    body = await request.json()
    content = (body.get("content") or "").strip()
    if not content:
        raise HTTPException(400, "متن خالی است")
    db = get_db()
    db.execute("INSERT INTO memories (user_id, content) VALUES (?, ?)", (user["id"], content[:4000]))
    db.commit()
    db.close()
    return {"ok": True}


@app.delete("/api/memories/{mem_id}")
def delete_memory(mem_id: int, request: Request):
    user = require_admin(request)
    db = get_db()
    db.execute("DELETE FROM memories WHERE id = ? AND user_id = ?", (mem_id, user["id"]))
    db.commit()
    db.close()
    return {"ok": True}


# ---------------- کتابخانه منابع پژوهشی (ارجاع‌دهی APA / IEEE) ----------------

REF_FIELDS = ("ref_type", "authors", "title", "year", "source",
              "volume", "issue", "pages", "url", "note")


def _fmt_authors(authors: str) -> str:
    """نویسندگان جداشده با ؛ یا ; → رشته‌ی خوانا با «و» فارسی"""
    parts = [a.strip() for a in re.split(r"[;؛]", authors or "") if a.strip()]
    if not parts:
        return "بی‌نام"
    if len(parts) == 1:
        return parts[0]
    return "، ".join(parts[:-1]) + " و " + parts[-1]


def format_reference(r: dict, style: str = "apa") -> str:
    """یک منبع را به سبک APA یا IEEE قالب‌بندی می‌کند"""
    authors = _fmt_authors(r.get("authors"))
    title = (r.get("title") or "").strip()
    year = (r.get("year") or "").strip()
    source = (r.get("source") or "").strip()
    vol = (r.get("volume") or "").strip()
    issue = (r.get("issue") or "").strip()
    pages = (r.get("pages") or "").strip()
    url = (r.get("url") or "").strip()
    rtype = (r.get("ref_type") or "article").strip()

    if style == "ieee":
        seg = [authors + "،", f"«{title}»,"]
        if source:
            seg.append(source + ("," if vol or issue or pages or year else ""))
        if vol:
            seg.append(f"دوره {vol},")
        if issue:
            seg.append(f"شماره {issue},")
        if pages:
            seg.append(f"ص {pages},")
        if year:
            seg.append(f"{year}.")
        if url:
            seg.append(url)
        return " ".join(s for s in seg if s).strip()

    # APA (پیش‌فرض)
    out = f"{authors} ({year or 'بی‌تا'}). {title}."
    if rtype == "book":
        if source:
            out += f" {source}."
    elif rtype in ("website",):
        if source:
            out += f" {source}."
    elif rtype == "thesis":
        out += " [پایان‌نامه]."
        if source:
            out += f" {source}."
    else:  # article / conference
        if source:
            vp = source
            if vol:
                vp += f"، {vol}"
                if issue:
                    vp += f"({issue})"
            if pages:
                vp += f"، {pages}"
            out += f" {vp}."
    if url:
        out += f" {url}"
    return out.strip()


@app.get("/api/references")
def list_references(request: Request):
    user = require_user(request)
    db = get_db()
    rows = db.execute(
        "SELECT * FROM bib_refs WHERE user_id = ? ORDER BY id DESC", (user["id"],)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.post("/api/reference")
async def add_reference(request: Request):
    user = require_user(request)
    body = await request.json()
    title = (body.get("title") or "").strip()
    if not title:
        raise HTTPException(400, "عنوان منبع الزامی است")
    vals = {f: (str(body.get(f) or "").strip())[:400] for f in REF_FIELDS}
    if vals["ref_type"] not in ("article", "book", "website", "thesis", "conference"):
        vals["ref_type"] = "article"
    db = get_db()
    cur = db.execute(
        """INSERT INTO bib_refs
           (user_id, ref_type, authors, title, year, source, volume, issue, pages, url, note)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (user["id"], vals["ref_type"], vals["authors"], vals["title"], vals["year"],
         vals["source"], vals["volume"], vals["issue"], vals["pages"], vals["url"], vals["note"]),
    )
    rid = cur.lastrowid
    db.commit()
    db.close()
    return {"id": rid, "ok": True}


@app.delete("/api/references/{ref_id}")
def delete_reference(ref_id: int, request: Request):
    user = require_user(request)
    db = get_db()
    db.execute("DELETE FROM bib_refs WHERE id = ? AND user_id = ?", (ref_id, user["id"]))
    db.commit()
    db.close()
    return {"ok": True}


@app.post("/api/references/parse")
async def parse_reference(request: Request):
    """یک ارجاع خام (کپی‌شده از مقاله) را با هوش مصنوعی به فیلدهای ساختارمند تبدیل می‌کند"""
    user = require_user(request)
    body = await request.json()
    raw = (body.get("text") or "").strip()
    if len(raw) < 8:
        raise HTTPException(400, "متن ارجاع را کامل بچسبانید")
    c = pick_model_for("article")
    if not c.get("key"):
        raise HTTPException(400, "برای تشخیص خودکار، مدیر باید کلید API را در پنل تنظیم کند")
    prompt = (
        "این یک ارجاع/منبع علمی است. آن را به JSON با این کلیدها تبدیل کن و فقط JSON خام برگردان "
        "(بدون توضیح، بدون ```): "
        '{"ref_type":"article|book|website|thesis|conference","authors":"نویسندگان جداشده با ؛",'
        '"title":"","year":"","source":"نام مجله یا ناشر","volume":"","issue":"","pages":"","url":""}\n\n'
        f"ارجاع:\n{raw[:1500]}"
    )
    try:
        out = await _call_model_once(c, prompt, max_tokens=500)
    except Exception:
        raise HTTPException(502, "تشخیص خودکار ممکن نشد؛ فیلدها را دستی وارد کنید")
    m = re.search(r"\{.*\}", out or "", re.DOTALL)
    if not m:
        raise HTTPException(422, "خروجی قابل خواندن نبود؛ فیلدها را دستی وارد کنید")
    try:
        data = json.loads(m.group(0))
    except json.JSONDecodeError:
        raise HTTPException(422, "خروجی قابل خواندن نبود؛ فیلدها را دستی وارد کنید")
    return {f: str(data.get(f, "") or "") for f in REF_FIELDS if f != "note"}


# ---------------- جست‌وجوی خودکار منابع علمی (Crossref + OpenAlex — رایگان و بدون کلید) ----------------

def _crossref_to_fields(item: dict) -> dict:
    authors = "؛ ".join(
        " ".join(p for p in [a.get("given", ""), a.get("family", "")] if p).strip()
        for a in (item.get("author") or [])
    )
    title = (item.get("title") or [""])[0]
    source = (item.get("container-title") or item.get("publisher") and [item.get("publisher")] or [""])
    source = source[0] if isinstance(source, list) else (source or "")
    dt = (item.get("issued") or {}).get("date-parts") or [[None]]
    year = str(dt[0][0]) if dt and dt[0] and dt[0][0] else ""
    rtype = "article"
    t = (item.get("type") or "").lower()
    if "book" in t:
        rtype = "book"
    elif "thesis" in t or "dissertation" in t:
        rtype = "thesis"
    elif "proceedings" in t or "conference" in t:
        rtype = "conference"
    return {
        "ref_type": rtype,
        "authors": authors,
        "title": title,
        "year": year,
        "source": source,
        "volume": str(item.get("volume") or ""),
        "issue": str(item.get("issue") or ""),
        "pages": str(item.get("page") or ""),
        "url": item.get("DOI") and f"https://doi.org/{item['DOI']}" or (item.get("URL") or ""),
    }


def _openalex_to_fields(w: dict) -> dict:
    authors = "؛ ".join(
        (a.get("author") or {}).get("display_name", "")
        for a in (w.get("authorships") or [])
    ).strip("؛ ")
    src = ((w.get("primary_location") or {}).get("source") or {}).get("display_name", "")
    bib = w.get("biblio") or {}
    doi = w.get("doi") or ""
    return {
        "ref_type": "article" if (w.get("type") or "") == "article" else (w.get("type") or "article"),
        "authors": authors,
        "title": w.get("title") or w.get("display_name") or "",
        "year": str(w.get("publication_year") or ""),
        "source": src,
        "volume": str(bib.get("volume") or ""),
        "issue": str(bib.get("issue") or ""),
        "pages": "-".join(p for p in [bib.get("first_page"), bib.get("last_page")] if p),
        "url": doi or (w.get("id") or ""),
    }


async def lookup_reference_online(query: str):
    """جست‌وجوی متادیتای مقاله در Crossref و سپس OpenAlex — تا ۵ نتیجه برمی‌گرداند"""
    query = (query or "").strip()
    if not query:
        return []
    results = []
    is_doi = bool(re.search(r"10\.\d{4,}/\S+", query))
    headers = {"User-Agent": "Antanu/1.0 (mailto:antanu@example.com)"}
    async with httpx.AsyncClient(timeout=httpx.Timeout(20, connect=10), headers=headers) as client:
        # Crossref
        try:
            if is_doi:
                doi = re.search(r"10\.\d{4,}/\S+", query).group(0)
                r = await client.get(f"https://api.crossref.org/works/{doi}")
                if r.status_code == 200:
                    results.append(_crossref_to_fields(r.json()["message"]))
            else:
                r = await client.get("https://api.crossref.org/works",
                                     params={"query.bibliographic": query, "rows": 5})
                if r.status_code == 200:
                    for it in r.json().get("message", {}).get("items", []):
                        results.append(_crossref_to_fields(it))
        except Exception:
            pass
        # OpenAlex (پشتیبان / تکمیلی)
        if len(results) < 3:
            try:
                if is_doi:
                    doi = re.search(r"10\.\d{4,}/\S+", query).group(0)
                    r = await client.get(f"https://api.openalex.org/works/https://doi.org/{doi}")
                    if r.status_code == 200:
                        results.append(_openalex_to_fields(r.json()))
                else:
                    r = await client.get("https://api.openalex.org/works",
                                         params={"search": query, "per-page": 5})
                    if r.status_code == 200:
                        for w in r.json().get("results", []):
                            results.append(_openalex_to_fields(w))
            except Exception:
                pass
    # حذف نتایج بی‌عنوان و تکراری
    seen, clean = set(), []
    for it in results:
        key = (it.get("title") or "").strip().lower()[:80]
        if not key or key in seen:
            continue
        seen.add(key)
        clean.append(it)
    return clean[:6]


@app.post("/api/references/lookup")
async def references_lookup(request: Request):
    """جست‌وجوی خودکار منبع بر اساس عنوان یا DOI و تکمیل خودکار فیلدهای APA"""
    user = require_user(request)
    body = await request.json()
    query = (body.get("query") or "").strip()
    if len(query) < 4:
        raise HTTPException(400, "عنوان مقاله یا DOI را وارد کنید (حداقل ۴ نویسه)")
    results = await lookup_reference_online(query)
    if not results:
        raise HTTPException(404, "منبعی پیدا نشد. عنوان دقیق‌تر یا DOI را امتحان کنید، یا دستی وارد کنید.")
    return {"results": results}


@app.post("/api/references/translate")
async def references_translate(request: Request):
    """ترجمه‌ی یک ارجاع انگلیسی به فارسی، بدون به‌هم‌ریختن اعداد، پرانتزها، DOI و سال"""
    user = require_user(request)
    body = await request.json()
    text = (body.get("text") or "").strip()
    if len(text) < 5:
        raise HTTPException(400, "متن ارجاع را وارد کنید")
    c = pick_model_for("translate")
    if not c.get("key"):
        raise HTTPException(400, "برای ترجمه، مدیر باید کلید API را در پنل تنظیم کند")
    prompt = (
        "این ارجاع علمی انگلیسی را به فارسی روان و آکادمیک ترجمه کن. قواعد بسیار مهم: "
        "۱) اعداد (سال، دوره، شماره، صفحات)، DOI و نشانی اینترنتی را دقیقاً و بدون تغییر و بدون جابه‌جایی نگه دار. "
        "۲) پرانتزها و علائم نگارشی سرجای خود بمانند و به‌هم نریزند. "
        "۳) نام نویسندگان را به فارسی آوانویسی کن ولی شکل لاتین را داخل پرانتز بیاور. "
        "۴) عنوان مقاله را ترجمه کن و عنوان اصلی انگلیسی را داخل پرانتز نگه دار. "
        "فقط خود ارجاع ترجمه‌شده را برگردان، بدون توضیح اضافه.\n\n"
        f"ارجاع:\n{text[:2000]}"
    )
    try:
        out = await _call_model_once(c, prompt, max_tokens=700)
    except Exception:
        raise HTTPException(502, "ترجمه ممکن نشد؛ بعداً تلاش کنید")
    return {"translated": clean_foreign_keep_latin(out or "")}


@app.post("/api/references/import")
async def references_import(request: Request):
    """استخراج فهرست منابع از یک فایل آپلودشده (Word/PDF/متن) — خطوطِ شبیه ارجاع را برمی‌گرداند"""
    user = require_user(request)
    body = await request.json()
    upload_id = body.get("upload_id")
    db = get_db()
    row = db.execute("SELECT content, kind FROM uploads WHERE id = ? AND user_id = ?",
                     (int(upload_id), user["id"])).fetchone()
    db.close()
    if not row or row["kind"] != "text":
        raise HTTPException(404, "فایل متنی یافت نشد؛ ابتدا یک فایل Word/PDF/متن آپلود کنید")
    text = row["content"] or ""
    lines = [ln.strip() for ln in re.split(r"[\r\n]+", text) if ln.strip()]
    # خطوطی که شبیه ارجاع‌اند: دارای سال داخل پرانتز یا DOI یا الگوی «نویسنده (سال)»
    cands = []
    for ln in lines:
        if len(ln) < 25:
            continue
        if re.search(r"\(\s*\d{4}\s*\)|10\.\d{4,}/|\d{4}\.\s|، \d{4}", ln) or \
           re.search(r"[A-Z][a-z]+,\s*[A-Z]\.", ln):
            cands.append(ln[:500])
    # حذف تکراری‌ها
    seen, out = set(), []
    for ln in cands:
        k = ln[:60].lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(ln)
    return {"candidates": out[:60], "count": len(out)}


@app.post("/api/references/bibliography")
async def build_bibliography(request: Request):
    """فهرست منابع مرتب‌شده به سبک APA یا IEEE می‌سازد"""
    user = require_user(request)
    body = await request.json()
    style = (body.get("style") or "apa").lower()
    if style not in ("apa", "ieee"):
        style = "apa"
    ids = body.get("ids") or []
    db = get_db()
    if ids:
        marks = ",".join("?" for _ in ids)
        rows = db.execute(
            f"SELECT * FROM bib_refs WHERE user_id = ? AND id IN ({marks})",
            (user["id"], *[int(i) for i in ids]),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT * FROM bib_refs WHERE user_id = ?", (user["id"],)
        ).fetchall()
    db.close()
    refs = [dict(r) for r in rows]
    if style == "ieee":
        # IEEE: به ترتیب ورود، شماره‌دار
        refs.sort(key=lambda r: r["id"])
        lines = [f"[{i}] {format_reference(r, 'ieee')}" for i, r in enumerate(refs, 1)]
    else:
        # APA: مرتب بر اساس نام نویسنده
        refs.sort(key=lambda r: (r.get("authors") or r.get("title") or "").strip())
        lines = [format_reference(r, "apa") for r in refs]
    return {"style": style, "count": len(refs), "text": "\n\n".join(lines)}


# ---------------- جستجوی وب (رایگان و بدون کلید) ----------------

def _search_web_sync(query: str) -> str:
    """جستجوی وب با موتور DuckDuckGo — رایگان، بدون نیاز به کلید"""
    from ddgs import DDGS
    lines = []
    with DDGS() as d:
        for r in d.text(query, max_results=5):
            title = r.get("title", "")
            body = r.get("body", "")
            href = r.get("href", "")
            lines.append(f"- {title}: {body}\n  منبع: {href}")
    return "\n".join(lines)


# ---------------- انتخاب زبان سایت ----------------

@app.middleware("http")
async def _lang_middleware(request: Request, call_next):
    """زبان این درخواست را از کوکی/مرورگر می‌خواند و در دسترس بقیه‌ی کد می‌گذارد.
    (بدون کوئری دیتابیس تا سرعت کم نشود؛ کوکی همیشه با انتخاب کاربر هم‌گام است.)"""
    try:
        cookie = request.cookies.get(LANG_COOKIE)
        if cookie and i18n.is_supported(cookie):
            _lang_ctx.set(cookie)
        else:
            _lang_ctx.set(i18n.pick_from_header(request.headers.get("accept-language", "")))
    except Exception:
        pass
    return await call_next(request)


@app.get("/api/languages")
def api_languages(request: Request):
    """فهرست زبان‌های موجود + زبان فعلی (برای منوی انتخاب زبان)."""
    user = current_user(request)
    return {"current": resolve_lang(request, user),
            "default": i18n.DEFAULT_LANG,
            "languages": i18n.language_list()}


def _apply_lang_cookie(response, code: str):
    """زبان را یک سال در کوکی نگه می‌دارد تا در بازدید بعدی هم همان باشد."""
    response.set_cookie(LANG_COOKIE, code, max_age=365 * 24 * 3600,
                        httponly=False, samesite="lax", path="/")
    return response


@app.post("/api/lang")
async def api_set_lang(request: Request):
    """تغییر زبان سایت. برای کاربر وارد‌شده در حساب هم ذخیره می‌شود
    تا روی همه‌ی دستگاه‌هایش یکسان باشد."""
    body = await request.json()
    code = i18n.normalize(body.get("lang"), fallback="")
    if not code:
        raise HTTPException(400, "Unsupported language.")
    user = current_user(request)
    if user:
        db = get_db()
        db.execute("UPDATE users SET lang = ? WHERE id = ?", (code, user["id"]))
        db.commit()
        db.close()
    out = JSONResponse({"ok": True, "lang": code, "dir": i18n.direction(code),
                        "message": i18n.t("lang.saved", code)})
    return _apply_lang_cookie(out, code)


@app.get("/lang/redirect")
def set_lang_from_query(request: Request, code: str = ""):
    """مسیر بدون جاوااسکریپت (noscript) — زبان از پارامتر code خوانده می‌شود."""
    return set_lang_redirect(code, request)


@app.get("/lang/{code}")
def set_lang_redirect(code: str, request: Request):
    """تغییر زبان بدون جاوااسکریپت (برای صفحه‌ی ورود) — بعدش به همان صفحه برمی‌گردد."""
    norm = i18n.normalize(code, fallback="")
    if not norm:
        norm = i18n.DEFAULT_LANG
    back = request.headers.get("referer") or "/"
    if not back.startswith("/") and "://" in back:
        # فقط به مسیرهای همین سایت برگرد (جلوگیری از ریدایرکت به بیرون)
        try:
            from urllib.parse import urlparse
            p = urlparse(back)
            back = p.path or "/"
            if p.query:
                back += "?" + p.query
        except Exception:
            back = "/"
    user = current_user(request)
    if user:
        db = get_db()
        db.execute("UPDATE users SET lang = ? WHERE id = ?", (norm, user["id"]))
        db.commit()
        db.close()
    return _apply_lang_cookie(RedirectResponse(back, status_code=303), norm)


# ---------------- API فهرست مدل‌ها و آپلود فایل ----------------

@app.get("/api/models")
def list_models(request: Request):
    require_user(request)
    # مدل‌های واقعی از دید کاربر پنهان است؛ فقط «آنتانو (خودکار)» نمایش داده می‌شود.
    # درخواست با شناسه‌ی auto به‌صورت خودکار روی بهترین مدل با قابلیت failover اجرا می‌شود.
    return [{"id": "auto", "name": "آنتانو (خودکار)"}]


ALLOWED_TEXT_EXT = {".txt", ".md", ".csv", ".json", ".py", ".html", ".xml", ".log"}


@app.post("/api/upload")
async def upload_file(request: Request, file: UploadFile = File(...)):
    user = require_user(request)
    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل نباید بیشتر از ۱۰ مگابایت باشد")

    name = file.filename or "file"
    ext = os.path.splitext(name)[1].lower()
    text = ""

    try:
        if ext in ALLOWED_TEXT_EXT:
            text = raw.decode("utf-8", errors="ignore")
        elif ext == ".pdf":
            import io
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif ext == ".docx":
            import io
            from docx import Document
            doc = Document(io.BytesIO(raw))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif ext in (".xlsx", ".xlsm"):
            import io
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"[برگه: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
            text = "\n".join(parts)
        elif ext in (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"):
            # عکس: فشرده‌سازی و تبدیل به base64 تا مدل‌های بینادار (مثل Gemini) آن را ببینند
            import io
            import base64
            from PIL import Image
            img = Image.open(io.BytesIO(raw))
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.thumbnail((1024, 1024))
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=82)
            b64 = base64.b64encode(buf.getvalue()).decode()
            # بندانگشتی کوچک برای نمایش در چت
            thumb = img.copy()
            thumb.thumbnail((320, 320))
            tbuf = io.BytesIO()
            thumb.save(tbuf, format="JPEG", quality=70)
            thumb_b64 = base64.b64encode(tbuf.getvalue()).decode()
            db = get_db()
            cur = db.execute(
                "INSERT INTO uploads (user_id, filename, content, kind) VALUES (?, ?, ?, 'image')",
                (user["id"], name[:200], b64),
            )
            db.commit()
            upload_id = cur.lastrowid
            db.close()
            return {"id": upload_id, "filename": name, "kind": "image",
                    "preview": f"data:image/jpeg;base64,{thumb_b64}"}
        elif ext in (".mp4", ".mov", ".avi", ".mkv", ".webm", ".3gp"):
            raise HTTPException(
                400,
                "تحلیل ویدیو هنوز توسط سرویس‌های هوش مصنوعی رایگان پشتیبانی نمی‌شود. "
                "می‌توانید عکس (اسکرین‌شات از ویدیو) بفرستید تا تحلیل شود.",
            )
        else:
            raise HTTPException(
                400,
                "این نوع فایل فعلاً پشتیبانی نمی‌شود. فرمت‌های مجاز: عکس (jpg/png)، pdf، docx، xlsx، txt، csv، md، json",
            )
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "خواندن محتوای فایل ممکن نشد. از سالم بودن فایل مطمئن شوید.")

    text = text.strip()
    if not text:
        raise HTTPException(400, "متنی در این فایل پیدا نشد.")

    db = get_db()
    cur = db.execute(
        "INSERT INTO uploads (user_id, filename, content, kind) VALUES (?, ?, ?, 'text')",
        (user["id"], name[:200], text[:60000]),
    )
    db.commit()
    upload_id = cur.lastrowid
    db.close()
    return {"id": upload_id, "filename": name, "kind": "text", "chars": len(text)}


# ---------------- API چت (استریم — چندمدلی + جستجوی وب + فایل) ----------------

# لحن و سبک پاسخ که کاربر از نوار ابزار انتخاب می‌کند
TONE_INSTRUCTIONS = {
    "academic": "لحن پاسخ را رسمی، علمی و دانشگاهی نگه دار؛ از اصطلاحات دقیق و ساختار منظم استفاده کن.",
    "simple": "پاسخ را با زبان ساده، روان و خودمانی بنویس تا برای همه قابل‌فهم باشد؛ از اصطلاحات پیچیده پرهیز کن.",
    "concise": "پاسخ را کوتاه، مختصر و مستقیم بده؛ فقط نکات کلیدی و بدون حاشیه.",
    "detailed": "پاسخ را کامل، مفصل و با جزئیات، مثال و توضیح گام‌به‌گام ارائه بده.",
    "creative": "پاسخ را خلاقانه، جذاب و با نگاه تازه بنویس؛ در صورت تناسب از تشبیه و مثال‌های ملموس استفاده کن.",
}


# زبان‌هایی که تشخیص‌دهنده می‌شناسد ولی در فهرست رابط سایت نیستند؛
# اگر کاربر با یکی از این‌ها بنویسد، باز هم باید به همان زبان پاسخ بگیرد.
_EXTRA_LANG_NAMES = {
    "ru": "Russian", "uk": "Ukrainian", "it": "Italian", "pt": "Portuguese",
    "nl": "Dutch", "pl": "Polish", "sv": "Swedish", "he": "Hebrew",
    "ur": "Urdu", "th": "Thai", "vi": "Vietnamese", "id": "Indonesian",
}


def _lang_display_name(code: str) -> str:
    """نام انگلیسیِ زبان برای نوشتن در پرامپت."""
    c = (code or "").strip().lower()
    if i18n.is_supported(c):
        return i18n.english_name(c)
    return _EXTRA_LANG_NAMES.get(c, c.upper() or "English")


def build_system_prompt(user, memories, tone: str | None = None, lang: str | None = None,
                        reply_lang: str | None = None) -> str:
    """پرامپت سیستمی.

    lang       = زبان انتخاب‌شده‌ی سایت (پیش‌فرض، وقتی زبان پیام مبهم است)
    reply_lang = زبانی که کاربر واقعاً با آن نوشته؛ پاسخ باید به همین زبان باشد.
    """
    prompt = BASE_SYSTEM_PROMPT
    # زبان پاسخ: آنتانو به همان زبانی جواب می‌دهد که کاربر پیام داده،
    # حتی اگر زبان رابط سایت چیز دیگری باشد.
    code = i18n.normalize(reply_lang or lang)
    if code != "fa":
        lname = _lang_display_name(reply_lang or lang)
        prompt += (
            f"\n\n[RESPONSE LANGUAGE — HIGHEST PRIORITY] The user wrote to you in {lname}. "
            f"Write every answer in {lname}, using that language's own script and natural style. "
            f"If the user later writes in a different language, switch to that language too — "
            f"always mirror the language of the user's latest message. "
            f"This overrides any earlier instruction that told you to answer in Persian: "
            f"ignore those Persian-only writing rules and apply the same care "
            f"(clean grammar, correct punctuation, no mixed-in foreign words, no broken sentences) to {lname} instead. "
            f"Only switch away from {lname} when the user explicitly asks for another language "
            f"or asks you to translate something. Keep the [[TR: ...]] format rule for translations."
        )
    else:
        prompt += (
            "\n\n[زبان پاسخ] کاربر به فارسی نوشته است، پس پاسخ را کامل به فارسی بنویس. "
            "اگر در پیام بعدی به زبان دیگری نوشت، به همان زبان جواب بده."
        )
    prompt += (f"\n\nتاریخ و ساعت کنونی: {now_string()}. "
               "هر جا درباره تاریخ، روز، سال یا ساعت پرسیده شد، دقیقاً از همین استفاده کن و نگو که نمی‌دانی.")
    if user["stars"] >= 4:
        prompt += (" وقتی کاربر توضیح کامل یا محتوای بلند خواست، پاسخ را عمیق و جامع ارائه بده — "
                   "ولی برای پرسش‌های ساده همچنان کوتاه و مستقیم جواب بده.")
    else:
        prompt += " پاسخ‌ها را متناسب با پرسش، مختصر و مفید نگه دار."
    if tone and tone in TONE_INSTRUCTIONS:
        prompt += "\n\nسبک پاسخ (خواسته‌ی کاربر): " + TONE_INSTRUCTIONS[tone]
    if memories:
        prompt += "\n\nحافظه بلندمدت این کاربر (همیشه در نظر بگیر):\n"
        prompt += "\n".join(f"- {m['content']}" for m in memories)
    return prompt


async def _learn_glossary_words(message: str):
    """کلمات ارمنیِ به‌کاررفته که در واژه‌نامه نیستند را با کمک مدل ترجمه و در «مغز واژگان» ثبت می‌کند.
    در پس‌زمینه اجرا می‌شود و هیچ اثری روی سرعت پاسخ کاربر ندارد."""
    try:
        import glossary as gl
        _, unknown = gl.known_and_unknown(message, "hy")
        unknown = [w for w in unknown if len(w) >= 2][:12]
        if not unknown:
            return
        c = pick_model_for("fast")
        if not c or not c.get("key"):
            return
        prompt = (
            "برای هر کلمه/عبارت ارمنی زیر یک شیء JSON بده. خروجی فقط یک آرایه‌ی JSON معتبر باشد، بدون هیچ توضیح. "
            'قالب هر عضو: {"term":"کلمه ارمنی","translation":"ترجمه فارسی","pronunciation":"تلفظ با حروف انگلیسی"}\n\n'
            "کلمات:\n" + "\n".join(unknown)
        )
        out = await _call_model_once(c, prompt, max_tokens=900)
        import re as _re
        m = _re.search(r"\[.*\]", out or "", _re.S)
        if not m:
            return
        arr = json.loads(m.group(0))
        for item in arr:
            if isinstance(item, dict) and item.get("term") and item.get("translation"):
                gl.add_term(item["term"], item["translation"], item.get("pronunciation", ""),
                            lang="hy", source="auto", status="done")
    except Exception:
        pass


class ModelError(Exception):
    def __init__(self, status: int, body: str = ""):
        self.status = status
        self.body = body


def _save_gen_image(dataurl: str):
    """عکس/ویدیوی تولیدشده مدل را به فایل واقعی تبدیل می‌کند و نامش را برمی‌گرداند"""
    try:
        import base64
        import export_utils
        header, b64 = dataurl.split(",", 1)
        ext = "png"
        if "jpeg" in header or "jpg" in header:
            ext = "jpg"
        elif "webp" in header:
            ext = "webp"
        elif "mp4" in header:
            ext = "mp4"
        elif "webm" in header:
            ext = "webm"
        name = f"antanu-img-{secrets.token_hex(6)}.{ext}"
        with open(os.path.join(export_utils.EXPORT_DIR, name), "wb") as f:
            f.write(base64.b64decode(b64))
        return name
    except Exception:
        return None


async def stream_model(messages, stars: int, model: str, base: str, key: str,
                       filter_foreign: bool = True):
    """استریم پاسخ از هر سرویس سازگار با OpenAI — هر مدل با آدرس و کلید خودش.
    filter_foreign=False برای زبان‌های غیرفارسی، تا خط آن زبان پاک نشود."""
    payload = {"model": model, "messages": messages, "stream": True}
    if MAX_TOKENS.get(stars, -1) > 0:
        payload["max_tokens"] = MAX_TOKENS[stars]

    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    ffilter = ForeignFilter(filter_foreign)   # فیلتر حالت‌دار: بلوک‌های [[TR: ...]] را دست‌نخورده رد می‌کند

    async with httpx.AsyncClient(timeout=httpx.Timeout(120, connect=15)) as client:
        async with client.stream(
            "POST", f"{base}/chat/completions", json=payload, headers=headers
        ) as r:
            if r.status_code != 200:
                body = (await r.aread()).decode(errors="ignore")[:400]
                raise ModelError(r.status_code, body)
            async for line in r.aiter_lines():
                if not line.startswith("data:"):
                    continue
                data = line[5:].strip()
                if data == "[DONE]":
                    break
                try:
                    obj = json.loads(data)
                except json.JSONDecodeError:
                    continue
                choices = obj.get("choices") or [{}]
                delta = choices[0].get("delta") or {}
                chunk = ffilter.feed(delta.get("content") or "")
                # عکس‌های تولیدشده توسط مدل‌های عکس‌ساز (مثل gemini-2.5-flash-image)
                for im in (delta.get("images") or []) + (delta.get("videos") or []):
                    url = ((im or {}).get("image_url") or (im or {}).get("video_url") or {}).get("url") or ""
                    if url.startswith("data:image"):
                        fname = _save_gen_image(url)
                        if fname:
                            chunk += f"\n\n[[ANTANU_IMG:/download/{fname}]]\n\n"
                    elif url.startswith("data:video"):
                        fname = _save_gen_image(url)
                        if fname:
                            chunk += f"\n\n[[ANTANU_VID:/download/{fname}]]\n\n"
                if chunk:
                    yield chunk
    tail = ffilter.flush()
    if tail:
        yield tail


@app.post("/api/chat")
async def api_chat(request: Request):
    user = require_user(request)
    if not chat_rate_ok(user["id"]):
        raise HTTPException(429, "پیام‌ها را خیلی سریع می‌فرستید. چند لحظه صبر کنید و دوباره تلاش کنید.")
    body = await request.json()
    message = (body.get("message") or "").strip()
    conv_id = body.get("conversation_id")
    selected_ids = body.get("models") or ["auto"]
    web_on = bool(body.get("web"))
    research = bool(body.get("research"))
    tone = (body.get("tone") or "").strip()
    attachment_ids = body.get("attachments") or []
    if not message:
        raise HTTPException(400, "پیام خالی است")

    # ---------- دستور ذخیره در حافظه:  \save متن  یا  /save ----------
    save_match = re.match(r"^[\\/]\s*save\b[:\s]*", message, re.IGNORECASE)
    if save_match:
        remainder = message[save_match.end():].strip()
        db = get_db()
        if conv_id:
            conv = db.execute(
                "SELECT id FROM conversations WHERE id = ? AND user_id = ?", (conv_id, user["id"])
            ).fetchone()
            if not conv:
                db.close()
                raise HTTPException(404, "گفتگو یافت نشد")
        else:
            cur = db.execute(
                "INSERT INTO conversations (user_id, title) VALUES (?, ?)",
                (user["id"], "🧠 ذخیره در حافظه"),
            )
            conv_id = cur.lastrowid
        db.execute(
            "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'user', ?)",
            (conv_id, message),
        )
        if not remainder:
            row = db.execute(
                "SELECT content FROM messages WHERE conversation_id = ? AND role = 'assistant' "
                "ORDER BY id DESC LIMIT 1",
                (conv_id,),
            ).fetchone()
            remainder = (row["content"] if row else "").strip()
        if remainder:
            db.execute(
                "INSERT INTO memories (user_id, content) VALUES (?, ?)",
                (user["id"], remainder[:4000]),
            )
            reply = "🧠 در حافظه بلندمدت آنتانو ذخیره شد و هرگز فراموش نمی‌شود."
        else:
            reply = "⚠️ چیزی برای ذخیره پیدا نشد. بنویسید: \\save متن موردنظر — یا بعد از پاسخ ربات فقط \\save بفرستید تا همان پاسخ ذخیره شود."
        db.execute(
            "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'assistant', ?)",
            (conv_id, reply),
        )
        db.commit()
        db.close()

        async def save_gen():
            yield reply

        return StreamingResponse(
            save_gen(),
            media_type="text/plain; charset=utf-8",
            headers={"X-Conversation-Id": str(conv_id)},
        )

    # مدل‌های انتخاب‌شده توسط کاربر (یکی یا چندتا)
    catalog = get_ai_catalog()
    chosen = [c for c in catalog if c["id"] in selected_ids]
    # اگر کاربر خودش مدل مشخصی (غیر از «خودکار») انتخاب کرده، مسیریاب دخالت نمی‌کند
    explicit_pick = bool(chosen) and not (len(chosen) == 1 and chosen[0]["id"] == "auto")
    if not chosen:
        chosen = [catalog[0]]

    # حالت تحقیق گروهی: همه هوش مصنوعی‌ها دست‌به‌دست هم پژوهش را کامل می‌کنند + جستجوی وب
    if research:
        chosen = list(catalog)
        web_on = True

    # تشخیص نوع درخواست (چت / ساخت عکس / ساخت ویدیو) برای سهمیه
    kind = "chat"
    if re.search(r"(عکس|تصویر|نقاشی|طرح)\s*(.{0,20})?(بساز|درست کن|تولید کن|بکش|ایجاد کن)", message) \
            or re.search(r"(بساز|تولید کن).{0,12}(عکس|تصویر)", message):
        kind = "image"
    if re.search(r"(ویدیو|ویدئو|فیلم|کلیپ|انیمیشن)\s*(.{0,20})?(بساز|درست کن|تولید کن|ایجاد کن)", message) \
            or re.search(r"(بساز|تولید کن).{0,12}(ویدیو|ویدئو|فیلم|کلیپ)", message):
        kind = "video"

    db = get_db()
    quota_check(db, user, kind)

    # گفتگو
    if conv_id:
        conv = db.execute(
            "SELECT id FROM conversations WHERE id = ? AND user_id = ?", (conv_id, user["id"])
        ).fetchone()
        if not conv:
            db.close()
            raise HTTPException(404, "گفتگو یافت نشد")
    else:
        cur = db.execute(
            "INSERT INTO conversations (user_id, title) VALUES (?, ?)",
            (user["id"], message[:60]),
        )
        conv_id = cur.lastrowid

    # محتوای فایل‌های پیوست (متن و عکس جداگانه)
    extra_context = ""
    attach_names = []
    images_b64 = []
    for aid in attachment_ids[:5]:
        row = db.execute(
            "SELECT filename, content, kind FROM uploads WHERE id = ? AND user_id = ?",
            (int(aid), user["id"]),
        ).fetchone()
        if not row:
            continue
        attach_names.append(row["filename"])
        if row["kind"] == "image":
            images_b64.append(row["content"])
        else:
            extra_context += f"\n\n[محتوای فایل پیوست «{row['filename']}»]:\n{row['content'][:20000]}"

    # جستجوی وب — اگر سؤال درباره اطلاعات روز باشد، خودکار روشن می‌شود
    if not web_on and any(w in message for w in AUTO_SEARCH_WORDS):
        web_on = True

    web_note = ""
    if web_on:
        try:
            results = await run_in_threadpool(_search_web_sync, message[:300])
            if results:
                extra_context += (
                    "\n\n[نتایج جستجوی وب — برای پاسخ به‌روز از این‌ها استفاده کن و منبع را ذکر کن]:\n"
                    + results
                )
        except Exception:
            web_note = "\n\n_🌐 جستجوی وب در دسترس نبود؛ پاسخ از دانش خود مدل است._"

    # پیام نمایشی که در تاریخچه ذخیره می‌شود
    display = message
    if attach_names:
        display += "\n📎 " + "، ".join(attach_names)
    if web_on:
        display += "\n🌐 با جستجوی وب"

    db.execute(
        "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'user', ?)",
        (conv_id, display),
    )
    db.commit()

    history = db.execute(
        "SELECT role, content FROM messages WHERE conversation_id = ? ORDER BY id DESC LIMIT 24",
        (conv_id,),
    ).fetchall()[::-1]
    memories = db.execute(
        "SELECT content FROM memories WHERE user_id = ? ORDER BY id DESC LIMIT 30", (user["id"],)
    ).fetchall()
    db.close()

    # دانش خودآموز: پاسخ‌های مشابه قبلی از پایگاه دانش آنتانو
    kb = search_knowledge(message, limit=2)
    # زبان پاسخ = زبانی که کاربر با آن نوشته است. زبان انتخاب‌شده‌ی سایت فقط وقتی
    # ملاک است که پیام کوتاه یا مبهم باشد (مثل «ok» یا یک ایموجی).
    user_lang = resolve_lang(request, user)
    reply_lang, _sure = i18n.detect_message_lang(message, fallback=user_lang)
    # این خط حیاتی است: فیلترِ «حذف حروف غیرفارسی» از همین متغیر می‌خواند. اگر
    # زبان سایت فارسی باشد و کاربر چینی/روسی بنویسد، بدون این خط کل پاسخ پاک می‌شود.
    _lang_ctx.set(reply_lang)
    sys_content = build_system_prompt(user, memories, tone, lang=user_lang,
                                      reply_lang=reply_lang)
    if kb:
        sys_content += "\n\nدانش پیشین آنتانو (از گفتگوهای قبلی، در صورت مرتبط بودن استفاده کن):\n"
        sys_content += "\n".join(f"پرسش: {k['question']}\nپاسخ: {k['answer'][:800]}" for k in kb)

    # کتابخانه‌ی دائمی: بخش‌های مرتبط از کتاب‌ها و جزوه‌های ذخیره‌شده
    sys_content += _library_context(message)

    # اگر پیام «درخواست ترجمه» است، آنتانو نباید به سلام/احوالپرسیِ داخل متن پاسخ دهد
    _tr_target = None
    try:
        _tr_target = _detect_translate_intent(message)
        if _tr_target is not None:
            sys_content += TRANSLATE_ONLY_NOTE.format(
                tgt=f" — زبان مقصد: {_tr_target}" if _tr_target else ""
            )
    except Exception:
        pass

    # مغز واژگان: اگر پیام حاوی کلمه‌های زبان‌های واژه‌نامه‌ای (مثل ارمنی) باشد،
    # ترجمه‌ها و تلفظ‌های ثبت‌شده به مدل داده می‌شود تا خروجی سه‌ستونه و دقیق بدهد.
    try:
        import glossary as _gl
        _force = set()
        if _tr_target is not None:
            _force = {c for n, c in _lang_alias_map().items() if n == _tr_target} or set()
        _ghint = _gl.prompt_hint_for_text(message, force_langs=_force)
        if _ghint:
            sys_content += "\n" + _ghint
    except Exception:
        pass

    msgs = [{"role": "system", "content": sys_content}]
    msgs += [{"role": r["role"], "content": r["content"]} for r in history[:-1]]
    if images_b64:
        # پیام چندرسانه‌ای: متن + عکس‌ها (برای مدل‌های بینادار مثل Gemini و GPT-4o)
        img_instruction = message if message else "این تصویر را با دقت و جزئیات کامل به فارسی توصیف و تحلیل کن."
        img_instruction += ("\n\n[راهنما: تو یک مدل بینا هستی؛ تصویر را مستقیماً ببین و "
                            "همه اشیا، متن‌ها، اعداد، نمودارها و جزئیات آن را به فارسی روان توضیح بده. "
                            "هرگز نگو که نمی‌توانی عکس ببینی.]")
        parts = [{"type": "text", "text": img_instruction + extra_context}]
        for b64 in images_b64[:4]:
            parts.append({"type": "image_url",
                          "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})
        msgs.append({"role": "user", "content": parts})
    else:
        msgs.append({"role": "user", "content": message + extra_context})

    stars = user["stars"]

    # ---------- سوئیچ خودکار مدل بر اساس نوع درخواست ----------
    # نوع کار پیام تشخیص داده می‌شود (ترجمه، کد، آمار، مقاله، دیدن عکس، پاسخ سریع…)
    # و مدلی که مدیر برای همان کار در پنل تعیین کرده، اولِ صف می‌آید.
    # اگر برای آن کار مدلی تعیین نشده باشد، رفتار مثل قبل است (مدل اصلی اول).
    route_order = list(catalog)
    if not research and not explicit_pick:
        try:
            import router as _router
            _task = _router.detect_task(
                message,
                has_images=bool(images_b64),
                has_files=bool(attach_names),
                is_translate=_tr_target is not None,
                kind=kind,
            )
            route_order = _router.order_catalog(catalog, _task)
            if route_order and route_order[0].get("key"):
                chosen = [route_order[0]]
        except Exception:
            route_order = list(catalog)

    async def gen():
        full = ""
        c0 = chosen[0]

        async def try_stream(messages):
            """تلاش روی مدلِ مناسبِ همین کار و در صورت خطا، خودکار روی بقیه مدل‌های دارای کلید"""
            nonlocal full
            ordered = [c for c in route_order if c.get("key")]
            candidates = [c0] + [c for c in ordered if c["id"] != c0["id"]]
            last_err = None
            for c in candidates:
                buf = ""
                emitted = False
                try:
                    async for chunk in stream_model(messages, stars, c["model"], c["base"], c["key"],
                                                    filter_foreign=(user_lang == "fa")):
                        if not emitted:
                            buf += chunk
                            if len(buf) >= 60:
                                if JUNK_RE.match(buf.strip()):
                                    raise ModelError(590, "junk")
                                full += buf
                                yield buf
                                emitted = True
                            continue
                        full += chunk
                        yield chunk
                    # پایان استریم
                    if not emitted and buf.strip():
                        if JUNK_RE.match(buf.strip()):
                            raise ModelError(590, "junk")
                        full += buf
                        yield buf
                        emitted = True
                    if emitted:
                        return
                    last_err = ModelError(591, "empty")  # پاسخ خالی → مدل بعدی
                except ModelError as e:
                    if emitted:
                        return  # وسط پاسخ قطع شد؛ همان بخش را نگه می‌داریم
                    last_err = e
                    continue

            # هیچ مدلی پاسخ سالم نداد
            e = last_err
            if images_b64 and e and e.status in (400, 422):
                note = ("⚠️ آنتانو الان توان دیدن این عکس را ندارد. "
                        "ظرفیت پردازش تصویر پر شده است؛ کمی بعد دوباره تلاش کنید.")
            elif e and e.status in (401, 403):
                note = "⚠️ سرویس آنتانو موقتاً در دسترس نیست. کمی بعد دوباره تلاش کنید."
            elif e and e.status == 429:
                note = ("⚠️ ظرفیت پاسخ‌گویی آنتانو فعلاً پر شده است. "
                        "کمی بعد دوباره تلاش کنید (ظرفیت هر شب آزاد می‌شود).")
            elif e and e.status == 590:
                note = "⚠️ مدل‌های متصل پاسخ نامعتبر دادند. مدیر سیستم مدل‌های خراب را از پنل حذف کند."
            else:
                note = "⚠️ هیچ سرویس هوش مصنوعی‌ای در دسترس نبود. کمی بعد دوباره تلاش کنید."
            full += note
            yield note

        try:
            if not c0.get("key"):
                # حالت بدون API: از پایگاه دانش آنتانو جواب بده
                kb_ans = search_knowledge(message, limit=1)
                if kb_ans:
                    ans = "🧠 (از حافظه آنتانو)\n\n" + kb_ans[0]["answer"]
                    full += ans
                    yield ans
                else:
                    err = ("⚠️ در حال حاضر هیچ هوش مصنوعی‌ای متصل نیست و پاسخ این پرسش هم در حافظه آنتانو ذخیره نشده. "
                           "مدیر سیستم باید از پنل مدیریت یک کلید API وصل کند.")
                    full += err
                    yield err
            elif len(chosen) == 1 or images_b64:
                # تک‌مدلی یا وقتی عکس هست: استریم با failover خودکار
                async for chunk in try_stream(msgs):
                    yield chunk
            else:
                # چند هوش مصنوعی: پیش‌نویس موازی + پاسخ واحد از زبان آنتانو
                async def draft(c):
                    if not c.get("key"):
                        return ""
                    try:
                        out = await _call_model_once(c, messages=msgs, max_tokens=1200)
                        return "" if JUNK_RE.match((out or "").strip()) else out
                    except Exception:
                        return ""

                drafts = await asyncio.gather(*[draft(c) for c in chosen])
                drafts = [d.strip() for d in drafts if d and d.strip()]

                if not drafts:
                    async for chunk in try_stream(msgs):
                        yield chunk
                else:
                    combo = "\n\n═══ پیش‌نویس بعدی ═══\n\n".join(d[:3500] for d in drafts)
                    goal = ("یک پژوهش کامل، عمیق و ساختارمند با عنوان‌بندی" if research
                            else "یک پاسخ واحد، کامل و منسجم")
                    # زبان خروجیِ ترکیب هم باید همان زبان پیام کاربر باشد
                    _lname = _lang_display_name(reply_lang)
                    _lang_rule = (
                        "تکرارها را حذف کن، اشتباه‌ها را اصلاح کن و فقط به فارسیِ معیارِ روان و درست بنویس؛ "
                        "هیچ حرف یا واژه غیرفارسی به کار نبر مگر معادل انگلیسی داخل پرانتز. "
                        if reply_lang == "fa" else
                        f"Remove repetition, fix mistakes, and write the whole answer in {_lname} "
                        f"using clean grammar and that language's own script. "
                    )
                    synth_msgs = [
                        {"role": "system", "content": build_system_prompt(
                            user, memories, tone, lang=user_lang, reply_lang=reply_lang)},
                        {"role": "user", "content":
                            f"پرسش کاربر: {message}\n\n"
                            f"چند پیش‌نویس داخلی برای پاسخ آماده شده است:\n{combo}\n\n"
                            f"بر پایه بهترین نکات همه پیش‌نویس‌ها، {goal} بنویس. "
                            + _lang_rule +
                            "هرگز به وجود پیش‌نویس‌ها یا چند دستیار اشاره نکن — پاسخ فقط از زبان خودت (آنتانو) باشد."},
                    ]
                    async for chunk in try_stream(synth_msgs):
                        yield chunk

            if web_note:
                full += web_note
                yield web_note

        except (httpx.ConnectError, httpx.ReadError, httpx.ConnectTimeout, httpx.RemoteProtocolError):
            err = "⚠️ اتصال به سرویس هوش مصنوعی برقرار نشد. کمی بعد دوباره تلاش کنید."
            full += err
            yield err
        finally:
            d = get_db()
            d.execute(
                "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'assistant', ?)",
                (conv_id, full or "…"),
            )
            # ثبت مصرف — سهمیه فقط روی نتیجه‌ی واقعاً موفق حساب می‌شود
            if kind == "chat":
                est_tokens = max(1, (len(message) + len(full)) // 3)
                quota_add(d, user, "chat", est_tokens)
            elif kind == "image" and "[[ANTANU_IMG:" in full:
                quota_add(d, user, "image")
            elif kind == "video" and "[[ANTANU_VID:" in full:
                quota_add(d, user, "video")
            d.commit()
            d.close()
            # ذخیره خودکار در پایگاه دانش آنتانو (یادگیری از گفتگوها)
            if not images_b64:
                add_knowledge(message, full)
            # یادگیری خودکار کلمات جدید واژه‌نامه‌ای (مثل ارمنی) در پس‌زمینه
            try:
                asyncio.create_task(_learn_glossary_words(message))
            except Exception:
                pass

    return StreamingResponse(
        gen(),
        media_type="text/plain; charset=utf-8",
        headers={"X-Conversation-Id": str(conv_id)},
    )


# ---------------- تحلیل آماری (فصل چهارم رساله) ----------------

@app.post("/api/stats/upload")
async def stats_upload(request: Request, file: UploadFile = File(...)):
    """آپلود فایل داده آماری برای تحلیل"""
    user = require_user(request)
    check_subscription(user)
    raw = await file.read()
    if len(raw) > 20 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل داده نباید بیشتر از ۲۰ مگابایت باشد")
    name = file.filename or "data.csv"
    ext = os.path.splitext(name)[1].lower()
    # هر فرمتی مجاز است — از جمله PDF و Word (جدول‌ها/داده‌ها خودکار استخراج می‌شوند).
    # فقط فایل‌هایی که هیچ داده‌ی متنی/جدولی ندارند (عکس/ویدیو/صوت/آرشیو) رد می‌شوند.
    NON_DATA = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff",
                ".mp4", ".mov", ".avi", ".mkv", ".mp3", ".wav", ".zip", ".rar", ".7z", ".exe"}
    if ext in NON_DATA:
        raise HTTPException(
            400,
            "این فایل داده ندارد (عکس/ویدیو/صوت). برای تحلیل آماری، فایل داده (Excel، CSV، SPSS، Stata، JSON، "
            "PDF یا Word) بفرستید.",
        )
    if not ext:
        ext = ".csv"  # فایل بدون پسوند را به‌عنوان CSV تلاش می‌کنیم
    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
    fname = f"data-{secrets.token_hex(6)}{ext}"
    path = os.path.join(export_utils.EXPORT_DIR, fname)
    with open(path, "wb") as f:
        f.write(raw)
    # توصیف اولیه
    try:
        try:
            import stats_engine
        except ImportError as _e:
            raise HTTPException(500, f"کتابخانه‌های تحلیل آماری نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
        overview = stats_engine.run("overview", path)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"خواندن داده ممکن نشد: {e}")
    if isinstance(overview, dict) and overview.get("error"):
        hint = ""
        if ext in (".pdf", ".docx", ".doc"):
            hint = (" اگر جدول در این فایل به‌صورت متن پیوسته است، آن را به Excel/CSV تبدیل کنید "
                    "یا مطمئن شوید ستون‌ها با فاصله/تب از هم جدا شده‌اند.")
        raise HTTPException(400, f"خواندن داده‌ی جدولی از این فایل ممکن نشد.{hint}")
    return {"file": fname, "overview": overview}


@app.post("/api/stats/run")
async def stats_run(request: Request):
    """اجرای یک تحلیل آماری و تفسیر دانشگاهی با هوش مصنوعی"""
    user = require_user(request)
    check_subscription(user)
    body = await request.json()
    fname = body.get("file") or ""
    analysis = body.get("analysis") or "overview"
    params = body.get("params") or {}
    if not re.fullmatch(r"[A-Za-z0-9._-]+", fname):
        raise HTTPException(400, "نام فایل نامعتبر")
    try:
        import export_utils, stats_engine
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌ها نصب نیستند. دستور: pip install -r requirements.txt (جزئیات: {_e})")
    path = os.path.join(export_utils.EXPORT_DIR, fname)
    if not os.path.exists(path):
        raise HTTPException(404, "فایل داده پیدا نشد؛ دوباره آپلود کنید")

    result = await run_in_threadpool(stats_engine.run, analysis, path, **params)
    if "error" in result:
        return {"result": result, "interpretation": ""}

    # تفسیر دانشگاهی توسط هوش مصنوعی
    c = pick_model_for("stats")
    interpretation = ""
    if c.get("key"):
        try:
            prompt = (
                f"تو یک متخصص آمار و مشاور روش تحقیق هستی. این خروجی واقعی یک تحلیل «{analysis}» است "
                f"که با نرم‌افزار آماری روی داده‌های کاربر اجرا شده:\n\n{json.dumps(result, ensure_ascii=False)}\n\n"
                "این نتایج را به زبان فارسی دانشگاهی و دقیق برای فصل چهارم رساله تفسیر کن: "
                "معناداری‌ها را توضیح بده، فرضیه‌ها را تأیید یا رد کن، و در صورت وجود مشکل روش‌شناختی هشدار بده. "
                "نمونه لحن درست: «با توجه به ضریب مسیر ۰٫۴۲ و آماره t برابر ۳٫۱۸، اثر متغیر … معنادار است»."
            )
            interpretation = await _call_model_once(c, prompt, system=BASE_SYSTEM_PROMPT, max_tokens=1500)
            quota_add_db = get_db()
            quota_add(quota_add_db, user, "chat", 3000)
            quota_add_db.commit(); quota_add_db.close()
        except Exception:
            interpretation = ""

    # نمودارها (SmartPLS و سبک SPSS) — پس از تفسیر ساخته می‌شوند تا در پرامپت نیایند
    try:
        import sem_plot
        if analysis in ("sem_pls", "pls") and result.get("constructs"):
            png = await run_in_threadpool(sem_plot.render_pls_diagram, result,
                                          "مدل مسیر پژوهش (ضرایب استاندارد و بارهای عاملی)")
        elif analysis in ("overview", "assumptions"):
            png = await run_in_threadpool(sem_plot.render_histograms, path, params.get("cols"))
        elif analysis == "correlation":
            png = await run_in_threadpool(sem_plot.render_corr_heatmap, path,
                                          params.get("method", "pearson"), params.get("cols"))
        else:
            png = None
        if png:
            result["diagram"] = f"/download/{png}"
    except Exception as _e:
        print("[ANTANU] stats chart failed:", _e)

    return {"result": result, "interpretation": interpretation}


@app.get("/api/ticker")
async def api_ticker(request: Request, lang: str = ""):
    """متن نوار تبلیغاتی به زبان کاربر — در صورت نیاز ترجمه و ذخیره می‌شود.

    مدیر متن را یک بار به فارسی می‌نویسد و هر کاربر آن را به زبان خودش می‌بیند.
    ترجمه فقط یک بار برای هر زبان گرفته می‌شود و بعد از انبار می‌آید.
    """
    text = (get_setting("ticker_text", "") or "").strip()
    if not text:
        return {"html": "", "dir": "rtl", "speed": 25}
    try:
        speed = max(5, min(int(get_setting("ticker_speed", "25") or 25), 180))
    except (TypeError, ValueError):
        speed = 25

    target = i18n.normalize(lang or resolve_lang(request, current_user(request)))
    src = ticker_source_lang(text)
    if target == src:
        return _ticker_payload(text, speed)

    key = _ticker_cache_key(text, target)
    cached = get_setting(key, "")
    if cached:
        return _ticker_payload(cached, speed)

    # نشانی‌ها نباید ترجمه شوند: جایشان نشانه می‌گذاریم و بعد برمی‌گردانیم
    urls = []

    def _hold(m):
        urls.append(m.group(1))
        return f"⟦{len(urls) - 1}⟧"

    masked = _TICKER_URL_RE.sub(_hold, text)

    translated = ""
    c = pick_model_for("translate")
    if c.get("key"):
        try:
            tname = _lang_display_name(target)
            out = await _call_model_once(
                c,
                f"Translate the following advertising banner into {tname}. "
                f"Keep it short and natural, like an ad line. "
                f"Keep every placeholder such as ⟦0⟧ exactly as it is and in place. "
                f"Keep emojis. Reply with the translation only — no quotes, no explanation.\n\n"
                f"{masked}",
                system="You are a professional advertising translator. Output only the translation.",
                max_tokens=400, keep_foreign=True)
            translated = (out or "").strip().strip('"').strip()
        except Exception:
            translated = ""

    if not translated:
        return _ticker_payload(text, speed)      # ترجمه نشد → همان متن اصلی

    for i, u in enumerate(urls):
        translated = translated.replace(f"⟦{i}⟧", u).replace(f"[{i}]", u)
    # اگر مدل نشانه‌ها را خورد، نشانی‌ها را ته متن برگردان تا لینک از دست نرود
    for u in urls:
        if u not in translated:
            translated += " " + u

    set_setting(key, translated[:1000])
    return _ticker_payload(translated, speed)


@app.post("/api/stats/auto")
async def stats_auto(request: Request):
    """تحلیل خودکار: کاربر فایل را داده و به زبان ساده می‌گوید چه می‌خواهد.

    آنتانو ساختار داده را می‌شناسد، نقشه‌ی تحلیل می‌سازد، آن را اعتبارسنجی و اجرا
    می‌کند و گزارش مرتب می‌دهد. پاسخ جریانی است تا کاربر پیشرفت کار را ببیند.
    """
    user = require_user(request)
    check_subscription(user)
    body = await request.json()
    fname = body.get("file") or ""
    ask = (body.get("request") or "").strip()[:1200]
    want_files = bool(body.get("export", True))
    if not re.fullmatch(r"[A-Za-z0-9._-]+", fname):
        raise HTTPException(400, "نام فایل نامعتبر")
    try:
        import export_utils, analysis_planner, stats_report
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌ها نصب نیستند: {_e}")
    path = os.path.join(export_utils.EXPORT_DIR, fname)
    if not os.path.exists(path):
        raise HTTPException(404, "فایل داده پیدا نشد؛ دوباره آپلود کنید")

    db = get_db()
    cur = db.execute(
        "INSERT INTO conversations (user_id, title) VALUES (?, ?)",
        (user["id"], f"تحلیل آماری: {(ask or 'خودکار')[:60]}"),
    )
    conv_id = cur.lastrowid
    db.execute("INSERT INTO messages (conversation_id, role, content) VALUES (?, 'user', ?)",
               (conv_id, f"📊 تحلیل خودکار داده — درخواست: {ask or '(بدون توضیح)'}"))
    db.commit(); db.close()

    async def gen():
        full_log = ""

        def log(t):
            nonlocal full_log
            full_log += t
            return t

        try:
            yield log("📊 **تحلیل خودکار داده**\n\n⏳ در حال شناخت ساختار داده…\n")
            prof = await run_in_threadpool(analysis_planner.profile, path)
            if prof.get("error"):
                yield log(f"\n⚠️ {prof['error']}")
                return
            yield log(f"✅ {prof['تعداد سطر']} مشاهده و {prof['تعداد ستون']} متغیر شناسایی شد.\n")
            if prof.get("ساختار تابلویی"):
                yield log("   نوع داده: تابلویی (پانل)\n")
            elif prof.get("سری‌زمانی است"):
                yield log("   نوع داده: سری‌زمانی\n")
            if prof.get("سازه‌های حدسی"):
                yield log(f"   سازه‌های احتمالی: {'، '.join(prof['سازه‌های حدسی'].keys())}\n")

            # ---------- نقشه‌ی تحلیل ----------
            yield log("\n⏳ انتخاب آزمون‌های مناسب…\n")
            c = pick_model_for("stats")
            steps, dropped = [], []
            if c.get("key"):
                try:
                    raw = await _call_model_once(
                        c, analysis_planner.build_prompt(prof, ask or "تحلیل متعارف این داده"),
                        system=analysis_planner.PLAN_SYSTEM, max_tokens=1400)
                    steps, dropped = analysis_planner.validate_plan(
                        analysis_planner.parse_plan(raw), prof)
                except Exception:
                    steps = []
            if not steps:
                yield log("   (نقشه‌ی پیش‌فرض بر پایه‌ی ساختار داده استفاده شد)\n")
                steps, dropped = analysis_planner.validate_plan(
                    analysis_planner.fallback_plan(prof), prof)
            if not steps:
                yield log("\n⚠️ برای این داده تحلیل مناسبی پیدا نشد. "
                          "لطفاً فایل داده‌ی جدولی با ستون‌های عددی بفرستید.")
                return
            for d in dropped[:4]:
                yield log(f"   ↷ «{d['تحلیل']}» کنار گذاشته شد: {d['دلیل']}\n")
            yield log(f"✅ {len(steps)} تحلیل انتخاب شد: "
                      f"{'، '.join(s['title'] for s in steps)}\n\n")

            # ---------- اجرا ----------
            results = []
            for i, st in enumerate(steps, 1):
                yield log(f"⏳ ({round((i-1)*100/len(steps))}٪) اجرای «{st['title']}»…\n")
                r = await run_in_threadpool(
                    lambda s=st: analysis_planner.execute(path, [s])[0])
                results.append(r)
                if r["result"].get("error"):
                    yield log(f"   ⚠️ {r['result']['error'][:120]}\n")
            yield log("\n✅ همه‌ی تحلیل‌ها اجرا شد.\n")

            # ---------- تفسیر ----------
            interpretation = ""
            if c.get("key"):
                yield log("\n⏳ نوشتن تفسیر دانشگاهی…\n")
                try:
                    digest = stats_report.summarize_for_model(results)
                    interpretation = await _call_model_once(
                        c,
                        "تو متخصص آمار و مشاور روش تحقیق هستی. این خروجی‌های واقعی تحلیل آماری "
                        f"روی داده‌ی کاربر است (درخواست کاربر: «{ask}»):\n\n{digest}\n\n"
                        "برای فصل چهارم رساله، به فارسی دانشگاهی و دقیق تفسیر کن: معناداری‌ها را "
                        "توضیح بده، فرضیه‌ها را تأیید یا رد کن، اندازه‌ی اثر را تفسیر کن و اگر "
                        "پیش‌فرضی نقض شده هشدار بده. عدد‌ها را از همین خروجی بردار و چیزی از خودت "
                        "نساز.",
                        system=BASE_SYSTEM_PROMPT, max_tokens=2000)
                    qdb = get_db(); quota_add(qdb, user, "chat", 4000); qdb.commit(); qdb.close()
                except Exception:
                    interpretation = ""

            # ---------- گزارش ----------
            md = stats_report.build_report(results, prof, ask, interpretation)
            yield log("\n---\n\n" + md + "\n")

            # ---------- نمودار و فایل خروجی ----------
            try:
                import sem_plot
                png = None
                pls = next((r for r in results if r["analysis"] == "pls_sem"
                            and r["result"].get("سازه‌ها")), None)
                if pls:
                    png = await run_in_threadpool(sem_plot.render_pls_diagram,
                                                  pls["result"], "مدل مسیر پژوهش")
                elif any(r["analysis"] == "correlation" for r in results):
                    png = await run_in_threadpool(sem_plot.render_corr_heatmap, path,
                                                  "pearson", None)
                if png:
                    yield log(f"\n![نمودار](/download/{png})\n")
            except Exception:
                pass

            if want_files:
                yield log("\n⏳ ساخت فایل خروجی…\n")
                try:
                    blocks = export_utils.md_to_blocks(md)
                    links = []
                    name = await run_in_threadpool(export_utils.build_docx, blocks,
                                                   "Vazirmatn", 14, "گزارش تحلیل آماری")
                    links.append(f"[📄 دانلود Word](/download/{name})")
                    try:
                        xn = await run_in_threadpool(export_utils.build_xlsx, blocks,
                                                     "Vazirmatn", 12, "گزارش تحلیل آماری")
                        links.append(f"[📊 دانلود Excel](/download/{xn})")
                    except Exception:
                        pass
                    yield log("\n" + "  |  ".join(links) + "\n")
                except Exception as e:
                    yield log(f"⚠️ ساخت فایل خروجی ممکن نشد: {e}\n")
        except Exception as e:
            yield log(f"\n⚠️ خطای غیرمنتظره در تحلیل خودکار: {e}")
        finally:
            d = get_db()
            d.execute("INSERT INTO messages (conversation_id, role, content) VALUES (?, 'assistant', ?)",
                      (conv_id, full_log or "…"))
            d.commit(); d.close()

    return StreamingResponse(gen(), media_type="text/plain; charset=utf-8",
                             headers={"X-Conversation-Id": str(conv_id)})


# ---------------- خروجی Word / PDF و سازنده مقاله بلند ----------------

FONT_CHOICES = ["Vazirmatn", "B Nazanin", "B Zar", "IRANSans", "B Titr", "Tahoma", "Times New Roman", "Calibri", "Arial"]

DESIGN_STYLES = ("auto", "tech", "medical", "finance", "education", "kids", "academic", "creative")


async def make_design_spec(topic: str, style: str = "auto", smart: bool = True):
    """طرح اختصاصی می‌سازد. اگر هوشمند و سبک=auto باشد، هوش مصنوعی بهترین سبک و پالت را انتخاب می‌کند."""
    import design_engine
    overrides = None
    if style not in DESIGN_STYLES:
        style = "auto"
    if smart and style == "auto":
        c = pick_model_for("article")
        if c.get("key"):
            prompt = (
                "بر اساس این موضوع، بهترین سبک طراحی سند/ارائه و یک پالت رنگ حرفه‌ای انتخاب کن. "
                "فقط JSON خام برگردان (بدون توضیح): "
                '{"style":"tech|medical|finance|education|kids|academic|creative",'
                '"primary":"#RRGGBB","secondary":"#RRGGBB","accent":"#RRGGBB"}\n'
                "راهنما: فناوری→مدرن تیره، پزشکی→سفید و آبی، مالی→رسمی مینیمال، آموزشی→رنگی، کودک→شاد، دانشگاهی→کلاسیک، بازاریابی→خلاق.\n\n"
                f"موضوع: {topic[:300]}"
            )
            try:
                out = await _call_model_once(c, prompt, max_tokens=160)
                m = re.search(r"\{.*\}", out or "", re.DOTALL)
                if m:
                    data = json.loads(m.group(0))
                    style = (data.get("style") or "auto").strip()
                    overrides = {k: data.get(k) for k in ("primary", "secondary", "accent") if data.get(k)}
            except Exception:
                pass
    return design_engine.build_spec(style, topic, overrides)


@app.post("/api/pptx")
async def api_pptx(request: Request):
    """تبدیل متن به ارائه پاورپوینت (با طراحی هوشمند اختصاصی در صورت درخواست)"""
    user = require_user(request)
    check_subscription(user)
    body = await request.json()
    content = (body.get("content") or "").strip()
    title = (body.get("title") or "ارائه آنتانو").strip()
    subtitle = (body.get("subtitle") or "").strip()
    smart_design = body.get("smart_design", True)
    style = (body.get("style") or "auto").strip()
    if not content:
        raise HTTPException(400, "متنی برای ساخت ارائه نیست")
    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
    if smart_design:
        try:
            import design_engine
            spec = await make_design_spec(title + " " + content[:200], style, True)
            name = await run_in_threadpool(design_engine.build_designed_pptx, content, spec, title, subtitle)
        except Exception as _e:
            print("[ANTANU] designed pptx failed, fallback:", _e)
            name = await run_in_threadpool(export_utils.build_pptx, content, title)
    else:
        name = await run_in_threadpool(export_utils.build_pptx, content, title)
    return {"files": [{"label": "📊 دانلود پاورپوینت", "url": f"/download/{name}"}]}


@app.post("/api/convert")
async def api_convert(request: Request, file: UploadFile = File(...), target: str = Form("docx")):
    """تبدیل سند: Word/PDF/PowerPoint ↔ یکدیگر (بر پایه استخراج متن و بازسازی)"""
    user = require_user(request)
    check_subscription(user)
    target = (target or "docx").lower().strip()
    if target not in ("docx", "pdf", "pptx"):
        raise HTTPException(400, "فرمت هدف باید docx، pdf یا pptx باشد")
    raw = await file.read()
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(400, "حجم فایل نباید بیشتر از ۲۵ مگابایت باشد")
    name = file.filename or "file"
    ext = os.path.splitext(name)[1].lower()
    if ext not in (".docx", ".pdf", ".pptx", ".txt", ".md"):
        raise HTTPException(400, "فرمت ورودی پشتیبانی‌شده: Word (.docx)، PDF، PowerPoint (.pptx)، متن")
    if ext.lstrip(".") == target:
        raise HTTPException(400, "فرمت ورودی و خروجی یکسان است")
    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند: pip install -r requirements.txt ({_e})")
    src = os.path.join(export_utils.EXPORT_DIR, f"src-{secrets.token_hex(6)}{ext}")
    with open(src, "wb") as f:
        f.write(raw)
    title = os.path.splitext(name)[0][:80]
    try:
        out_name, err = await run_in_threadpool(export_utils.convert_document, src, target,
                                                "Vazirmatn", 14, "justify", title)
    except Exception as e:
        raise HTTPException(400, f"تبدیل ناموفق بود: {e}")
    finally:
        try:
            os.remove(src)
        except OSError:
            pass
    if not out_name:
        raise HTTPException(400, err or "تبدیل ناموفق بود")
    labels = {"docx": "📄 دانلود Word", "pdf": "📕 دانلود PDF", "pptx": "📊 دانلود پاورپوینت"}
    return {"files": [{"label": labels.get(target, "دانلود"), "url": f"/download/{out_name}"}]}


@app.get("/download/{fname}")
def download_file(fname: str, request: Request):
    require_user(request)
    if not re.fullmatch(r"[A-Za-z0-9._-]+", fname):
        raise HTTPException(400, "نام فایل نامعتبر")
    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
    path = os.path.join(export_utils.EXPORT_DIR, fname)
    if not os.path.exists(path):
        raise HTTPException(404, "فایل پیدا نشد یا منقضی شده است")
    return FileResponse(path, filename=fname)


@app.post("/api/export")
async def api_export(request: Request):
    """تبدیل یک پاسخ به فایل Word / PDF با فونت و سایز دلخواه"""
    user = require_user(request)
    body = await request.json()
    content = (body.get("content") or "").strip()
    if not content:
        raise HTTPException(400, "متنی برای خروجی وجود ندارد")
    font = body.get("font") or "Vazirmatn"
    if font not in FONT_CHOICES:
        font = "Vazirmatn"
    try:
        size = max(8, min(int(body.get("size") or 14), 36))
    except (TypeError, ValueError):
        size = 14
    formats = body.get("formats") or ["docx"]
    align = body.get("align") if body.get("align") in ("right", "left", "center", "justify") else "right"
    title = (body.get("title") or "").strip() or None
    toc = bool(body.get("toc"))
    numbering = bool(body.get("numbering"))
    smart_design = bool(body.get("smart_design"))
    style = (body.get("style") or "auto").strip()
    subtitle = (body.get("subtitle") or "").strip()

    try:
        import export_utils
    except ImportError as _e:
        raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
    blocks = export_utils.md_to_blocks(content)
    files, notes = [], []
    design_spec = None
    if smart_design:
        design_spec = await make_design_spec((title or "") + " " + content[:200], style, True)

    async def _make_docx():
        """ساخت فایل Word (با طراحی هوشمند در صورت انتخاب) و برگرداندن نام آن."""
        if smart_design and design_spec:
            import design_engine
            return await run_in_threadpool(design_engine.build_designed_docx, blocks, design_spec,
                                           title, subtitle, size, align, toc, numbering)
        return await run_in_threadpool(export_utils.build_docx, blocks, font, size, title, align, toc, numbering)

    docx_name = None
    if "docx" in formats or "pdf" in formats:
        docx_name = await _make_docx()
    if "docx" in formats and docx_name:
        files.append({"label": "📄 دانلود Word", "url": f"/download/{docx_name}"})
    if "pdf" in formats:
        # PDF را از روی همان فایل Wordِ تمیز می‌سازیم تا ظاهرش دقیقاً مثل Word باشد
        pdf_name = None
        if docx_name:
            pdf_name, _cerr = await run_in_threadpool(export_utils.docx_to_pdf, docx_name)
        if not pdf_name:
            # اگر تبدیل Word→PDF ممکن نبود، به روش مستقیم برمی‌گردیم
            pdf_name, err = await run_in_threadpool(export_utils.build_pdf, blocks, size, title, align)
            if err:
                notes.append(err)
        if pdf_name:
            files.append({"label": "📕 دانلود PDF", "url": f"/download/{pdf_name}"})
    if "xlsx" in formats:
        name = await run_in_threadpool(export_utils.build_xlsx, blocks, font, size, title, align)
        files.append({"label": "📊 دانلود Excel", "url": f"/download/{name}"})
    if "txt" in formats:
        name = await run_in_threadpool(export_utils.build_txt, blocks, title)
        files.append({"label": "📃 دانلود متن (TXT)", "url": f"/download/{name}"})
    if "md" in formats:
        name = await run_in_threadpool(export_utils.build_md, blocks, title)
        files.append({"label": "📝 دانلود Markdown", "url": f"/download/{name}"})
    return {"files": files, "notes": notes}


async def _call_model_once(c, prompt: str | None = None, system: str | None = None,
                           max_tokens: int = 1800, messages=None,
                           keep_foreign: bool = False) -> str:
    """یک فراخوانی بدون استریم — با دو تلاش مجدد در صورت شلوغی.
    keep_foreign=True برای ترجمه: خروجی عمداً غیرفارسی است و نباید فیلتر شود."""
    if messages is None:
        messages = ([{"role": "system", "content": system}] if system else []) + [
            {"role": "user", "content": prompt}
        ]
    payload = {"model": c["model"], "messages": messages, "max_tokens": max_tokens}
    headers = {"Authorization": f"Bearer {c['key']}", "Content-Type": "application/json"}
    for attempt in range(3):
        async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=15)) as client:
            r = await client.post(f"{c['base']}/chat/completions", json=payload, headers=headers)
        if r.status_code == 200:
            data = r.json()
            msg_obj = (data.get("choices") or [{}])[0].get("message") or {}
            raw_out = msg_obj.get("content", "") or ""
            out = raw_out if keep_foreign else clean_foreign_keep_tr(raw_out)
            for im in (msg_obj.get("images") or []):
                url = ((im or {}).get("image_url") or {}).get("url") or ""
                if url.startswith("data:image"):
                    fname = _save_gen_image(url)
                    if fname:
                        out += f"\n\n[[ANTANU_IMG:/download/{fname}]]\n\n"
            return out
        if r.status_code == 429 and attempt < 2:
            await asyncio.sleep(20)
            continue
        raise ModelError(r.status_code, r.text[:300])
    return ""


@app.post("/api/longdoc")
async def api_longdoc(request: Request):
    """سازنده مقاله بلند: فهرست بخش‌ها → نوشتن بخش‌به‌بخش → خروجی Word/PDF"""
    user = require_user(request)
    body = await request.json()
    topic = (body.get("topic") or "").strip()
    if not topic:
        raise HTTPException(400, "موضوع مقاله را بنویسید")
    try:
        pages = max(1, min(int(body.get("pages") or 10), 500))
    except (TypeError, ValueError):
        pages = 10
    font = body.get("font") or "Vazirmatn"
    if font not in FONT_CHOICES:
        font = "Vazirmatn"
    try:
        size = max(8, min(int(body.get("size") or 14), 36))
    except (TypeError, ValueError):
        size = 14
    formats = body.get("formats") or ["docx"]
    align = body.get("align") if body.get("align") in ("right", "left", "center", "justify") else "right"
    attachment_ids = body.get("attachments") or []
    use_web = bool(body.get("use_web"))
    ld_smart = bool(body.get("smart_design"))
    ld_style = (body.get("style") or "auto").strip()

    c = pick_model_for("article")
    if not c.get("key"):
        raise HTTPException(400, "ابتدا در پنل مدیریت، کلید API را تنظیم کنید")

    # سهمیه مقاله بلند
    dbq = get_db()
    try:
        quota_check(dbq, user, "article")
        quota_add(dbq, user, "article")
        dbq.commit()
    finally:
        dbq.close()

    # منابع کاربر: فایل‌های متنی و عکس‌ها
    source_text = ""
    source_images = []
    if attachment_ids:
        db0 = get_db()
        for aid in attachment_ids[:6]:
            row = db0.execute(
                "SELECT filename, content, kind FROM uploads WHERE id = ? AND user_id = ?",
                (int(aid), user["id"]),
            ).fetchone()
            if not row:
                continue
            if row["kind"] == "image":
                source_images.append(row["content"])
            else:
                source_text += f"\n\n[منبع: {row['filename']}]\n{row['content'][:12000]}"
        db0.close()

    n_sections = max(3, min(pages // 2 + 1, 250))
    sys_prompt = BASE_SYSTEM_PROMPT

    # ثبت در تاریخچه گفتگوها
    db = get_db()
    cur = db.execute(
        "INSERT INTO conversations (user_id, title) VALUES (?, ?)",
        (user["id"], f"📄 مقاله: {topic[:45]}"),
    )
    conv_id = cur.lastrowid
    db.execute(
        "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'user', ?)",
        (conv_id, f"📄 درخواست مقاله {pages} صفحه‌ای درباره: {topic}"),
    )
    db.commit()
    db.close()

    async def gen():
        full_log = ""

        def log(t):
            nonlocal full_log
            full_log += t
            return t

        try:
            yield log(f"📄 **ساخت مقاله «{topic}» — حدود {pages} صفحه**\n\n")

            # گام ۰: مطالعه منابع کاربر (فایل، عکس، وب)
            digest = ""
            if source_images:
                yield log("⏳ گام ۰: بررسی عکس‌های شما…\n")
                try:
                    parts = [{"type": "text", "text":
                              "این تصاویر منبع یک مقاله هستند. هر تصویر را با جزئیات کامل به فارسی توصیف کن "
                              "و هر متن، عدد یا نموداری که در آن هست را بنویس."}]
                    for b64 in source_images[:4]:
                        parts.append({"type": "image_url",
                                      "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})
                    desc = await _call_model_once(c, messages=[{"role": "user", "content": parts}], max_tokens=1000)
                    digest += "\n\n[توصیف عکس‌های کاربر]:\n" + desc
                except ModelError:
                    yield log("⚠️ این مدل عکس را پشتیبانی نمی‌کند؛ عکس‌ها نادیده گرفته شدند.\n")
            if source_text:
                digest += source_text
            if use_web:
                yield log("⏳ جستجوی وب برای منابع…\n")
                try:
                    results = await run_in_threadpool(_search_web_sync, topic[:300])
                    if results:
                        digest += "\n\n[نتایج جستجوی وب]:\n" + results
                except Exception:
                    pass
            if len(digest) > 6000:
                yield log("⏳ خلاصه‌سازی منابع…\n")
                try:
                    digest = await _call_model_once(
                        c,
                        f"این منابع را برای نگارش مقاله‌ای درباره «{topic}» در حداکثر ۱۰۰۰ کلمه فارسی خلاصه کن "
                        f"و همه داده‌ها و نکات مهم را نگه دار:\n{digest[:18000]}",
                        system=sys_prompt, max_tokens=1600,
                    )
                except ModelError:
                    digest = digest[:6000]
            src_note = (f"\n\nمنابع کاربر (حتماً مبنای مقاله قرار بده):\n{digest[:5000]}" if digest.strip() else "")

            # کتابخانه‌ی دائمی آنتانو: کتاب‌ها و جزوه‌های ذخیره‌شده درباره‌ی همین موضوع
            lib_note = _library_context(topic, limit=4, budget=3500)
            if lib_note:
                yield log("📚 منابع مرتبط از کتابخانه‌ی آنتانو پیدا شد.\n")

            yield log("⏳ گام ۱: طراحی فهرست بخش‌ها…\n")
            outline = await _call_model_once(
                c,
                f"برای یک مقاله علمی-پژوهشی {pages} صفحه‌ای فارسی درباره «{topic}» دقیقاً {n_sections} عنوان بخش بنویس. "
                "ساختار باید استاندارد مقاله دانشگاهی باشد: با چکیده و مقدمه شروع شود، سپس مبانی نظری و پیشینه پژوهش "
                "(داخلی و خارجی)، روش‌شناسی پژوهش، یافته‌ها و تحلیل داده‌ها، بحث و نتیجه‌گیری، و در پایان منابع. "
                "هر عنوان در یک خط جداگانه، بدون شماره و بدون توضیح اضافه. عنوان‌ها متنوع و بدون هم‌پوشانی باشند."
                + src_note + lib_note,
                system=sys_prompt,
                max_tokens=1200,
            )
            titles = [re.sub(r"^[\d\-.،*#)\s]+", "", t).strip() for t in outline.split("\n")]
            titles = [t for t in titles if 2 < len(t) < 120][:n_sections]
            if not titles:
                yield log("\n⚠️ فهرست بخش‌ها ساخته نشد. دوباره تلاش کنید.")
                return

            yield log(f"✅ {len(titles)} بخش طراحی شد.\n\n")
            article = f"# {topic}\n"
            done_titles = []

            for i, t in enumerate(titles, 1):
                if await request.is_disconnected():
                    yield log("\n⏹ ساخت مقاله توسط کاربر متوقف شد.")
                    break
                pct = round((i - 1) * 100 / len(titles))
                yield log(f"⏳ ({pct}٪) نوشتن بخش {i} از {len(titles)}: «{t}»…\n")
                # برای هر بخش، مرتبط‌ترین صفحه‌های کتابخانه به همان بخش را می‌آوریم
                sec_lib = _library_context(f"{topic} {t}", limit=3, budget=2600) or lib_note
                try:
                    part = await _call_model_once(
                        c,
                        f"مقاله‌ای فارسی درباره «{topic}» در حال نگارش است.\n"
                        f"بخش‌های نوشته‌شده تاکنون: {'، '.join(done_titles) if done_titles else 'هیچ'}.\n"
                        f"اکنون فقط بخش «{t}» را بنویس: حدود ۶۰۰ تا ۸۰۰ کلمه، علمی و ساختارمند. "
                        "از تکرار مطالب و واژه‌های بخش‌های قبلی جداً پرهیز کن و مطالب و واژگان کاملاً تازه بیاور. "
                        "فقط به فارسی معیار بنویس و هیچ واژه خارجی وسط متن نیاور. "
                        "خودِ عنوان بخش را ننویس؛ فقط متن."
                        + src_note + sec_lib,
                        system=sys_prompt,
                    )
                except ModelError as e:
                    yield log(f"⚠️ بخش «{t}» به دلیل خطای سرویس (کد {e.status}) رد شد.\n")
                    continue
                article += f"\n\n## {t}\n\n{part.strip()}"
                done_titles.append(t)
                await asyncio.sleep(1)

            yield log("\n⏳ گام پایانی: ساخت فایل‌ها…\n")
            try:
                import export_utils
            except ImportError as _e:
                raise HTTPException(500, f"کتابخانه‌های ساخت فایل نصب نیستند. دستور را اجرا کنید: pip install -r requirements.txt (جزئیات: {_e})")
            blocks = export_utils.md_to_blocks(article)
            links = []

            async def _make_article_docx():
                # مقاله‌ی بلند: فهرست خودکار + شماره‌گذاری سرفصل‌ها (سبک پایان‌نامه)
                if ld_smart:
                    import design_engine
                    spec = await make_design_spec(topic, ld_style, True)
                    return await run_in_threadpool(design_engine.build_designed_docx, blocks, spec,
                                                   topic, "", size, align, True, True)
                return await run_in_threadpool(export_utils.build_docx, blocks, font, size, topic, align, True, True)

            article_docx = None
            if "docx" in formats or "pdf" in formats:
                article_docx = await _make_article_docx()
            if "docx" in formats and article_docx:
                links.append(f"[📄 دانلود Word](/download/{article_docx})")
            if "pdf" in formats:
                # PDF از روی همان Word تمیز ساخته می‌شود تا ظاهرش یکسان باشد
                pdf_name = None
                if article_docx:
                    pdf_name, _ = await run_in_threadpool(export_utils.docx_to_pdf, article_docx)
                if not pdf_name:
                    pdf_name, err = await run_in_threadpool(export_utils.build_pdf, blocks, size, topic, align)
                    if err:
                        yield log(f"⚠️ {err}\n")
                if pdf_name:
                    links.append(f"[📕 دانلود PDF](/download/{pdf_name})")
            if "xlsx" in formats:
                name = await run_in_threadpool(export_utils.build_xlsx, blocks, font, size, topic, align)
                links.append(f"[📊 دانلود Excel](/download/{name})")
            if "txt" in formats:
                name = await run_in_threadpool(export_utils.build_txt, blocks, topic)
                links.append(f"[📃 دانلود متن](/download/{name})")
            if "md" in formats:
                name = await run_in_threadpool(export_utils.build_md, blocks, topic)
                links.append(f"[📝 دانلود Markdown](/download/{name})")
            yield log(f"\n✅ **مقاله آماده شد!** ({len(done_titles)} بخش)\n\n" + "  |  ".join(links))
        except ModelError as e:
            yield log(f"\n⚠️ سرویس هوش مصنوعی خطا داد (کد {e.status}). کمی بعد دوباره تلاش کنید.")
        except Exception:
            yield log("\n⚠️ خطای غیرمنتظره در ساخت مقاله.")
        finally:
            d = get_db()
            d.execute(
                "INSERT INTO messages (conversation_id, role, content) VALUES (?, 'assistant', ?)",
                (conv_id, full_log or "…"),
            )
            d.commit()
            d.close()

    return StreamingResponse(
        gen(),
        media_type="text/plain; charset=utf-8",
        headers={"X-Conversation-Id": str(conv_id)},
    )


# ---------------- API پنل مدیریت ----------------

@app.get("/admin/app_settings")
def admin_get_app_settings(request: Request):
    require_admin(request)
    return {"version": get_setting("app_version", "1.0"),
            "announcement": get_setting("announcement", ""),
            "ticker_text": get_setting("ticker_text", ""),
            "ticker_speed": get_setting("ticker_speed", "25")}


@app.post("/admin/app_settings")
async def admin_save_app_settings(request: Request):
    require_admin(request)
    body = await request.json()
    set_setting("app_version", (body.get("version") or "1.0").strip()[:20])
    set_setting("announcement", (body.get("announcement") or "").strip()[:300])
    set_setting("ticker_text", (body.get("ticker_text") or "").strip()[:600])
    try:
        spd = max(5, min(int(body.get("ticker_speed") or 25), 180))
    except (TypeError, ValueError):
        spd = 25
    set_setting("ticker_speed", str(spd))
    return {"ok": True}


@app.get("/admin/library_stats")
def admin_library_stats(request: Request):
    """آمار «کتابخانه‌ی دائمی»: چند سند و چند قطعه متن ذخیره شده."""
    require_admin(request)
    try:
        import knowledge_base as _kb
        return _kb.stats()
    except Exception:
        return {"sources": 0, "chunks": 0, "items": []}


@app.get("/admin/build_info")
def admin_get_build_info(request: Request):
    """نسخه و وضعیت ساختی که همین حالا روی سرور اجرا می‌شود.
    برای اینکه بعد از هر دیپلوی بشود مطمئن شد فایل‌های تازه واقعاً نشسته‌اند."""
    require_admin(request)
    return _build_report()


@app.get("/admin/ai_settings")
def admin_get_ai_settings(request: Request):
    require_admin(request)
    db = get_db()
    row = db.execute("SELECT value FROM settings WHERE key = 'ai_config'").fetchone()
    db.close()
    if row:
        try:
            return json.loads(row["value"])
        except json.JSONDecodeError:
            pass
    return {"provider": PROVIDER, "api_key": "", "model": "", "ais": []}


@app.post("/admin/ai_settings")
async def admin_save_ai_settings(request: Request):
    require_admin(request)
    body = await request.json()
    cfg = {
        "provider": (body.get("provider") or "groq").strip(),
        "api_key": (body.get("api_key") or "").strip(),
        "model": (body.get("model") or "").strip(),
        "main_role": (body.get("main_role") or "general").strip().lower(),
        "ais": [
            {
                "name": (a.get("name") or "").strip(),
                "service": (a.get("service") or "").strip(),
                "model": (a.get("model") or "").strip(),
                "key": (a.get("key") or "").strip(),
                "role": (a.get("role") or "").strip().lower(),
            }
            for a in (body.get("ais") or [])[:20]
        ],
    }
    db = get_db()
    db.execute(
        "INSERT INTO settings (key, value) VALUES ('ai_config', ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (json.dumps(cfg, ensure_ascii=False),),
    )
    db.commit()
    db.close()
    return {"ok": True, "models": len(get_ai_catalog())}


@app.post("/admin/test_ai")
async def admin_test_ai(request: Request):
    """تست زنده یک کلید — یک پیام کوتاه به سرویس می‌فرستد"""
    require_admin(request)
    body = await request.json()
    base = resolve_base(body.get("service"))
    model = (body.get("model") or "").strip()
    key = (body.get("key") or "").strip()
    if not (base and model and key):
        return {"ok": False, "msg": "سرویس، نام مدل و کلید را کامل وارد کنید"}
    payload = {"model": model, "messages": [{"role": "user", "content": "سلام"}], "max_tokens": 10}
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(30, connect=10)) as client:
            r = await client.post(
                f"{base}/chat/completions",
                json=payload,
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
            )
        if r.status_code == 200:
            return {"ok": True, "msg": "✅ کلید سالم است و مدل پاسخ داد"}
        if r.status_code in (401, 403):
            return {"ok": False, "msg": "❌ کلید نامعتبر یا منقضی است"}
        if r.status_code == 404:
            return {"ok": False, "msg": "❌ نام مدل در این سرویس پیدا نشد"}
        if r.status_code == 429:
            return {"ok": False, "msg": "⚠️ کلید درست است ولی ظرفیت رایگان فعلاً پر است"}
        if r.status_code >= 500:
            return {"ok": False, "msg": f"⚠️ سرویس موقتاً پاسخ نمی‌دهد (کد {r.status_code}) — چند دقیقه بعد دوباره تست کنید یا نام مدل دیگری بنویسید"}
        return {"ok": False, "msg": f"❌ خطای سرویس (کد {r.status_code})"}
    except Exception:
        return {"ok": False, "msg": "❌ اتصال به سرویس برقرار نشد"}


@app.get("/admin/knowledge_stats")
def admin_kb_stats(request: Request):
    require_admin(request)
    db = get_db()
    n = db.execute("SELECT COUNT(*) AS c FROM knowledge").fetchone()["c"]
    db.close()
    return {"count": n}


@app.post("/admin/knowledge_clear")
def admin_kb_clear(request: Request):
    require_admin(request)
    db = get_db()
    db.execute("DELETE FROM knowledge")
    db.commit()
    db.close()
    return {"ok": True}


@app.get("/admin/feedback")
def admin_feedback_list(request: Request):
    """فهرست ایده‌ها و نظرات کاربران برای مدیر."""
    require_admin(request)
    db = get_db()
    rows = db.execute(
        "SELECT id, username, category, content, status, created_at "
        "FROM feedback ORDER BY (status='new') DESC, id DESC LIMIT 300"
    ).fetchall()
    new_count = db.execute("SELECT COUNT(*) AS c FROM feedback WHERE status='new'").fetchone()["c"]
    db.close()
    return {"items": [dict(r) for r in rows], "new_count": new_count}


@app.post("/admin/feedback_status")
async def admin_feedback_status(request: Request):
    """به‌روزرسانی وضعیت یک نظر (seen/done) یا حذف آن."""
    require_admin(request)
    body = await request.json()
    fid = int(body.get("id", 0))
    action = (body.get("action") or "").strip()
    db = get_db()
    if action == "delete":
        db.execute("DELETE FROM feedback WHERE id = ?", (fid,))
    elif action in ("seen", "done", "new"):
        db.execute("UPDATE feedback SET status = ? WHERE id = ?", (action, fid))
    db.commit()
    db.close()
    return {"ok": True}


@app.get("/admin/analytics")
def admin_analytics(request: Request):
    """آمار مصرف: کاربران فعال، روند روزانه و پرمصرف‌ترین کاربران (بر پایه usage_log)"""
    require_admin(request)
    db = get_db()

    def iran_today_offset(days_ago: int = 0) -> str:
        row = db.execute(
            "SELECT date('now', '+3 hours', '+30 minutes', ?) AS d", (f"-{days_ago} days",)
        ).fetchone()
        return row["d"]

    today = iran_today_offset(0)

    today_active = db.execute(
        "SELECT COUNT(DISTINCT user_id) AS c FROM usage_log "
        "WHERE date(created_at, '+3 hours', '+30 minutes') = ?", (today,),
    ).fetchone()["c"]

    week_active = db.execute(
        "SELECT COUNT(DISTINCT user_id) AS c FROM usage_log "
        "WHERE date(created_at, '+3 hours', '+30 minutes') >= ?", (iran_today_offset(6),),
    ).fetchone()["c"]

    total_users = db.execute("SELECT COUNT(*) AS c FROM users WHERE is_admin = 0").fetchone()["c"]

    rows = db.execute(
        """SELECT date(created_at, '+3 hours', '+30 minutes') AS d,
                  SUM(CASE WHEN kind='chat' THEN 1 ELSE 0 END) AS messages,
                  SUM(CASE WHEN kind='image' THEN amount ELSE 0 END) AS images,
                  SUM(CASE WHEN kind='video' THEN amount ELSE 0 END) AS videos,
                  SUM(CASE WHEN kind='article' THEN amount ELSE 0 END) AS articles
           FROM usage_log
           WHERE date(created_at, '+3 hours', '+30 minutes') >= ?
           GROUP BY d ORDER BY d""",
        (iran_today_offset(13),),
    ).fetchall()
    by_day = {r["d"]: dict(r) for r in rows}
    daily = []
    for i in range(13, -1, -1):
        d = iran_today_offset(i)
        r = by_day.get(d)
        daily.append({
            "date": d,
            "messages": (r["messages"] if r else 0) or 0,
            "images": (r["images"] if r else 0) or 0,
            "videos": (r["videos"] if r else 0) or 0,
            "articles": (r["articles"] if r else 0) or 0,
        })

    today_messages = by_day.get(today, {}).get("messages") or 0

    top_users = db.execute(
        """SELECT u.username AS username, u.stars AS stars,
                  SUM(CASE WHEN l.kind='chat' THEN 1 ELSE 0 END) AS messages,
                  SUM(CASE WHEN l.kind='image' THEN l.amount ELSE 0 END) AS images,
                  SUM(CASE WHEN l.kind='video' THEN l.amount ELSE 0 END) AS videos,
                  SUM(CASE WHEN l.kind='article' THEN l.amount ELSE 0 END) AS articles,
                  MAX(l.created_at) AS last_active
           FROM usage_log l JOIN users u ON u.id = l.user_id
           WHERE l.created_at >= datetime('now', '-30 days')
           GROUP BY u.id
           ORDER BY messages DESC
           LIMIT 10""",
    ).fetchall()
    db.close()

    return {
        "today_active": today_active,
        "week_active": week_active,
        "total_users": total_users,
        "today_messages": today_messages,
        "daily": daily,
        "top_users": [dict(u) for u in top_users],
    }


@app.get("/admin/backup")
def admin_backup(request: Request):
    """دانلود نسخه‌ی پشتیبان کامل پایگاه داده (کاربران، کدها، گفتگوها و…)"""
    require_admin(request)
    if not os.path.exists(DB_PATH):
        raise HTTPException(404, "فایل پایگاه داده پیدا نشد")
    # نسخه‌ی سازگار با SQLite (حتی هنگام باز بودن) با API رسمی backup ساخته می‌شود
    import sqlite3
    import datetime as _d
    try:
        import export_utils
        out_dir = export_utils.EXPORT_DIR
    except ImportError:
        out_dir = "."
    stamp = _d.datetime.now().strftime("%Y%m%d-%H%M%S")
    out_name = f"antanu-backup-{stamp}.db"
    out_path = os.path.join(out_dir, out_name)
    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(out_path)
    try:
        with dst:
            src.backup(dst)
    finally:
        src.close()
        dst.close()
    return FileResponse(out_path, filename=out_name, media_type="application/octet-stream")


@app.get("/admin/telegram_settings")
def admin_telegram_get(request: Request):
    """خواندن تنظیمات پشتیبان‌گیری تلگرام (توکن به‌صورت پنهان برنمی‌گردد)."""
    require_admin(request)
    token = get_setting("tg_bot_token", "") or os.environ.get("ANTANU_TG_BOT_TOKEN", "")
    chat = get_setting("tg_backup_chat", "") or os.environ.get("ANTANU_TG_BACKUP_CHAT", "")
    return {"has_token": bool(token.strip()), "chat": chat}


@app.post("/admin/telegram_settings")
async def admin_telegram_set(request: Request):
    """ذخیره‌ی توکن ربات و آیدی کانال تلگرام از پنل مدیریت (بدون نیاز به .env)."""
    require_admin(request)
    body = await request.json()
    token = (body.get("token") or "").strip()
    chat = (body.get("chat") or "").strip()
    # اگر توکن خالی فرستاده شد، توکن قبلی را نگه دار (تا با ذخیره‌ی دوباره پاک نشود)
    if token:
        set_setting("tg_bot_token", token)
    set_setting("tg_backup_chat", chat)
    return {"ok": True, "message": "✅ تنظیمات تلگرام ذخیره شد."}


@app.get("/admin/tts_settings")
def admin_tts_get(request: Request):
    """خواندن تنظیمات صدای حرفه‌ای (کلید پنهان برنمی‌گردد)."""
    require_admin(request)
    key = get_setting("elevenlabs_key", "") or os.environ.get("ELEVENLABS_API_KEY", "")
    return {
        "has_key": bool(key.strip()),
        "voice_self": get_setting("tts_voice_self", ""),
        "voice_female": get_setting("tts_voice_female", ""),
        "voice_male": get_setting("tts_voice_male", ""),
    }


@app.post("/admin/tts_settings")
async def admin_tts_set(request: Request):
    """ذخیره‌ی کلید ElevenLabs و شناسه‌ی صداها (آنتانو/خانم/آقا) از پنل مدیریت."""
    require_admin(request)
    body = await request.json()
    key = (body.get("key") or "").strip()
    if key:
        set_setting("elevenlabs_key", key)
    set_setting("tts_voice_self", (body.get("voice_self") or "").strip())
    set_setting("tts_voice_female", (body.get("voice_female") or "").strip())
    set_setting("tts_voice_male", (body.get("voice_male") or "").strip())
    return {"ok": True, "message": "✅ تنظیمات صدای حرفه‌ای ذخیره شد."}


@app.get("/admin/stt_settings")
def admin_stt_get(request: Request):
    """خواندن تنظیمات تبدیل صدا به متن (کلید پنهان برنمی‌گردد)."""
    require_admin(request)
    key = get_setting("stt_key", "") or os.environ.get("ANTANU_STT_KEY", "")
    return {
        "has_key": bool(key.strip()),
        "base": get_setting("stt_base", "") or "https://api.groq.com/openai/v1",
        "model": get_setting("stt_model", "") or "whisper-large-v3",
    }


@app.post("/admin/stt_settings")
async def admin_stt_set(request: Request):
    """ذخیره‌ی کلید، آدرس و مدل سرویس تبدیل صدا به متن (Whisper) از پنل مدیریت."""
    require_admin(request)
    body = await request.json()
    key = (body.get("key") or "").strip()
    if key:
        set_setting("stt_key", key)
    set_setting("stt_base", (body.get("base") or "").strip())
    set_setting("stt_model", (body.get("model") or "").strip())
    return {"ok": True, "message": "✅ تنظیمات تبدیل صدا به متن ذخیره شد."}


@app.get("/admin/services_settings")
def admin_services_get(request: Request):
    """خواندن کلیدهای سرویس‌های جانبی (آهنگ/ویدیو/عکس/تماس). کلیدها پنهان برمی‌گردند."""
    require_admin(request)
    return {
        "has_replicate": bool((get_setting("replicate_key", "") or os.environ.get("REPLICATE_API_TOKEN", "")).strip()),
        "has_music": bool(get_setting("music_key", "").strip()),
        "music_base": get_setting("music_base", ""),
        "has_image": bool(get_setting("image_key", "").strip()),
        "image_base": get_setting("image_base", ""),
        "twilio_sid": get_setting("twilio_sid", ""),
        "has_twilio_token": bool(get_setting("twilio_token", "").strip()),
        "twilio_phone": get_setting("twilio_phone", ""),
        "phone_greeting": get_setting("phone_greeting", ""),
        "animal_sounds_text": _animal_map_to_text(get_setting("animal_sounds", "")),
    }


@app.post("/admin/services_settings")
async def admin_services_set(request: Request):
    """ذخیره‌ی کلیدهای سرویس‌های جانبی از پنل مدیریت (کلید خالی = بدون تغییر)."""
    require_admin(request)
    body = await request.json()

    def _save_key(field, key):
        v = (body.get(field) or "").strip()
        if v:
            set_setting(key, v)

    _save_key("replicate_key", "replicate_key")
    _save_key("music_key", "music_key")
    _save_key("image_key", "image_key")
    _save_key("twilio_token", "twilio_token")
    # فیلدهای غیرمحرمانه همیشه ذخیره می‌شوند
    set_setting("music_base", (body.get("music_base") or "").strip())
    set_setting("image_base", (body.get("image_base") or "").strip())
    set_setting("twilio_sid", (body.get("twilio_sid") or "").strip())
    set_setting("twilio_phone", (body.get("twilio_phone") or "").strip())
    set_setting("phone_greeting", (body.get("phone_greeting") or "").strip())
    if "animal_sounds_text" in body:
        set_setting("animal_sounds", _animal_text_to_map(body.get("animal_sounds_text") or ""))
    return {"ok": True, "message": "✅ کلیدهای سرویس‌های جانبی ذخیره شد."}


@app.post("/admin/backup_telegram")
async def admin_backup_telegram(request: Request):
    """همین حالا یک پشتیبان بساز و به کانال تلگرام بفرست (برای تست تنظیمات)."""
    require_admin(request)
    path = await run_in_threadpool(_backup_db_now)
    ok, msg = await run_in_threadpool(_send_backup_to_telegram, path)
    return {"ok": ok, "message": msg}


@app.get("/admin/data")
def admin_data(request: Request):
    require_admin(request)
    db = get_db()
    codes = db.execute("SELECT * FROM codes ORDER BY id DESC LIMIT 500").fetchall()
    users = db.execute(
        "SELECT id, username, stars, is_admin, device_fp, code_used, created_at FROM users ORDER BY id DESC"
    ).fetchall()
    db.close()
    return {"codes": [dict(c) for c in codes], "users": [dict(u) for u in users]}


@app.post("/admin/codes")
async def admin_generate_codes(request: Request):
    require_admin(request)
    body = await request.json()
    stars = int(body.get("stars", 1))
    count = max(1, min(int(body.get("count", 1)), 200))
    if stars not in (1, 2, 3, 4, 5):
        raise HTTPException(400, "ستاره باید بین ۱ تا ۵ باشد")
    db = get_db()
    new_codes = []
    for _ in range(count):
        code = generate_code()
        db.execute("INSERT INTO codes (code, stars) VALUES (?, ?)", (code, stars))
        new_codes.append(code)
    db.commit()
    db.close()
    return {"codes": new_codes, "stars": stars}


@app.post("/admin/delete_code")
async def admin_delete_code(request: Request):
    """حذف یک کد ثبت‌نام (یا همه‌ی کدهای استفاده‌نشده) توسط مدیر"""
    require_admin(request)
    body = await request.json()
    db = get_db()
    if body.get("all_unused"):
        cur = db.execute("DELETE FROM codes WHERE used = 0")
        n = cur.rowcount
    else:
        code_id = body.get("code_id")
        if code_id is None:
            db.close()
            raise HTTPException(400, "شناسه کد مشخص نشده")
        cur = db.execute("DELETE FROM codes WHERE id = ?", (int(code_id),))
        n = cur.rowcount
    db.commit()
    db.close()
    return {"ok": True, "deleted": n}


@app.post("/admin/create_user")
async def admin_create_user(request: Request):
    """ساخت دستی کاربر توسط ادمین (بدون نیاز به کد)"""
    require_admin(request)
    body = await request.json()
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    stars = int(body.get("stars", 1))
    if len(username) < 3 or len(password) < 6 or stars not in (1, 2, 3, 4):
        raise HTTPException(400, "اطلاعات نامعتبر (نام ≥ ۳ حرف، گذرواژه ≥ ۶ حرف، ستاره ۱ تا ۴)")
    db = get_db()
    if db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
        db.close()
        raise HTTPException(400, "این نام کاربری وجود دارد")
    db.execute(
        "INSERT INTO users (username, password, stars, expires_at) "
        "VALUES (?, ?, ?, datetime('now', '+30 days'))",
        (username, hash_pw(password), stars),
    )
    db.commit()
    db.close()
    return {"ok": True}


@app.post("/admin/reset_device")
async def admin_reset_device(request: Request):
    """آزاد کردن قفل دستگاه یک کاربر (برای انتقال به گوشی/کامپیوتر جدید)"""
    require_admin(request)
    body = await request.json()
    db = get_db()
    db.execute("UPDATE users SET device_fp = NULL WHERE id = ?", (int(body["user_id"]),))
    db.commit()
    db.close()
    return {"ok": True}


@app.post("/admin/set_stars")
async def admin_set_stars(request: Request):
    require_admin(request)
    body = await request.json()
    stars = int(body.get("stars", 1))
    if stars not in (1, 2, 3, 4, 5):
        raise HTTPException(400, "ستاره باید بین ۱ تا ۵ باشد")
    db = get_db()
    db.execute(
        "UPDATE users SET stars = ?, expires_at = datetime('now', '+30 days') "
        "WHERE id = ? AND is_admin = 0",
        (stars, int(body["user_id"])),
    )
    db.commit()
    db.close()
    return {"ok": True}


@app.post("/admin/reset_password")
async def admin_reset_password(request: Request):
    """ساخت رمز جدید برای کاربری که رمزش را فراموش کرده — ادمین رمز را به او می‌دهد"""
    require_admin(request)
    body = await request.json()
    temp = "".join(secrets.choice("ABCDEFGHJKMNPQRSTUVWXYZ23456789") for _ in range(8))
    db = get_db()
    db.execute("UPDATE users SET password = ? WHERE id = ? AND is_admin = 0",
               (hash_pw(temp), int(body["user_id"])))
    db.execute("DELETE FROM sessions WHERE user_id = ?", (int(body["user_id"]),))
    db.commit()
    db.close()
    return {"password": temp}


@app.post("/admin/delete_user")
async def admin_delete_user(request: Request):
    require_admin(request)
    body = await request.json()
    db = get_db()
    db.execute("DELETE FROM users WHERE id = ? AND is_admin = 0", (int(body["user_id"]),))
    db.commit()
    db.close()
    return {"ok": True}
